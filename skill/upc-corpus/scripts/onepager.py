#!/usr/bin/env python3
"""A one-pager: one printable page for one audience, from a UPC corpus, where every number on
the page is a figure whose passage is stored in the corpus and checked again at every build.

    python3 onepager.py init    --corpus <dir> --slug S --audience <a.yaml> --project <p.yaml> [--title T]
                                [--layout executive|technical|policy] [--distribution internal|partner|public]
                                [--date YYYY-MM-DD] [--show-writer-as WRITER=LABEL]
    python3 onepager.py figures add    --corpus <dir> --from <figures.json> [--by "<who>"] [--replace] [--plan]
    python3 onepager.py figures import --corpus <dir> --from-dossier [--plan]
    python3 onepager.py figures list   --corpus <dir> [--slug S]
    python3 onepager.py code    --corpus <dir> --slug S [--figure ID ...] [--redo] [--plan]
    python3 onepager.py review  --corpus <dir> --slug S --out <sheet.csv|.xlsx|.md>
    python3 onepager.py review  --corpus <dir> --slug S --apply <sheet.csv|.xlsx> --coder <handle>
    python3 onepager.py reply   --corpus <dir> --from <reply.md|.txt> --by <handle> [--title T] [--date D] [--plan]
    python3 onepager.py signoff --corpus <dir> --figure ID [--figure ID ...] --by <handle>
                                (--check confirmed|corrected|query | --clearance cleared|internal-only|withheld)
                                [--ref <source>] [--note "<text>"]
    python3 onepager.py write   --corpus <dir> --slug S [--plan]
    python3 onepager.py add     --corpus <dir> --slug S --from <page.json> --writer "<who>" [--model M] [--plan]
    python3 onepager.py build   --corpus <dir> --slug S --out <dir> [--draft] [--no-pdf] [--no-docx]
                                [--playwright-dir <dir>]
    python3 onepager.py check   --corpus <dir> --slug S [--draft]

It reads the same two profiles as a dossier (references/profiles/), plus the optional fields in
references/profiles/one-pager.schema.json: what the readers decide, their objections, the numbers they
think in, the page's ask and how far it may travel (`distribution`); the project's brand, images,
reviewers and the numbers that are its own facts.

The objects are UPC's own:

  * a FIGURE is an anchored passage (`upc anchor --type data_point`) that contains its value. A value
    read off a chart is anchored in a readings table stored as its own source (spec/05's OCR
    pattern); two readers are two readings. A derived figure names its inputs and its rule. The
    register (.upc/upc-corpus/one-pager/figures.json) holds each figure's label, value and the forms
    a page may print it in.
  * what a figure IS and who may see it are CODINGS in closed codebooks, one per question: how much
    it matters to each audience (onepager-rel-<audience>: lead, support, background, avoid), its class
    (measured, modelled, literature, derived), a reviewer's check (confirmed, corrected, query) and its
    release (cleared, internal-only, withheld). A model, Ellie and the engineers code as different
    coders; a person's answer outranks a model's, and every answer stays visible.
  * a PAGE is an `x-one-pager` synthesis: its words, with a bare citation for every figure it uses
    and the figures' passages quoted under hop C. A revision supersedes the one before.

`build` refuses a final page that uses a figure not judged lead or support for its audience, not
confirmed by a reviewer, or (for a partner or public page) not cleared; that prints a number which is
not one of its figures; that holds a never-publish phrase or a local path; or whose PDF is not exactly
one page. `--draft` builds anyway, stamped "Internal draft, not for distribution", and lists what is
open. It writes the page (HTML, and a PDF through headless Chromium) and the figures note: every
number, its passage, its readings, who checked it and what is still open.

The page and figure shapes are in references/one-pager.md.
"""

from __future__ import annotations

import argparse
import base64
import csv
import datetime as dt
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import common as C  # noqa: E402
import dossier as D  # noqa: E402
from corpus_view import (check_extraction, esc, exact, extractions, find_span, read_jsonl,  # noqa: E402
                         suppressed_in, validate, value_in)
from make_codebooks import build_codebook, write_codebooks  # noqa: E402

TOOL = {"tool": "upc-corpus/onepager", "tool_version": "1.0.0"}
PAGE = "x-one-pager"
HERE = Path(__file__).resolve().parent
ASSETS = HERE.parent / "assets" / "one-pager"
OVERLAY = D.SCHEMAS / "one-pager.schema.json"
RENDER = HERE / "onepager_render.mjs"
LAYOUTS = ("executive", "technical", "policy")
DISTRIBUTIONS = ("internal", "partner", "public")
MODEL_CODER = "onepager-model-v1"
SLUG = re.compile(r"^[a-z0-9][a-z0-9-]*$")
DRAFT_STAMP = "Internal draft, not for distribution"

REL = (("lead", "Lead", "Headline-worthy for this reader: a number they would act on, repeat to a colleague or put on "
                        "a slide, in units they already use."),
       ("support", "Support", "Backs up a lead point for this reader: worth a place on their page, but not the first "
                              "thing they need to see."),
       ("background", "Background", "True and relevant to the project, but this reader would not miss it; it belongs "
                                    "in the notes, not on their page."),
       ("avoid", "Avoid", "Accurate, but wrong for this reader: in units they distrust, at a scale they do not work "
                          "at, or likely to raise an objection the page cannot answer."))
CLASS = (("measured", "Measured", "Observed in a test, trial or plant by the project or the source: a lab or field "
                                  "result, not a model or an estimate."),
         ("modelled", "Modelled", "The output of a model, a simulation or an engineering estimate, however careful; not "
                                  "yet observed."),
         ("literature", "Literature", "Reported by someone else's published work: a paper, a government statistic, a "
                                      "standard or a regulation."),
         ("derived", "Derived", "Worked out from other figures by a stated rule (a difference, a ratio, a rounding), "
                                "not stated as such in any one source."))
CHECK = (("confirmed", "Confirmed", "A named reviewer checked the figure against the work it comes from, and it stands as "
                                    "shown, with its wording."),
         ("corrected", "Corrected", "A reviewer found the figure wrong or overstated; it must change before it is used "
                                    "again."),
         ("query", "Query", "A reviewer has an open question about the figure; it stays off final pages until the "
                            "question is answered."))
CLEAR = (("cleared", "Cleared", "The person who may release this result agreed it can leave the team, for the "
                                "audiences the rationale names."),
         ("internal-only", "Internal only", "The result may be used inside the team and with the funder, but not on a "
                                            "page that goes to partners or the public."),
         ("withheld", "Withheld", "The result must not be used on any page, draft or final, until this is lifted."))
BOOKS = {"class": ("onepager-figure-class", "One-pager: what kind of figure it is", CLASS),
         "check": ("onepager-figure-check", "One-pager: a reviewer's check of a figure", CHECK),
         "clearance": ("onepager-figure-clearance", "One-pager: whether a figure may leave the team", CLEAR)}
CODES = {"rel": [c for c, _, _ in REL], "class": [c for c, _, _ in CLASS], "check": [c for c, _, _ in CHECK],
         "clearance": [c for c, _, _ in CLEAR]}
KIND_RANK = {"adjudicator": 3, "human": 2, "rule": 1, "ensemble": 1, "model": 0}

NUM = re.compile(r"(?<![\w.,])\d{1,3}(?:,\d{3})+(?:\.\d+)?(?![\w,])|(?<![\w.,])\d+(?:\.\d+)?|[½⅓⅔¼¾⅕⅖⅗⅘⅙⅚⅛⅜⅝⅞]")
YEAR = re.compile(r"^(19|20)\d\d$")
FRAC = {"½": 1 / 2, "⅓": 1 / 3, "⅔": 2 / 3, "¼": 1 / 4, "¾": 3 / 4, "⅕": 1 / 5, "⅖": 2 / 5, "⅗": 3 / 5, "⅘": 4 / 5,
        "⅙": 1 / 6, "⅚": 5 / 6, "⅛": 1 / 8, "⅜": 3 / 8, "⅝": 5 / 8, "⅞": 7 / 8}
P = D._prose  # the writer's words in a UPC output: a straight double quotation mark would read as a quotation


def rel_slug(aud_id: str) -> str:
    return f"onepager-rel-{aud_id}"


# ------------------------------------------------------------------ profiles, config, the register
def load_profiles(cfg: dict) -> tuple[dict, dict]:
    """The dossier's two profiles, then the one-pager's own fields checked against the overlay schema."""
    aud, proj = D.load_profile(cfg["audience"], "audience"), D.load_profile(cfg["project"], "project")
    overlay = C.load_json(OVERLAY, {})
    errs = [f"audience {e}" for e in D.schema_errors(overlay["$defs"]["audience"], aud, root=overlay)]
    errs += [f"project {e}" for e in D.schema_errors(overlay["$defs"]["project"], proj, root=overlay)]
    ids = [a.get("id") for a in (proj.get("one_pager") or {}).get("assets") or []]
    errs += [f"project one_pager.assets: {i!r} is used twice" for i in sorted({i for i in ids if ids.count(i) > 1})]
    if errs:
        raise SystemExit("the profiles do not fit references/profiles/one-pager.schema.json:\n  " + "\n  ".join(errs[:12]))
    return aud, proj


def opt(profile: dict) -> dict:
    return profile.get("one_pager") or {}


def sdir(corpus: Path, *parts: str) -> Path:
    return C.state_dir(corpus, "one-pager", *parts)


def page_path(corpus: Path, slug: str) -> Path:
    return sdir(corpus, "pages") / f"{slug}.json"


def load_cfg(corpus: Path, slug: str) -> dict:
    cfg = C.load_json(page_path(corpus, slug), None)
    if not cfg:
        raise SystemExit(f"no one-pager {slug!r} here yet: run `onepager.py init --corpus {corpus} --slug {slug} …` first")
    return cfg


def reg_path(corpus: Path) -> Path:
    return sdir(corpus) / "figures.json"


def load_reg(corpus: Path) -> dict[str, dict]:
    return (C.load_json(reg_path(corpus), {}) or {}).get("figures") or {}


def save_reg(corpus: Path, figs: dict) -> None:
    C.save_json(reg_path(corpus), {"figures": dict(sorted(figs.items()))})


def namespace(corpus: Path) -> str:
    """The corpus's namespace (its settings), else the one this format first used here, kept in its
    state so the sets keep their ids when the corpus folder is moved or copied."""
    ns = C.corpus_settings(corpus).get("namespace")
    if ns:
        return ns
    st = sdir(corpus) / "settings.json"
    ns = (C.load_json(st, {}) or {}).get("namespace")
    if not ns:
        ns = re.sub(r"[^a-z0-9]+", "-", corpus.name.lower()).strip("-") or "onepager"
        C.save_json(st, {"namespace": ns})
    return ns


def exs_id(corpus: Path) -> str:
    return f"exs-{namespace(corpus)}-onepager"


def cds_id(corpus: Path) -> str:
    return f"cds-{namespace(corpus)}-onepager"


def digest(obj) -> str:
    return hashlib.sha256(json.dumps(obj, sort_keys=True, default=str).encode()).hexdigest()[:16]


def now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


# ------------------------------------------------------------------ anchoring
def anchor(corpus: Path, info: dict, quotes: list[str | None], *, typ: str, model: str, prompt: str,
           dry: bool = False) -> list[dict]:
    """dossier.anchor_quotes with this format's own extraction set and type: each quotation, repaired
    to the document's own characters (case, curly marks, dashes, spacing; widened when it occurs
    twice), found by `upc anchor`. -> {quote, extraction_id} or {quote, status}. `--source` keeps a
    passage on its own source when two share one reading copy."""
    text = Path(info["text_path"]).read_text(encoding="utf-8")
    fixed = []
    for q in quotes:
        q = re.sub(r"[ \t]+", " ", str(q or "")).strip()
        if q and q not in text:
            q = find_span(q, text) or q
        if q and C.occurrences(text, q) > 1:
            q = C.widen_to_unique(text, q) or q
        fixed.append(q)
    out = [{"quote": q, "status": "no quotation given"} for q in fixed]
    todo = [i for i, q in enumerate(fixed) if q]
    if not todo:
        return out
    args = ["anchor", "--corpus", str(corpus), "--rep", info["text_rep"], "--source", info["source_id"], "--normalize",
            "--set", exs_id(corpus), "--type", typ, "--tool", TOOL["tool"], "--model", model or "unknown",
            "--prompt-version", prompt]
    res = C.upc_jsonl(*args, *(["--dry-run"] if dry else []), records=[{"quote": fixed[i]} for i in todo])
    if not isinstance(res, dict) or "results" not in res:
        raise C.Unavailable(f"upc anchor gave no results: {str(res)[:200]}")
    by_i = {r["i"]: r for r in res["results"]}
    exts = {} if dry else extractions(corpus)
    for j, i in enumerate(todo):
        r = by_i.get(j) or {}
        if r.get("status") in ("anchored", "refound", "exists"):
            eid = r.get("extraction_id")
            stored = (exts.get(eid) or {}).get("direct_quote") or r.get("stored_quote") or fixed[i]
            out[i] = {"quote": stored, "extraction_id": eid}
        else:
            out[i] = {"quote": fixed[i], "status": r.get("status") or "no result"}
    return out


# ------------------------------------------------------------------ numbers
def num_tokens(s) -> list[str]:
    """Every number in a text, as written (thousands commas dropped). A number inside a word (CO2,
    DE-EE0011216) is not one."""
    return [t.replace(",", "") for t in NUM.findall(str(s or ""))]


def number_of(s) -> float | None:
    toks = num_tokens(s)
    if not toks:
        return None
    return FRAC.get(toks[0]) or float(toks[0])


def shown_forms(f: dict) -> list[str]:
    return list(f.get("display") or [f["value"]])


# ------------------------------------------------------------------ figures
FIG_KEYS = ("label", "value", "unit", "display", "as_of", "tolerance", "derived", "needs_clearance", "note")


def _inputs(f: dict) -> dict:
    """What a person gave for a figure, to tell a repeat from a change."""
    keep = {k: f.get(k) for k in FIG_KEYS + ("quote", "readings") if f.get(k) not in (None, "", [], {})}
    if f.get("source_id") or f.get("source"):
        keep["source"] = f.get("source_id") or f.get("source")
    keep["readings"] = [{"quote": r.get("quote"), "value": r.get("value"), **({"rep": r["rep"]} if r.get("rep") else {})}
                        for r in keep.get("readings") or []]
    if f.get("rep"):
        keep["rep"] = f["rep"]
    return keep


def text_copy(corpus: Path, info: dict, rep: str | None) -> dict | str:
    """The source as anchor() needs it, reading `rep` (one of its text representations, e.g. the
    readings of one slide, whose parent is the slide's picture) instead of its reading copy. -> the
    info, or what is wrong."""
    if not rep:
        return info
    r = next((x for x in (info.get("record") or {}).get("representations") or [] if x.get("representation_id") == rep), None)
    if not r:
        return f"source {info['source_id']} has no representation {rep}"
    if not str(r.get("media_type", "")).startswith("text/") or not r.get("path"):
        return f"representation {rep} is not text ({r.get('media_type')})"
    return {**info, "text_rep": rep, "text_path": str(corpus / r["path"])}


def cmd_figures_add(corpus: Path, items: list, *, by: str, replace: bool, plan: bool) -> int:
    reg, idx = load_reg(corpus), D.source_index(corpus)
    ids = [f.get("id") if isinstance(f, dict) else None for f in items]
    problems, todo = [], []
    for i, f in enumerate(items):
        where = f"figures[{i}]" + (f" {f.get('id')!r}" if isinstance(f, dict) else "")
        if not isinstance(f, dict) or not SLUG.match(str(f.get("id") or "")):
            problems.append(f"{where}: an id of lower-case letters, digits and dashes is needed")
            continue
        if ids.count(f["id"]) > 1:
            problems.append(f"{where}: the id is used twice in this file")
            continue
        if not str(f.get("label") or "").strip() or not str(f.get("value") or "").strip():
            problems.append(f"{where}: a label and a value are needed")
            continue
        der = f.get("derived")
        if der is not None:
            if not isinstance(der, dict) or not der.get("from") or not der.get("rule"):
                problems.append(f"{where}: derived needs from (the ids of the figures it is worked out from) and rule")
                continue
            unknown = [x for x in der["from"] if x not in reg and x not in ids]
            if unknown:
                problems.append(f"{where}: derived.from names figures that are not in the register: {', '.join(unknown)}")
                continue
        if der is None or f.get("quote"):
            if not idx.get(f.get("source")):
                problems.append(f"{where}: no source {f.get('source')!r} with a reading copy in this corpus")
                continue
            if not value_in(f["value"], f.get("quote")):
                problems.append(f"{where}: its value {f['value']!r} is not in its quote; a figure's passage must state it")
                continue
            tc = text_copy(corpus, idx[f["source"]], f.get("rep"))
            if isinstance(tc, str):
                problems.append(f"{where}: {tc}")
                continue
        for j, r in enumerate(f.get("readings") or []):
            if isinstance(r, dict) and idx.get(r.get("source")) and isinstance(text_copy(corpus, idx[r["source"]], r.get("rep")), str):
                problems.append(f"{where}.readings[{j}]: {text_copy(corpus, idx[r['source']], r.get('rep'))}")
            if not isinstance(r, dict) or not idx.get(r.get("source")) or not value_in(r.get("value"), r.get("quote")):
                problems.append(f"{where}.readings[{j}]: a reading needs a source in this corpus and a quote that holds its value")
        for d in f.get("display") or []:
            if not str(d).strip():
                problems.append(f"{where}: an empty display form")
        old = reg.get(f["id"])
        if old and not replace and old.get("input_digest") != digest(_inputs(f)):
            problems.append(f"{where}: already in the register with different content; pass --replace to change it")
            continue
        todo.append(f)
    if problems:
        raise SystemExit("the figures file does not fit (references/one-pager.md, 'A figure'); nothing was stored:\n  "
                         + "\n  ".join(problems[:20]))

    asks: dict[tuple, list] = {}
    for k, f in enumerate(todo):
        if f.get("quote"):
            asks.setdefault((idx[f["source"]]["source_id"], f.get("rep")), []).append((k, None, f["quote"]))
        for j, r in enumerate(f.get("readings") or []):
            asks.setdefault((idx[r["source"]]["source_id"], r.get("rep")), []).append((k, j, r["quote"]))
    found: dict[tuple, dict] = {}
    for (sid, rep), lst in asks.items():
        info = text_copy(corpus, idx[sid], rep)
        if isinstance(info, str):
            raise SystemExit(f"figures add: {info}")
        got = anchor(corpus, info, [q for *_, q in lst], typ="data_point", model=by, prompt="onepager-figure@1", dry=plan)
        for (k, j, _), g in zip(lst, got):
            found[(k, j)] = g

    batch = {f["id"]: f for f in todo}
    by_ext = {v["extraction_id"]: fid for fid, v in reg.items() if v.get("extraction_id")}
    kept, dropped = {}, []
    for k, f in enumerate(todo):
        rec = {"id": f["id"], **{key: f[key] for key in FIG_KEYS if f.get(key) not in (None, "", [], {})}}
        rec["display"] = [str(d) for d in f.get("display") or [f["value"]]]
        rec["input_digest"] = digest(_inputs(f))
        ok = True
        if f.get("quote"):
            g = found[(k, None)]
            eid = g.get("extraction_id")
            other = by_ext.get(eid) if eid else None
            if not eid and not (plan and "status" not in g):
                dropped.append((f["id"], "its quote", g.get("status")))
                ok = False
            elif other and other != f["id"]:
                dropped.append((f["id"], "its quote", f"the same passage as figure {other}; give each figure its own span"))
                ok = False
            else:
                rec.update(source_id=idx[f["source"]]["source_id"], quote=g["quote"], extraction_id=eid,
                           **({"rep": f["rep"]} if f.get("rep") else {}))
        for j, r in enumerate(f.get("readings") or []):
            g = found[(k, j)]
            if not g.get("extraction_id") and not (plan and "status" not in g):
                dropped.append((f["id"], f"reading {j + 1}", g.get("status")))
                ok = False
                continue
            rd = {"source_id": idx[r["source"]]["source_id"], "quote": g["quote"], "extraction_id": g.get("extraction_id"),
                  "value": str(r["value"]), **({"rep": r["rep"]} if r.get("rep") else {})}
            a, b = number_of(r["value"]), number_of(f["value"])
            if f.get("tolerance") is not None and a is not None and b is not None:
                rd["agrees"] = abs(a - b) <= float(f["tolerance"]) + 1e-9
            rec.setdefault("readings", []).append(rd)
        der = rec.get("derived")
        if der:
            tol = max([float((reg.get(x) or batch.get(x) or {}).get("tolerance") or 0) for x in der["from"]] or [0])
            if der.get("margin") is not None and float(der["margin"]) <= tol:
                rec["thin"] = (f"its margin ({der['margin']}) is within the readings' tolerance ({tol}): "
                               "a second reading could reverse it")
        vt = set(num_tokens(rec["value"]))
        rounded = [d for d in rec["display"] if set(num_tokens(d)) - vt]
        if rounded:
            rec["rounded"] = rounded
        if reg.get(f["id"], {}).get("judgements") and not rec.get("extraction_id"):
            rec["judgements"] = reg[f["id"]]["judgements"]
        if ok:
            kept[f["id"]] = rec
            if rec.get("extraction_id"):
                by_ext[rec["extraction_id"]] = f["id"]

    C.counts_header("Figures: add" + (" (plan: nothing stored)" if plan else ""),
                    [("in the file", len(items)), ("kept", len(kept)), ("not stored", len({d[0] for d in dropped})),
                     ("in the register after this", len({**reg, **kept}) if not plan else len(reg))])
    for fid, rec in kept.items():
        extra = []
        if rec.get("readings"):
            n_ok = sum(1 for r in rec["readings"] if r.get("agrees") is not False)
            extra.append(f"{len(rec['readings'])} reading(s), {n_ok} agreeing")
        if rec.get("thin"):
            extra.append("THIN margin")
        if rec.get("rounded"):
            extra.append(f"shown as {', '.join(rec['rounded'])}")
        moved = reg.get(fid, {}).get("extraction_id") not in (None, rec.get("extraction_id"))
        print(f"  kept {fid}: {rec['value']}" + (f" ({'; '.join(extra)})" if extra else "")
              + ("  NOTE: a new passage, so its codings start again" if moved else ""))
    for fid, what, status in dropped:
        print(f"  NOT FOUND {fid}, {what}: {status}")
    if not plan and kept:
        reg.update(kept)
        save_reg(corpus, reg)
        C.upc("regen", str(corpus), check=False)
    return 1 if dropped else 0


def cmd_figures_import(corpus: Path, plan: bool) -> int:
    """The figures of a dossier in this corpus, whose passages are anchored already."""
    reg = load_reg(corpus)
    by_ext = {v["extraction_id"]: fid for fid, v in reg.items() if v.get("extraction_id")}
    new, skipped = {}, []
    for e in D.entries(corpus):
        for it in e.get("figures") or []:
            eid = it.get("extraction_id")
            if not eid or eid in by_ext:
                continue
            value = it.get("value") if value_in(it.get("value"), it.get("quote")) else exact(it.get("value"), it.get("quote"))
            if not value:
                skipped.append(f"{it.get('label')!r} ({e['title'][:50]}): its value {it.get('value')!r} is not in its "
                               "quote; add it by hand with figures add")
                continue
            base = D.slug(it.get("label") or "figure", 40)
            fid, n = base, 2
            while fid in reg or fid in new:
                fid, n = f"{base}-{n}", n + 1
            new[fid] = {"id": fid, "label": it.get("label") or fid, "value": value, "display": [value],
                        "source_id": e["_source_id"], "quote": it["quote"], "extraction_id": eid, "from_dossier": e["_gen"]}
            by_ext[eid] = fid
    C.counts_header("Figures: import from the dossier" + (" (plan)" if plan else ""),
                    [("new figures", len(new)), ("skipped", len(skipped)), ("already in the register", len(reg))])
    for fid, f in new.items():
        print(f"  {fid}: {f['value']}  ({f['label'][:60]})")
    for s in skipped:
        print(f"  SKIPPED {s}")
    if not plan and new:
        reg.update(new)
        save_reg(corpus, reg)
    return 0


# ------------------------------------------------------------------ codebooks and codings
def _book(slug: str, title: str, codes, question: str) -> dict:
    return {"slug": slug, "title": title, "closed": True, "multi_label": False, "unit": "passage", "tool": TOOL["tool"],
            "question": question, "codes": [{"code": c, "label": lab, "definition": d} for c, lab, d in codes]}


def book_def(slug: str) -> dict:
    for key, (s, title, codes) in BOOKS.items():
        if s == slug:
            q = {"class": "What kind of figure is this?", "check": "Has a reviewer checked this figure?",
                 "clearance": "May this figure leave the team?"}[key]
            return _book(s, title, codes, q)
    if slug.startswith("onepager-rel-"):
        a = slug.removeprefix("onepager-rel-")
        return _book(slug, f"One-pager: how much a figure matters to {a}", REL,
                     f"How much does this figure matter on a one-pager for the audience {a}?")
    raise ValueError(f"not a one-pager codebook: {slug}")


def book_ids(corpus: Path) -> dict[str, str]:
    """This format's codebooks in the corpus: slug -> codebook id."""
    out = {}
    for f in (corpus / "codebooks").glob("*.json"):
        d = C.load_json(f, {}) or {}
        if str(d.get("slug", "")).startswith("onepager-"):
            out[d["slug"]] = d["codebook_id"]
    return out


def ensure_books(corpus: Path, slugs: list[str]) -> dict[str, str]:
    have = book_ids(corpus)
    missing = [s for s in dict.fromkeys(slugs) if s not in have]
    if missing:
        ns = namespace(corpus)
        rows = [build_codebook(corpus, book_def(s), ns)[:3] for s in missing]
        write_codebooks(corpus, rows)
        have.update({o["slug"]: o["codebook_id"] for o, _, _ in rows})
    return have


def codings_paths(corpus: Path) -> tuple[Path, Path]:
    sec = ((C.load_json(corpus / "corpus.json", {}) or {}).get("sections") or {}).get("codings") or "codings/"
    d = corpus / sec / cds_id(corpus)
    return d / "items.jsonl", d / "manifest.json"


def judgements(corpus: Path, reg: dict) -> dict[tuple[str, str], list[dict]]:
    """(figure id, codebook slug) -> every active answer {code, coder, kind, rationale, at, coding_id}.
    A derived figure has no passage of its own for UPC to code, so its answers live in the register."""
    items, man = codings_paths(corpus)
    kinds = {c.get("coder"): c.get("kind", "model") for c in (C.load_json(man, {}) or {}).get("coders", [])}
    slug_of = {v: k for k, v in book_ids(corpus).items()}
    fig_of = {f["extraction_id"]: fid for fid, f in reg.items() if f.get("extraction_id")}
    out: dict[tuple[str, str], list[dict]] = {}
    for c in read_jsonl(items) if items.exists() else []:
        if (c.get("status") or "active") != "active":
            continue
        fid, book = fig_of.get((c.get("target") or {}).get("id")), slug_of.get(c.get("codebook_ref"))
        if fid and book:
            out.setdefault((fid, book), []).append({
                "code": c.get("code"), "coder": c["coder"], "kind": kinds.get(c["coder"], "model"),
                "rationale": c.get("rationale", ""), "at": (c.get("provenance") or {}).get("created_at", ""),
                "coding_id": c.get("coding_id")})
    for fid, f in reg.items():
        for j in f.get("judgements") or []:
            if j.get("status", "active") == "active":
                out.setdefault((fid, j["book"]), []).append(j)
    return out


def resolve(answers: list[dict]) -> tuple[str | None, bool]:
    """The answer that counts, and whether coders disagree. A person (or an adjudicator) outranks a
    model; among equals the newest stands. The others stay visible in the review sheet and the notes."""
    if not answers:
        return None, False
    top = max(answers, key=lambda a: (KIND_RANK.get(a.get("kind"), 0), a.get("at") or ""))
    return top["code"], len({a["code"] for a in answers}) > 1


def write_codes(corpus: Path, reg: dict, records: list[dict], *, coder: str, kind: str, model: str | None = None,
                prompt: str | None = None) -> dict:
    """records: [{figure, book, code, rationale}] -> the tally. A figure with a passage gets a UPC
    coding (`upc code`, which supersedes the same coder's earlier answer); a derived one without a
    passage gets the same in the register."""
    have = ensure_books(corpus, [r["book"] for r in records])
    upc_recs, tally, changed = [], {}, False
    for r in records:
        f = reg[r["figure"]]
        if f.get("extraction_id"):
            upc_recs.append({"codebook_ref": have[r["book"]], "code": r["code"],
                             "target": {"kind": "extraction", "id": f["extraction_id"]}, "coder": coder,
                             **({"rationale": r["rationale"]} if r.get("rationale") else {})})
            continue
        js = f.setdefault("judgements", [])
        same = [j for j in js if j["book"] == r["book"] and j["coder"] == coder and j.get("status", "active") == "active"]
        if same and same[0]["code"] == r["code"]:
            tally["unchanged"] = tally.get("unchanged", 0) + 1
            continue
        for j in same:
            j["status"] = "superseded"
        js.append({"book": r["book"], "code": r["code"], "coder": coder, "kind": kind, "rationale": r.get("rationale", ""),
                   "at": now()})
        tally["coded"] = tally.get("coded", 0) + 1
        changed = True
    if upc_recs:
        res = C.upc_jsonl("code", "--corpus", str(corpus), "--set", cds_id(corpus), "--coder-kind", kind,
                          "--tool", TOOL["tool"], "--method", "model" if kind == "model" else "manual",
                          *(["--model", model] if model else []), *(["--prompt-version", prompt] if prompt else []),
                          records=upc_recs)
        if not isinstance(res, dict):
            raise C.Unavailable(f"upc code gave no result: {str(res)[:200]}")
        bad = [x for x in res.get("results", []) if x.get("status") in ("error", "invalid")]
        if bad:
            raise C.Unavailable(f"upc code refused {len(bad)} coding(s): {bad[0].get('detail')}")
        for k, v in (res.get("tally") or {}).items():
            tally[k] = tally.get(k, 0) + v
    if changed:
        save_reg(corpus, reg)
    C.upc("regen", str(corpus), check=False)
    return tally


def status_of(fid: str, reg: dict, J: dict, aud_id: str, _seen: frozenset = frozenset()) -> dict:
    f = reg[fid]
    rel, rel_d = resolve(J.get((fid, rel_slug(aud_id)), []))
    cls, cls_d = resolve(J.get((fid, BOOKS["class"][0]), []))
    chk, chk_d = resolve(J.get((fid, BOOKS["check"][0]), []))
    clr, clr_d = resolve(J.get((fid, BOOKS["clearance"][0]), []))
    if not cls and f.get("derived"):
        cls = "derived"
    inputs = {x: status_of(x, reg, J, aud_id, _seen | {fid}) for x in (f.get("derived") or {}).get("from") or []
              if x in reg and x not in _seen}
    return {"rel": rel, "class": cls, "check": chk, "clearance": clr, "inputs": inputs,
            "disagree": [k for k, d in (("relevance", rel_d), ("class", cls_d), ("check", chk_d), ("clearance", clr_d)) if d]}


def gate(fid: str, reg: dict, J: dict, aud_id: str, distribution: str) -> tuple[list[str], list[str]]:
    """-> (always blocking, blocking on a final page). A draft builds with the second list as warnings."""
    s, f = status_of(fid, reg, J, aud_id), reg[fid]
    hard, final = [], []
    if s["clearance"] == "withheld" or any(i["clearance"] == "withheld" for i in s["inputs"].values()):
        hard.append("withheld from every page")
    if s["rel"] not in ("lead", "support"):
        final.append(f"relevance for this audience is {s['rel'] or 'not judged yet'}")
    if not s["class"]:
        final.append("no class yet (measured, modelled, literature or derived)")
    if s["check"] != "confirmed":
        final.append(f"not confirmed by a reviewer ({s['check'] or 'not checked yet'})")
    unchecked = [x for x, i in s["inputs"].items() if i["check"] != "confirmed"]
    if unchecked:
        final.append("not every figure it is worked out from is confirmed")
    if distribution != "internal":
        if s["clearance"] != "cleared":
            final.append(f"not cleared to leave the team ({s['clearance'] or 'no release recorded'})")
        if any(i["clearance"] != "cleared" for i in s["inputs"].values()):
            final.append("not every figure it is worked out from is cleared to leave the team")
    if any(r.get("agrees") is False for r in f.get("readings") or []):
        final.append("its readings disagree by more than the tolerance")
    return hard, final


def closure(ids: list[str], reg: dict) -> list[str]:
    """The figures and, for a derived one, the figures it is worked out from."""
    out, todo = [], list(ids)
    while todo:
        x = todo.pop(0)
        if x in out or x not in reg:
            continue
        out.append(x)
        todo += (reg[x].get("derived") or {}).get("from") or []
    return out


# ------------------------------------------------------------------ init
def cmd_init(corpus: Path, a) -> int:
    if not SLUG.match(a.slug or ""):
        raise SystemExit("--slug: lower-case letters, digits and dashes")
    cfg = C.load_json(page_path(corpus, a.slug), {}) or {}
    for key, val in (("title", a.title), ("date", a.date), ("layout", a.layout), ("distribution", a.distribution)):
        if val is not None:
            cfg[key] = val
    if a.audience:
        cfg["audience"] = str(Path(a.audience).expanduser().resolve())
    if a.project:
        cfg["project"] = str(Path(a.project).expanduser().resolve())
    for pair in getattr(a, "show_writer_as", None) or []:
        real, sep, label = pair.partition("=")
        if not sep or not real.strip():
            raise SystemExit("--show-writer-as: WRITER=LABEL, e.g. 'Claude (Opus 5.5)=Hermes'")
        labels = cfg.setdefault("writer_labels", {})
        if label.strip():
            labels[real.strip()] = label.strip()
        else:
            labels.pop(real.strip(), None)
    for need in ("audience", "project"):
        if not cfg.get(need):
            raise SystemExit(f"init needs --{need} the first time")
    aud, proj = load_profiles(cfg)
    cfg["slug"] = a.slug
    cfg.setdefault("layout", opt(aud).get("layout") or "technical")
    cfg.setdefault("distribution", opt(aud).get("distribution") or "internal")
    if cfg["layout"] not in LAYOUTS or cfg["distribution"] not in DISTRIBUTIONS:
        raise SystemExit(f"--layout is one of {', '.join(LAYOUTS)}; --distribution one of {', '.join(DISTRIBUTIONS)}")
    cfg.setdefault("title", f"{proj['name']}: for {aud['name']}")
    cfg.setdefault("date", dt.date.today().isoformat())
    C.save_json(page_path(corpus, a.slug), cfg)
    books = ensure_books(corpus, [b[0] for b in BOOKS.values()] + [rel_slug(aud["id"])])
    C.counts_header("One-pager: config", [("config", page_path(corpus, a.slug)), ("title", cfg["title"]),
                                          ("audience", f"{aud['name']} ({aud['id']})"), ("project", proj["name"]),
                                          ("layout", cfg["layout"]), ("distribution", cfg["distribution"]),
                                          ("codebooks", len(books))])
    return 0


# ------------------------------------------------------------------ the profile, for a prompt
def reader_text(aud: dict) -> str:
    def line(label, xs):
        xs = [x for x in xs if x]
        return f"{label}: {'; '.join(xs)}." if xs else ""
    o = opt(aud)
    nums = aud.get("numbers") or {}
    reading = aud.get("reading")
    parts = [
        f"The readers: {aud['name']}" + (f" ({aud['organization']})" if aud.get("organization") else "") + ": "
        + "; ".join(f"{r['name']} ({r['role']})" for r in aud["readers"]) + ".",
        aud.get("description") or "",
        line("They already know (do not explain)", aud.get("knows") or []),
        line("They look for", aud.get("wants") or []),
        line("They decide", aud.get("decides") or []),
        line("They care about", aud.get("cares_about") or []),
        line("They will push back on", [o_["objection"] + (f" (answer: {o_['answer_hint']})" if o_.get("answer_hint") else "")
                                        for o_ in aud.get("objections") or []]),
        line("What persuades them", aud.get("trusts") or []),
        line("Numbers they think in", nums.get("prefer") or []),
        line("Keep off their page", nums.get("avoid") or []),
        line("Their own words (use where they fit, never in quotation marks)", [w["phrase"] for w in aud.get("words") or []]),
        line("How they read it", reading if isinstance(reading, list) else [reading] if reading else []),
        f"How they like to read: {D.style_text(aud)}.",
        line("Words to avoid", o.get("banned_words") or []),
    ]
    return "\n".join(p for p in parts if p)


def project_text(proj: dict) -> str:
    return "\n".join(x for x in (
        f"The project: {proj['name']}" + (f" ({proj['full_name']})" if proj.get("full_name") else "") + ".",
        ("Aims: " + "; ".join(proj["aims"]) + ".") if proj.get("aims") else "",
        ("Project rules: " + "; ".join(proj["rules"]) + ".") if proj.get("rules") else "",
    ) if x)


# ------------------------------------------------------------------ code (the local model judges relevance)
def code_prompt(aud: dict, proj: dict) -> str:
    rel = "\n".join(f"  - {c}: {d}" for c, _, d in REL)
    cls = "\n".join(f"  - {c}: {d}" for c, _, d in CLASS)
    return f"""You judge which figures matter to one group of readers, for a one-page brief about a project.

{reader_text(aud)}

{project_text(proj)}

For every figure listed, give:
- code, how much it matters on a page for these readers:
{rel}
- kind, what sort of figure it is:
{cls}
- rationale: one sentence on why, for these readers.
- angle: how these readers would see it, in their own terms (a short phrase).
Judge only from the figure, its label and its passage. Output JSON only."""


def cmd_code(corpus: Path, slug: str, only: list[str] | None, redo: bool, plan: bool, coder: str = MODEL_CODER) -> int:
    cfg = load_cfg(corpus, slug)
    aud, proj = load_profiles(cfg)
    reg = load_reg(corpus)
    J = judgements(corpus, reg)
    book = rel_slug(aud["id"])
    unknown = [x for x in only or [] if x not in reg]
    if unknown:
        raise SystemExit(f"--figure: not in the register: {', '.join(unknown)}")
    # `code` fills gaps. A figure someone has already judged for this audience is left alone, whoever judged
    # it: among model answers the newest counts, so a second pass would silently replace the first. --redo
    # judges them again.
    asked = only or sorted(reg)
    todo = [fid for fid in asked if redo or not J.get((fid, book))]
    C.counts_header("Code: plan", [("audience", f"{aud['name']} ({aud['id']})"), ("figures", len(reg)),
                                   ("to judge now", len(todo)), ("judged already (--redo to judge again)", len(asked) - len(todo)),
                                   ("coder", coder),
                                   ("model", f"{C.DEFAULT_MODEL} at {C.CHAT_URL}")])
    if plan or not todo:
        return 0
    system = code_prompt(aud, proj)
    recs, failed = [], 0
    for at in range(0, len(todo), 24):
        chunk = todo[at:at + 24]
        lines = []
        for fid in chunk:
            f = reg[fid]
            said = f.get("quote") or "; ".join(f"{x} = {reg[x]['value']}" for x in (f.get("derived") or {}).get("from", []))
            lines.append(f"- {fid}: {f['label']} = {f['value']}{' ' + f['unit'] if f.get('unit') else ''}. "
                         f"Passage: {str(said)[:400]}")
        item = {"type": "object", "required": ["figure", "code", "kind", "rationale", "angle"],
                "properties": {"figure": {"type": "string", "enum": chunk}, "code": {"type": "string", "enum": CODES["rel"]},
                               "kind": {"type": "string", "enum": CODES["class"]}, "rationale": {"type": "string"},
                               "angle": {"type": "string"}}}
        schema = {"type": "object", "required": ["judgements"], "properties": {"judgements": {"type": "array", "items": item}}}
        try:
            out = C.llm_chat(system, "Figures:\n" + "\n".join(lines), schema=schema, max_tokens=4000, temperature=0.1)
        except C.Unavailable as e:
            print(f"  UNAVAILABLE {len(chunk)} figure(s): {e}")
            failed += len(chunk)
            continue
        got = {j.get("figure"): j for j in (C.parse_json_loose(out["content"]) or {}).get("judgements") or []
               if isinstance(j, dict)}
        for fid in chunk:
            j = got.get(fid)
            if not j or j.get("code") not in CODES["rel"]:
                print(f"  NO ANSWER {fid}")
                failed += 1
                continue
            why = " ".join(x for x in (str(j.get("rationale") or "").strip(), f"Angle: {j['angle'].strip()}." if j.get("angle") else "") if x)
            recs.append({"figure": fid, "book": book, "code": j["code"], "rationale": why})
            if j.get("kind") in CODES["class"] and not J.get((fid, BOOKS["class"][0])):
                recs.append({"figure": fid, "book": BOOKS["class"][0], "code": j["kind"], "rationale": why})
    tally = write_codes(corpus, reg, recs, coder=coder, kind="model", model=C.DEFAULT_MODEL, prompt="onepager-code@1") if recs else {}
    C.counts_header("Code: result", [("codings written", tally.get("coded", 0)), ("unchanged", tally.get("unchanged", 0)),
                                     ("figures with no answer", failed)])
    return 1 if failed else 0


# ------------------------------------------------------------------ review (a person's answers) and signoff
REVIEW_COLS = ["figure", "label", "value", "shown as", "class", "relevance", "class answers", "relevance answers",
               "check", "clearance", "passage", "source"]


def _answers(xs: list[dict]) -> str:
    return " | ".join(f"{a['coder']}: {a['code']}" + (f" ({a['rationale'][:160]})" if a.get("rationale") else "")
                      for a in sorted(xs, key=lambda a: (-KIND_RANK.get(a.get("kind"), 0), a["coder"])))


def review_rows(corpus: Path, slug: str) -> tuple[list[list[str]], dict]:
    cfg = load_cfg(corpus, slug)
    aud, _ = load_profiles(cfg)
    reg = load_reg(corpus)
    J = judgements(corpus, reg)
    idx = D.source_index(corpus)
    rows = []
    for fid in sorted(reg):
        f, s = reg[fid], status_of(fid, reg, J, aud["id"])
        src = (idx.get(f.get("source_id") or "") or {}).get("title") or ""
        passage = f.get("quote") or ("Worked out: " + (f.get("derived") or {}).get("rule", ""))
        rows.append([fid, f["label"], f["value"], " / ".join(shown_forms(f)), s["class"] or "", s["rel"] or "",
                     _answers(J.get((fid, BOOKS["class"][0]), [])), _answers(J.get((fid, rel_slug(aud["id"])), [])),
                     s["check"] or "", s["clearance"] or "", passage, src])
    return rows, aud


def cmd_review_out(corpus: Path, slug: str, out: Path) -> int:
    rows, aud = review_rows(corpus, slug)
    out = out.expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.suffix == ".csv":
        with open(out, "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(REVIEW_COLS)
            w.writerows(rows)
    elif out.suffix == ".xlsx":
        from openpyxl import Workbook  # noqa: PLC0415
        from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
        from openpyxl.worksheet.datavalidation import DataValidation  # noqa: PLC0415
        wb = Workbook()
        ws = wb.active
        ws.title = "Figures"
        ws.append(REVIEW_COLS)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E6EAEE")
        for r in rows:
            ws.append(r)
        for col, codes in (("E", CODES["class"]), ("F", CODES["rel"])):
            dv = DataValidation(type="list", formula1='"' + ",".join(codes) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{max(2, len(rows) + 1)}")
            ws[f"{col}1"].fill = PatternFill("solid", fgColor="FFE699")
        for i, wdt in enumerate([24, 34, 12, 18, 13, 13, 40, 50, 12, 14, 60, 40]):
            ws.column_dimensions[chr(65 + i)].width = wdt
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        wb.save(out)
    elif out.suffix == ".md":
        L = [f"# Figures for {aud['name']}", "",
             "Edit the class and relevance columns in the .csv or .xlsx version, then `review --apply` it.", "",
             "| " + " | ".join(REVIEW_COLS[:10]) + " |", "|" + "---|" * 10]
        L += ["| " + " | ".join(str(c).replace("|", "/").replace("\n", " ") for c in r[:10]) + " |" for r in rows]
        out.write_text("\n".join(L) + "\n", encoding="utf-8")
    else:
        raise SystemExit("--out: a .csv, .xlsx or .md file")
    C.counts_header("Review: sheet", [("figures", len(rows)), ("audience", aud["id"]), ("sheet", out)],
                    ["Edit the class and relevance cells (or leave a cell empty for no answer), then apply it."])
    return 0


def cmd_review_apply(corpus: Path, slug: str, sheet: Path, coder: str) -> int:
    if not SLUG.match(coder or "") or coder == MODEL_CODER:
        raise SystemExit("--coder: the person's own handle (lower-case letters, digits, dashes), not the model's")
    sheet = sheet.expanduser()
    if sheet.suffix == ".csv":
        with open(sheet, encoding="utf-8", newline="") as fh:
            table = list(csv.reader(fh))
    elif sheet.suffix == ".xlsx":
        from openpyxl import load_workbook  # noqa: PLC0415
        ws = load_workbook(sheet, read_only=True).worksheets[0]
        table = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        raise SystemExit("--apply: a .csv or .xlsx sheet made by `review --out`")
    if not table or [h.strip().lower() for h in table[0][:len(REVIEW_COLS)]] != REVIEW_COLS:
        raise SystemExit(f"{sheet.name}: the columns are not review --out's ({', '.join(REVIEW_COLS)})")
    cfg = load_cfg(corpus, slug)
    aud, _ = load_profiles(cfg)
    reg = load_reg(corpus)
    recs, problems = [], []
    for n, r in enumerate(table[1:], 2):
        r = (r + [""] * len(REVIEW_COLS))[:len(REVIEW_COLS)]
        fid, cls, rel = r[0].strip(), r[4].strip().lower(), r[5].strip().lower()
        if not fid:
            continue
        if fid not in reg:
            problems.append(f"row {n}: no figure {fid!r} in the register")
            continue
        if cls:
            if cls not in CODES["class"]:
                problems.append(f"row {n} ({fid}): class {cls!r} is not one of {', '.join(CODES['class'])}")
            else:
                recs.append({"figure": fid, "book": BOOKS["class"][0], "code": cls, "rationale": f"Set by {coder} in review."})
        if rel:
            if rel not in CODES["rel"]:
                problems.append(f"row {n} ({fid}): relevance {rel!r} is not one of {', '.join(CODES['rel'])}")
            else:
                recs.append({"figure": fid, "book": rel_slug(aud["id"]), "code": rel, "rationale": f"Set by {coder} in review."})
    if problems:
        raise SystemExit("the sheet has answers that are not codes; nothing was applied:\n  " + "\n  ".join(problems[:20]))
    tally = write_codes(corpus, reg, recs, coder=coder, kind="human") if recs else {}
    C.counts_header("Review: applied", [("coder", coder), ("answers in the sheet", len(recs)),
                                        ("new or changed", tally.get("coded", 0)), ("unchanged", tally.get("unchanged", 0))])
    return 0


def cmd_reply(corpus: Path, src: Path, by: str, title: str | None, date: str | None, plan: bool) -> str | None:
    """A reviewer's reply (an email or a note, saved as Markdown or text) as a source, so a check or
    a clearance can cite it (`signoff --ref`). -> its source_id."""
    if not SLUG.match(by or ""):
        raise SystemExit("--by: the reviewer's handle (lower-case letters, digits, dashes)")
    src = Path(src).expanduser().resolve()
    if src.suffix.lower() not in (".md", ".txt") or not src.is_file():
        raise SystemExit(f"--from: a .md or .txt file (save the email's text first): {src}")
    if not src.read_text(encoding="utf-8").strip():
        raise SystemExit(f"--from: {src.name} is empty")
    day = date or dt.date.today().isoformat()
    md = src.suffix.lower() == ".md"
    rec = {"slug": f"reply-{by}-{day}-{hashlib.sha256(src.read_bytes()).hexdigest()[:6]}",
           "title": title or f"Reply from {by}, {day}", "source_kind": "document",
           "retrieval": {"fetch_status": "success", "fetch_method": "manual",
                         "content_type": "text/markdown" if md else "text/plain", "fetched_at": now(),
                         "sha256": "sha256:" + hashlib.sha256(src.read_bytes()).hexdigest()},
           "discovery": {"query": f"A reviewer's reply about one-pager figures ({by}, {day})"},
           "tags": ["one-pager", "reviewer-reply"],
           "files": [{"path": str(src), "role": "markdown" if md else "text",
                      "media_type": "text/markdown" if md else "text/plain", "produced_by": "import"}],
           "provenance": {"produced_by": {**TOOL, "method": "import"}, "created_at": now()},
           "ext": {"upc_corpus": {"one_pager": {"reply_by": by, "date": day}}}}
    res = C.upc("add", "source", "--corpus", str(corpus), "--tool", TOOL["tool"], *(["--dry-run"] if plan else []),
                check=False, stdin=json.dumps(rec))
    r = ((res or {}).get("results") or [{}])[0] if isinstance(res, dict) else {}
    if r.get("status") not in ("added", "planned", "exists"):
        raise SystemExit(f"the reply was not saved: {r.get('detail') or r.get('error') or res}")
    if not plan:
        C.upc("regen", str(corpus), check=False)
    C.counts_header("Reply" + (" (plan)" if plan else ""), [("from", by), ("date", day), ("status", r["status"]),
                                                          ("source", r.get("source_id"))],
                    [f"Cite it: signoff --figure <id> --by {by} --check confirmed --ref {r.get('source_id')}"])
    return r.get("source_id")


def cmd_signoff(corpus: Path, figures: list[str], by: str, check: str | None, clearance: str | None,
                ref: str | None, note: str | None) -> int:
    if not SLUG.match(by or ""):
        raise SystemExit("--by: the reviewer's handle (lower-case letters, digits, dashes)")
    if bool(check) == bool(clearance):
        raise SystemExit("give one of --check or --clearance")
    reg = load_reg(corpus)
    unknown = [x for x in figures if x not in reg]
    if unknown:
        raise SystemExit(f"--figure: not in the register: {', '.join(unknown)}")
    why = (note or "").strip()
    if ref:
        info = D.source_index(corpus).get(ref)
        if not info:
            raise SystemExit(f"--ref: no source {ref!r} in this corpus (save the reply as a source first)")
        why = (why + " " if why else "") + f"(See {info['source_id']}: {info['title'][:120]}.)"
    book, code = (BOOKS["check"][0], check) if check else (BOOKS["clearance"][0], clearance)
    tally = write_codes(corpus, reg, [{"figure": x, "book": book, "code": code, "rationale": why or f"Recorded for {by}."}
                                      for x in figures], coder=by, kind="human")
    C.counts_header("Signoff", [("by", by), ("what", "check" if check else "clearance"), ("answer", code),
                                ("figures", len(figures)), ("new or changed", tally.get("coded", 0)),
                                ("unchanged", tally.get("unchanged", 0))])
    return 0


# ------------------------------------------------------------------ a page: its shape and its numbers
def page_figures(page: dict) -> list[str]:
    """Every figure a page cites, in the order it cites them."""
    ids = [t.get("figure") for t in page.get("tiles") or []]
    for st in (page.get("how") or {}).get("steps") or []:
        ids += st.get("figures") or []
    ids += (page.get("status") or {}).get("figures") or []
    for blk in ("hero", "about"):
        ch = (page.get(blk) or {}).get("chart")
        for s in (ch or {}).get("series") or []:
            ids += [p for p in s.get("points") or [] if p]
        ids += [r.get("figure") for r in (ch or {}).get("refs") or []]
    ids += page.get("figures") or []
    return list(dict.fromkeys(x for x in ids if x))


def page_texts(page: dict) -> list[tuple[str, str]]:
    """(where, text): every piece of the page's own words a reader sees. A verbatim quotation is the
    source's words, checked at hop C, so it is not here; axis labels and ticks are not either."""
    out = [("headline", page.get("headline")), ("subhead", page.get("subhead"))]
    for i, t in enumerate(page.get("tiles") or []):
        out += [(f"tiles[{i}].{k}", t.get(k)) for k in ("big", "unit", "line", "source_line")]
    how = page.get("how") or {}
    out.append(("how.title", how.get("title")))
    for i, s in enumerate(how.get("steps") or []):
        out += [(f"how.steps[{i}].title", s.get("title")), (f"how.steps[{i}].text", s.get("text"))]
    st = page.get("status") or {}
    out += [(f"status.{k}", st.get(k)) for k in ("shown", "underway", "next")]
    out += [(f"status.labels[{i}]", x) for i, x in enumerate(st.get("labels") or [])]
    cta = page.get("cta") or {}
    out += [("cta.ask", cta.get("ask")), ("cta.detail", cta.get("detail"))]
    if page.get("voice"):
        out.append(("voice.attribution", page["voice"].get("attribution")))
    for blk in ("hero", "about"):
        b = page.get(blk) or {}
        out += [(f"{blk}.{k}", b.get(k)) for k in ("title", "sub", "note", "caption")]
        ch = b.get("chart") or {}
        out += [(f"{blk}.chart.{k}", ch.get(k)) for k in ("title", "sub", "note")]
        out += [(f"{blk}.chart.series[{i}].label", s.get("label")) for i, s in enumerate(ch.get("series") or [])]
        out += [(f"{blk}.chart.refs[{i}].label", r.get("label")) for i, r in enumerate(ch.get("refs") or [])]
    out += [(f"about.items[{i}]", x) for i, x in enumerate((page.get("about") or {}).get("items") or [])]
    return [(w, str(t)) for w, t in out if t]


def check_shape(page: dict, reg: dict, proj: dict) -> list[str]:
    errs = []
    if not isinstance(page, dict):
        return ["a page is a JSON object (references/one-pager.md)"]
    for k in ("headline", "subhead"):
        if not str(page.get(k) or "").strip():
            errs.append(f"{k}: missing")
    tiles = page.get("tiles") or []
    if not 1 <= len(tiles) <= 3:
        errs.append("tiles: one to three")
    for i, t in enumerate(tiles):
        if t.get("figure") not in reg:
            errs.append(f"tiles[{i}].figure: {t.get('figure')!r} is not in the register")
        if not str(t.get("big") or "").strip() or not str(t.get("line") or "").strip():
            errs.append(f"tiles[{i}]: big and line are needed")
    for fid in page_figures(page):
        if fid not in reg:
            errs.append(f"{fid!r}: cited, but not in the register")
    for blk in ("hero", "about"):
        b = page.get(blk) or {}
        ch = b.get("chart")
        if ch:
            if ch.get("type") not in ("bars", "lines"):
                errs.append(f"{blk}.chart.type: bars or lines")
            cats = ch.get("groups") if ch.get("type") == "bars" else ch.get("x")
            if not cats:
                errs.append(f"{blk}.chart: {'groups' if ch.get('type') == 'bars' else 'x'} (the category labels) are needed")
            for i, s in enumerate(ch.get("series") or []):
                pts = s.get("points") or []
                if len(pts) != len(cats or []):
                    errs.append(f"{blk}.chart.series[{i}]: {len(pts)} points for {len(cats or [])} categories")
                for p in pts:
                    if p and p in reg and number_of(reg[p]["value"]) is None:
                        errs.append(f"{blk}.chart.series[{i}]: figure {p} has no number to plot")
            if not ch.get("series"):
                errs.append(f"{blk}.chart.series: at least one")
            for i, r in enumerate(ch.get("refs") or []):
                if not isinstance(r, dict) or r.get("figure") not in reg or number_of(reg[r["figure"]]["value"]) is None:
                    errs.append(f"{blk}.chart.refs[{i}]: a figure in the register, with a number to draw the line at")
        if b.get("image") and b["image"] not in {x["id"] for x in opt(proj).get("assets") or []}:
            errs.append(f"{blk}.image: {b['image']!r} is not an asset in the project profile")
    if (page.get("how") or {}).get("image") and page["how"]["image"] not in {x["id"] for x in opt(proj).get("assets") or []}:
        errs.append(f"how.image: {page['how']['image']!r} is not an asset in the project profile")
    steps = (page.get("how") or {}).get("steps") or []
    if steps and not 2 <= len(steps) <= 4:
        errs.append("how.steps: two to four")
    v = page.get("voice")
    if v and not all(str(v.get(k) or "").strip() for k in ("source", "quote", "attribution")):
        errs.append("voice: source, quote and attribution are needed")
    for i, p in enumerate(page.get("plain_numbers") or []):
        if not isinstance(p, dict) or not str(p.get("text") or "").strip() or not str(p.get("why") or "").strip():
            errs.append(f"plain_numbers[{i}]: text and why are needed")
    return errs


def number_problems(page: dict, reg: dict, proj: dict) -> list[str]:
    """Every number the page's words print must be one of its figures (a value or a display form),
    one of the project's own facts (`allow_numbers`), or a plain number the page names with a reason.
    A tile's big number must be its own figure's."""
    figs = [f for f in page_figures(page) if f in reg]
    allowed: set[str] = set()
    for fid in figs:
        for s in [reg[fid]["value"], *shown_forms(reg[fid])]:
            allowed |= set(num_tokens(s))
    for s in opt(proj).get("allow_numbers") or []:
        allowed |= set(num_tokens(s))
    for p in page.get("plain_numbers") or []:
        allowed |= set(num_tokens((p or {}).get("text")))
    probs = []
    for where, text in page_texts(page):
        bad = [t for t in num_tokens(text) if t not in allowed and not YEAR.match(t)]
        if bad:
            probs.append(f"{where}: {', '.join(dict.fromkeys(bad))} in {text[:90]!r} is not a figure on this page")
    for i, t in enumerate(page.get("tiles") or []):
        f = reg.get(t.get("figure"))
        if f:
            own = set().union(*(set(num_tokens(s)) for s in [f["value"], *shown_forms(f)]))
            bad = [x for x in num_tokens(t.get("big")) if x not in own]
            if bad:
                probs.append(f"tiles[{i}].big: {', '.join(bad)} is not a form of its figure {t['figure']} "
                             f"({' / '.join(shown_forms(f))})")
    return probs


def current_page(corpus: Path, slug: str) -> dict | None:
    best = None
    for sj in corpus.glob("syntheses/*/synthesis.json"):
        y = C.load_json(sj, {}) or {}
        if y.get("type") != PAGE or y.get("superseded_by"):
            continue
        if (((y.get("ext") or {}).get("upc_corpus") or {}).get("one_pager") or {}).get("slug") != slug:
            continue
        if best is None or (y.get("provenance") or {}).get("created_at", "") > (best.get("provenance") or {}).get("created_at", ""):
            best = y
    return best


def slot_lines(page: dict) -> dict[str, str]:
    return {w: re.sub(r"\s+", " ", t).strip() for w, t in page_texts(page)}


def page_markdown(page: dict, reg: dict, idx: dict, aud: dict, proj: dict, writer: str, rev: int,
                  voice: dict | None) -> tuple[str, list[dict], set[str]]:
    """The synthesis output: the page's words with a bare citation per figure, then each figure's
    passage quoted under hop C. -> (markdown, claims, source ids)."""
    used: set[str] = set()

    def ev(fids: list[str]) -> tuple[list[str], list[str]]:
        eids, sids = [], []
        for x in closure(fids, reg):
            f = reg[x]
            for e, s in [(f.get("extraction_id"), f.get("source_id"))] + [(r["extraction_id"], r["source_id"]) for r in f.get("readings") or []]:
                if e and e not in eids:
                    eids.append(e)
                    sids.append(s)
        return eids, sorted(set(sids))

    def cite(fids):
        return " ".join(f"[{e}]" for e in ev(fids)[0])

    claims = []

    def claim(text, fids):
        eids, sids = ev(fids)
        if eids:
            used.update(sids)
            claims.append({"claim_id": f"cl-{len(claims) + 1:04d}", "text": P(text)[:600], "evidence_ids": eids,
                           "source_ids": sids, "confidence": "medium"})

    L = [f"# {P(page['headline'])}", "",
         f"A one-pager for {P(aud['name'])} on {P(proj['name'])}, revision {rev}. Words by {P(writer)}.", "",
         P(page["subhead"]), "", "## The numbers", ""]
    for t in page.get("tiles") or []:
        text = f"{t['big']} {t.get('unit') or ''}: {t['line']}"
        L.append(f"- {P(text)} {cite([t['figure']])}")
        claim(text, [t["figure"]])
    how = page.get("how") or {}
    if how.get("steps"):
        L += ["", f"## {P(how.get('title') or 'How it works')}", ""]
        for i, s in enumerate(how["steps"], 1):
            text = f"{s.get('title', '')}: {s.get('text', '')}"
            L.append(f"{i}. {P(text)} {cite(s.get('figures') or [])}".rstrip())
            claim(text, s.get("figures") or [])
    st = page.get("status") or {}
    if any(st.get(k) for k in ("shown", "underway", "next")):
        labels = st.get("labels") or ["Shown so far", "Underway", "Next"]
        L += ["", "## Where it stands", ""]
        L += [f"- {P(lab)}: {P(st[k])}" for lab, k in zip(labels, ("shown", "underway", "next")) if st.get(k)]
        if st.get("figures"):
            L.append(f"  (Figures: {cite(st['figures'])})")
            claim(" ".join(st.get(k) or "" for k in ("shown", "underway", "next")), st["figures"])
    cta = page.get("cta") or {}
    if cta.get("ask"):
        L += ["", "## The ask", "", P(cta["ask"]) + (" " + P(cta["detail"]) if cta.get("detail") else "")]
    if voice:
        L += ["", "## In their words", "", D.marker(voice["quote"], voice["extraction_id"]), f"({P(voice['attribution'])})"]
        used.add(voice["source_id"])
        claims.append({"claim_id": f"cl-{len(claims) + 1:04d}", "text": P(f"In their words: {voice['attribution']}"),
                       "evidence_ids": [voice["extraction_id"]], "source_ids": [voice["source_id"]], "confidence": "high"})
    L += ["", "## The figures", ""]
    for x in closure(page_figures(page), reg):
        f = reg[x]
        L += [f"### {P(f['label'])}", "", f"{P(f['value'])}{' ' + P(f['unit']) if f.get('unit') else ''}"
              + (f", shown as {P(' / '.join(shown_forms(f)))}" if shown_forms(f) != [f["value"]] else ""), ""]
        if f.get("extraction_id"):
            L += [D.marker(f["quote"], f["extraction_id"]), ""]
            used.add(f["source_id"])
        for r in f.get("readings") or []:
            L += [f"A reading: {P(r['value'])}", "", D.marker(r["quote"], r["extraction_id"]), ""]
            used.add(r["source_id"])
        if f.get("derived"):
            L += [f"Worked out from {', '.join(f['derived']['from'])}: {P(f['derived']['rule'])}", ""]
    L += ["## Sources", ""] + [f"- {sid}: {P((idx.get(sid) or {}).get('title') or sid)}" for sid in sorted(used)] + [""]
    return "\n".join(L), claims, used


def add_page(corpus: Path, slug: str, page: dict, *, writer: str, model: str | None = None, method: str = "manual",
             plan: bool = False, prompt: str = "onepager-page@1") -> dict:
    """Check a page's words and figures and store it as the next revision (an x-one-pager synthesis)."""
    cfg = load_cfg(corpus, slug)
    aud, proj = load_profiles(cfg)
    reg, idx = load_reg(corpus), D.source_index(corpus)
    errs = check_shape(page, reg, proj)
    if errs:
        raise SystemExit("the page does not fit (references/one-pager.md, 'A page'):\n  " + "\n  ".join(errs[:20]))
    nums = number_problems(page, reg, proj)
    if nums:
        raise SystemExit("refusing the page: it prints numbers that are not its figures (add the figure, use one of its "
                         "display forms, or name the number in plain_numbers with a reason):\n  " + "\n  ".join(nums[:20]))
    J = judgements(corpus, reg)
    figs = page_figures(page)
    hard = {x: gate(x, reg, J, aud["id"], cfg["distribution"])[0] for x in closure(figs, reg)}
    blocked = {x: h for x, h in hard.items() if h}
    if blocked:
        raise SystemExit("refusing the page: it uses figures withheld from every page:\n  "
                         + "\n  ".join(f"{x}: {'; '.join(h)}" for x, h in blocked.items()))
    voice = None
    if page.get("voice"):
        v = page["voice"]
        info = idx.get(v["source"])
        if not info:
            raise SystemExit(f"voice.source: no source {v['source']!r} with a reading copy in this corpus")
        g = anchor(corpus, info, [v["quote"]], typ="evidence", model=model or writer, prompt=prompt, dry=plan)[0]
        if not g.get("extraction_id") and not (plan and "status" not in g):
            raise SystemExit(f"voice.quote is not in its source ({g.get('status')}): {v['quote'][:100]!r}")
        voice = {"source_id": info["source_id"], "quote": g["quote"], "extraction_id": g.get("extraction_id"),
                 "attribution": v["attribution"]}
    warn = {x: gate(x, reg, J, aud["id"], cfg["distribution"])[1] for x in closure(figs, reg)}
    report = {"figures": len(figs), "open": {x: w for x, w in warn.items() if w}}
    if plan:
        report["status"] = "plan"
        return report
    prior = current_page(corpus, slug)
    prior_op = (((prior or {}).get("ext") or {}).get("upc_corpus") or {}).get("one_pager") or {}
    rev = int(prior_op.get("revision") or 0) + 1
    before, after = slot_lines(prior_op.get("page") or {}), slot_lines(page)
    changes = [{"where": w, "before": before.get(w), "after": after.get(w)} for w in sorted(set(before) | set(after))
               if before.get(w) != after.get(w)] if prior else []
    md, claims, used = page_markdown(page, reg, idx, aud, proj, writer, rev, voice)
    if not claims:
        raise SystemExit("the page cites no figure with a stored passage; a one-pager needs at least one")
    md_path = sdir(corpus, "pages") / f"{slug}.md"
    md_path.write_text(md, encoding="utf-8")
    snap = {x: {k: reg[x].get(k) for k in ("label", "value", "unit", "display", "extraction_id", "source_id", "quote",
                                           "readings", "derived")} for x in closure(figs, reg)}
    coding_ids = sorted({a["coding_id"] for x in closure(figs, reg) for (fid, _), xs in J.items() if fid == x
                         for a in xs if a.get("coding_id")})
    ext = {"slug": slug, "revision": rev, "page": page, "figures": snap, "voice": voice, "writer": writer,
           "audience": {"id": aud["id"], "name": aud["name"], "digest": digest(aud)},
           "project": {"id": proj["id"], "name": proj["name"], "digest": digest(proj)},
           "layout": cfg["layout"], "distribution": cfg["distribution"], "codings": coding_ids, "changes": changes}
    produced = {**TOOL, "method": method, "prompt_version": prompt, **({"model": model} if model else {})}
    res = C.upc("add", "synthesis", "--corpus", str(corpus), "--output", str(md_path), "--tool", TOOL["tool"],
                *(["--supersedes", prior["synthesis_id"]] if prior else []), check=False, stdin=json.dumps({
                    "type": PAGE, "title": f"{cfg['title']} (revision {rev})", "claims": claims,
                    "question": f"One-pager {slug} r{rev}: {aud['name']}, {proj['name']}",
                    "provenance": {"produced_by": produced}, "ext": {"upc_corpus": {"one_pager": ext}}}))
    if not (isinstance(res, dict) and res.get("status") == "ok"):
        fails = (res or {}).get("failures", []) if isinstance(res, dict) else []
        raise SystemExit(f"the page was refused: {(res or {}).get('detail') if isinstance(res, dict) else res}\n  "
                         + "\n  ".join(f"{f.get('marker')}: {f.get('reason')}" for f in fails[:8]))
    cfg["revision"] = rev
    C.save_json(page_path(corpus, slug), cfg)
    C.upc("regen", str(corpus), check=False)
    report.update(status="ok", synthesis_id=res["synthesis_id"], revision=rev, changes=len(changes))
    return report


def open_groups(open_: dict[str, list[str]]) -> list[tuple[str, list[str]]]:
    """What keeps figures off a final page, one line per reason with the figures it holds back,
    the commonest first, so 30 unchecked figures read as one line, not 30."""
    by: dict[str, list[str]] = {}
    for fid, reasons in open_.items():
        for r in reasons:
            by.setdefault(r, []).append(fid)
    return sorted(by.items(), key=lambda kv: (-len(kv[1]), kv[0]))


def print_open(open_: dict[str, list[str]]) -> None:
    for reason, fids in open_groups(open_):
        print(f"  OPEN ({len(fids)}) {reason}: {', '.join(fids)}")


def print_page_report(rep: dict) -> None:
    C.counts_header("Page", [("status", rep.get("status")), ("revision", rep.get("revision", "—")),
                             ("synthesis", rep.get("synthesis_id", "—")), ("figures cited", rep["figures"]),
                             ("figures still open for a final page", len(rep["open"])),
                             ("wording changes since the last revision", rep.get("changes", "—"))])
    print_open(rep["open"])


# ------------------------------------------------------------------ write (the local model drafts the page)
LAYOUT_GUIDE = {
    "executive": "Open with the reader's own problem and what the technology does for their product, costs, "
                 "downtime and people; say only briefly how it works. Few numbers, each a reason to act. No method on the page.",
    "technical": "Lead with the result, then how it works in three steps and where it stands. Numbers can be "
                 "precise, in the units the reader uses on the floor.",
    "policy": "Lead with impact at the reader's scale (the state or the nation: energy, emissions, jobs, the "
              "industry), then what the project has shown and what support would speed it. The ask is a program, "
              "a meeting or a pilot.",
}


def write_prompt(aud: dict, proj: dict, cfg: dict, allowed: list[str]) -> str:
    o = opt(aud)
    return f"""You write a one-page brief for one group of readers, from figures that have been checked against their sources.

{reader_text(aud)}

{project_text(proj)}

Layout ({cfg['layout']}): {LAYOUT_GUIDE[cfg['layout']]}
The ask: {o.get('cta') or 'the next step for these readers'}{' (' + o['cta_detail'] + ')' if o.get('cta_detail') else ''}.

Output JSON only, in your own words:
- headline: at most 8 words, a benefit these readers care about.
- subhead: two sentences, at most 45 words.
- tiles: three (fewer if fewer figures are listed): figure (an id from the list), big (the number, as one of
  that figure's display forms, copied exactly), unit (short, or empty), line (one sentence of at most 22 words on
  why it matters to these readers), source_line (who found it and when, from the figure's source).
- how: title (one sentence) and three steps, each a title (at most 6 words), a text (at most 18 words) and the
  figures it uses.
- status: shown, underway, next (one or two short sentences each) and the figures they use.
- cta: ask (one sentence) and detail (one sentence: what the reader would share or do).
- limits: what the page leaves out that a careful reader should know (they go in the notes, not on the page).
- questions: up to three questions for the engineers about these figures.
Rules: every number you write must be one of the display forms listed for a figure you cite, copied exactly, or
one of these: {', '.join(allowed) or 'none'}. Write no other number, and no number as a word. Use no quotation
marks. Never name a person or a company."""


def page_schema(fig_ids: list[str]) -> dict:
    s = {"type": "string"}
    fig = {"type": "string", "enum": fig_ids}

    def obj(**p):
        return {"type": "object", "properties": p, "required": list(p)}

    return obj(headline=s, subhead=s,
               tiles={"type": "array", "minItems": min(3, len(fig_ids)), "maxItems": 3,
                      "items": obj(figure=fig, big=s, unit=s, line=s, source_line=s)},
               how=obj(title=s, steps={"type": "array", "minItems": 3, "maxItems": 3,
                                        "items": obj(title=s, text=s, figures={"type": "array", "items": fig})}),
               status=obj(shown=s, underway=s, next=s, figures={"type": "array", "items": fig}),
               cta=obj(ask=s, detail=s), limits={"type": "array", "items": s},
               questions={"type": "array", "items": s, "maxItems": 3})


def cmd_write(corpus: Path, slug: str, plan: bool) -> int:
    cfg = load_cfg(corpus, slug)
    aud, proj = load_profiles(cfg)
    reg, idx = load_reg(corpus), D.source_index(corpus)
    J = judgements(corpus, reg)
    elig = [x for x in sorted(reg) if not gate(x, reg, J, aud["id"], cfg["distribution"])[0]
            and status_of(x, reg, J, aud["id"])["rel"] in ("lead", "support")]
    elig.sort(key=lambda x: status_of(x, reg, J, aud["id"])["rel"] != "lead")
    C.counts_header("Write: plan", [("audience", aud["name"]), ("layout", cfg["layout"]),
                                    ("figures judged lead or support", len(elig)),
                                    ("model", f"{C.DEFAULT_MODEL} at {C.CHAT_URL}")])
    if not elig:
        raise SystemExit(f"no figure is judged lead or support for {aud['id']} yet: run `code` or `review --apply` first")
    for x in elig:
        print(f"  {x}: {reg[x]['value']} ({status_of(x, reg, J, aud['id'])['rel']})")
    if plan:
        return 0
    allowed = list(opt(proj).get("allow_numbers") or [])
    lines = []
    for x in elig:
        f, s = reg[x], status_of(x, reg, J, aud["id"])
        why = next((a["rationale"] for a in J.get((x, rel_slug(aud["id"])), []) if a.get("rationale")), "")
        src = (idx.get(f.get("source_id") or "") or {}).get("title") or "worked out from " + ", ".join((f.get("derived") or {}).get("from", []))
        lines.append(f"- {x}: {f['label']} = {f['value']}{' ' + f['unit'] if f.get('unit') else ''}; display forms: "
                     f"{' | '.join(shown_forms(f))}; class: {s['class'] or 'not set'}; for these readers: {s['rel']}"
                     f"{' (' + why + ')' if why else ''}; source: {src}")
    try:
        out = C.llm_chat(write_prompt(aud, proj, cfg, allowed), "Figures you may use:\n" + "\n".join(lines),
                         schema=page_schema(elig), max_tokens=3000, temperature=0.3)
    except C.Unavailable as e:
        raise SystemExit(f"UNAVAILABLE: {e}") from None
    d = C.parse_json_loose(out["content"]) or {}
    draft = sdir(corpus, "pages") / f"{slug}.draft.json"
    C.save_json(draft, d)
    writer = f"{C.AGENT_NAME} ({C.DEFAULT_MODEL})"
    try:
        rep = add_page(corpus, slug, d, writer=writer, model=C.DEFAULT_MODEL, method="model", prompt="onepager-write@1")
    except SystemExit as e:
        print(f"  NOT STORED: {e}\n  The model's draft is in {draft}: edit it, then run `add --from {draft}`.")
        return 1
    print_page_report(rep)
    return 0


# ------------------------------------------------------------------ rendering: charts, the process flow, the page
TONE = {"base": "var(--d-base)", "hot": "var(--d-hot)", "cool": "var(--d-cool)"}


def _fmt(v: float) -> str:
    return f"{v:g}"


def _nice(x: float) -> float:
    if x <= 0:
        return 1
    e = 10 ** math.floor(math.log10(x))
    f = x / e
    return (1 if f <= 1 else 2 if f <= 2 else 5 if f <= 5 else 10) * e


def _ticks(ch: dict, vals: list[float]) -> tuple[float, float, list[tuple[float, str]]]:
    y = ch.get("y") or {}
    lo = float(y.get("min", 0))
    hi = float(y.get("max") or (max(vals + [0]) * 1.15 or 1))
    if y.get("ticks"):
        return lo, hi, [(float(t[0]), str(t[1])) if isinstance(t, list) else (float(t), _fmt(float(t))) for t in y["ticks"]]
    step = _nice((hi - lo) / 4)
    ticks, v = [], lo
    while v <= hi + 1e-9:
        ticks.append((v, _fmt(round(v, 10))))
        v += step
    return lo, hi, ticks


def _legend(series: list[dict], W: int, refs: list[dict] = ()) -> tuple[list[str], int]:
    """A key for each series, then a dashed key for each reference line (a target or a limit)."""
    out, x, y = [], 0, 14
    for s in list(series) + [{**r, "_ref": True} for r in refs]:
        sw = 24 if s.get("_ref") else 12
        w = sw + 10 + len(str(s.get("label", ""))) * 6.6 + 16
        if x and x + w > W:
            x, y = 0, y + 18
        out.append((f'<line x1="{x:.0f}" x2="{x + sw:.0f}" y1="{y - 3}" y2="{y - 3}" stroke="var(--ink-2)" stroke-width="1.5" stroke-dasharray="5 3"/>'
                    if s.get("_ref") else
                    f'<rect x="{x:.0f}" y="{y - 9}" width="12" height="12" rx="2" fill="{TONE.get(s.get("tone"), TONE["cool"])}"/>')
                   + f'<text x="{x + sw + 6:.0f}" y="{y + 1}" class="lab">{esc(s.get("label", ""))}</text>')
        x += w
    return out, y + 14


def chart_svg(ch: dict, reg: dict) -> str:
    """A bar or line chart drawn from figures: every plotted value is a figure's own value, and each
    bar or point is labelled with it."""
    W = 380
    series = ch.get("series") or []
    vals = [[number_of(reg[p]["value"]) if p else None for p in s.get("points") or []] for s in series]
    flat = [v for vs in vals for v in vs if v is not None]
    lo, hi, ticks = _ticks(ch, flat)
    leg, top = _legend(series, W, ch.get("refs") or [])
    m = {"l": max(30, 10 + 7 * max([len(str(lab)) for _, lab in ticks] or [1])), "r": 8, "t": top + 10, "b": 34}
    ph = 132
    H = m["t"] + ph + m["b"]
    pw = W - m["l"] - m["r"]

    def Y(v):
        return m["t"] + ph - (max(lo, min(hi, v)) - lo) / ((hi - lo) or 1) * ph

    g = []
    for v, lab in ticks:
        g.append(f'<line x1="{m["l"]}" x2="{W - m["r"]}" y1="{Y(v):.1f}" y2="{Y(v):.1f}" stroke="var(--rule)" stroke-width="1"/>'
                 f'<text x="{m["l"] - 6}" y="{Y(v) + 4:.1f}" text-anchor="end">{esc(lab)}</text>')
    for r in ch.get("refs") or []:  # a target or a limit, drawn at its own figure's value
        y = Y(number_of(reg[r["figure"]]["value"]))
        g.append(f'<line x1="{m["l"]}" x2="{W - m["r"]}" y1="{y:.1f}" y2="{y:.1f}" stroke="var(--ink-2)" stroke-width="1.5" '
                 'stroke-dasharray="5 3"/>')
    if ch.get("type") == "bars":
        groups = ch.get("groups") or []
        gw = pw / max(1, len(groups))
        bw = min(40.0, gw * 0.8 / max(1, len(series)))  # wide enough for a label like "≥ 5.05"
        for gi, name in enumerate(groups):
            x0 = m["l"] + gi * gw + (gw - bw * len(series)) / 2
            for si, s in enumerate(series):
                v = vals[si][gi]
                if v is None:
                    continue
                f = reg[s["points"][gi]]
                x = x0 + si * bw
                g.append(f'<rect x="{x + 2:.1f}" y="{Y(v):.1f}" width="{bw - 4:.1f}" height="{Y(lo) - Y(v):.1f}" '
                         f'fill="{TONE.get(s.get("tone"), TONE["cool"])}"/>'
                         f'<text x="{x + bw / 2:.1f}" y="{Y(v) - 4:.1f}" text-anchor="middle" class="val">{esc(f["value"])}</text>')
            g.append(f'<text x="{m["l"] + gi * gw + gw / 2:.1f}" y="{H - 14}" text-anchor="middle" class="lab">{esc(name)}</text>')
    else:
        xs = ch.get("x") or []
        step = pw / max(1, len(xs) - 1)

        def X(i):
            return m["l"] + i * step if len(xs) > 1 else m["l"] + pw / 2

        for i, name in enumerate(xs):
            g.append(f'<text x="{X(i):.1f}" y="{H - 14}" text-anchor="middle">{esc(name)}</text>')
        for si, s in enumerate(series):
            pts = [(X(i), Y(v)) for i, v in enumerate(vals[si]) if v is not None]
            col = TONE.get(s.get("tone"), TONE["cool"])
            g.append(f'<path d="{"".join(("L" if j else "M") + f"{x:.1f},{y:.1f}" for j, (x, y) in enumerate(pts))}" '
                     f'fill="none" stroke="{col}" stroke-width="{3 if s.get("tone") == "cool" else 2.2}" stroke-linejoin="round"/>')
            g += [f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.6" fill="{col}" stroke="var(--sheet)" stroke-width="1.5"/>' for x, y in pts]
    desc = "; ".join(f"{s.get('label')}: " + ", ".join(reg[p]["value"] for p in s.get("points") or [] if p) for s in series)
    desc += "".join(f"; {r.get('label')}: {reg[r['figure']]['value']}" for r in ch.get("refs") or [])
    return (f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="{esc((ch.get("title") or "Chart") + ". " + desc)}">'
            + "".join(leg) + "".join(g) + "</svg>")


# A number and its unit stay on one line, on screen and on paper ("55 °C", "28 days"), and so does a
# range ("1 to 2").
UNIT_GAP = re.compile(r"(\d) +(°[CF]?|%|days?\b|weeks?\b|minutes?\b|min\b|hours?\b|ppm\b|mL\b|log\b|kWh\b|MWh\b|kW\b|MMBtu\b|therms?\b|gpm\b|GPM\b|gallons?\b|cfm\b|CFM\b|psig?\b|Pa\b)")
RANGE_GAP = re.compile(r"\b(\d[\d.,]*) (to|or|–) (?=\d)")  # \b: not the 2 of CO2
NUM_HYPHEN = re.compile(r"(?<=\d)-(?=[A-Za-z])")  # "5-log": a word joiner after the hyphen keeps it whole


def keep_units(v):
    """The page's words with a no-break space between each number and its unit, through lists and dicts."""
    if isinstance(v, str):
        return NUM_HYPHEN.sub("-\u2060", UNIT_GAP.sub("\\1\u00a0\\2", RANGE_GAP.sub("\\1\u00a0\\2\u00a0", v)))
    if isinstance(v, list):
        return [keep_units(x) for x in v]
    if isinstance(v, dict):
        return {k: keep_units(x) for k, x in v.items()}
    return v


def chart_table(ch: dict, reg: dict) -> str:
    """The values a chart plots, as a table: each one its figure's own value, with the unit they share."""
    series = ch.get("series") or []
    cats = (ch.get("groups") if ch.get("type") == "bars" else ch.get("x")) or []
    units = {reg[p].get("unit") or "" for s in series for p in s.get("points") or [] if p}
    cap = (ch.get("title") or "The chart") + (f", {units.pop()}" if len(units) == 1 and "" not in units else "")
    head = "<th scope=\"col\"></th>" + "".join(f'<th scope="col">{esc(s.get("label"))}</th>' for s in series)
    def cell(s, i):
        pts = s.get("points") or []
        return esc(reg[pts[i]]["value"]) if i < len(pts) and pts[i] else ""
    rows = "".join(f'<tr><th scope="row">{esc(c)}</th>' + "".join(f"<td>{cell(s, i)}</td>" for s in series) + "</tr>"
                   for i, c in enumerate(cats))
    return f'<div class="vals-wrap"><table class="vals"><caption>{esc(cap)}</caption><tr>{head}</tr>{rows}</table></div>'


def _wrap(text: str, width: int) -> list[str]:
    out, cur = [], ""
    for w in str(text or "").split(" "):  # not split(): that would break at a no-break space too
        if not w:
            continue
        if cur and len(cur) + 1 + len(w) > width:
            out.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    return out + ([cur] if cur else [])


def flow_svg(steps: list[dict]) -> str:
    """The steps as a simple process flow: numbered boxes, left to right, with arrows between."""
    n, W, gap = len(steps), 720, 28
    bw = (W - gap * (n - 1)) / n
    tl = [_wrap(s.get("title"), int((bw - 50) / 7.9)) for s in steps]
    dl = [_wrap(s.get("text"), int((bw - 20) / 6.6)) for s in steps]
    H = 20 + 17 * max(len(t) for t in tl) + 15.5 * max(len(d) for d in dl) + 13
    g = []
    for i, s in enumerate(steps):
        x = i * (bw + gap)
        g.append(f'<rect x="{x:.1f}" y="1" width="{bw:.1f}" height="{H - 2:.1f}" rx="6" fill="var(--tint-2)" stroke="var(--rule)"/>'
                 f'<circle cx="{x + 21:.1f}" cy="21" r="12" fill="var(--brand)"/>'
                 f'<text x="{x + 21:.1f}" y="26" text-anchor="middle" class="n">{i + 1}</text>')
        g += [f'<text x="{x + 42:.1f}" y="{26 + 17 * j}" class="t">{esc(line)}</text>' for j, line in enumerate(tl[i])]
        y0 = 26 + 17 * len(tl[i]) + 4
        g += [f'<text x="{x + 12:.1f}" y="{y0 + 15.5 * j:.1f}" class="d">{esc(line)}</text>' for j, line in enumerate(dl[i])]
        if i < n - 1:
            ax = x + bw + 5
            g.append(f'<path d="M{ax:.1f},{H / 2 - 9:.1f} L{ax + gap - 10:.1f},{H / 2:.1f} L{ax:.1f},{H / 2 + 9:.1f} Z" fill="var(--accent)"/>')
    label = " Then ".join(f"{i + 1}. {s.get('title')}: {s.get('text')}" for i, s in enumerate(steps))
    return f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="{esc(label)}">{"".join(g)}</svg>'


MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".svg": "image/svg+xml",
        ".gif": "image/gif", ".woff2": "font/woff2", ".woff": "font/woff", ".ttf": "font/ttf", ".otf": "font/otf"}


def data_uri(path: Path) -> str:
    mime = MIME.get(path.suffix.lower())
    if not mime:
        raise SystemExit(f"{path.name}: not an image or font this page can embed ({', '.join(MIME)})")
    return f"data:{mime};base64,{base64.b64encode(path.read_bytes()).decode()}"


def asset_path(cfg: dict, a: dict) -> Path:
    p = Path(str(a["path"])).expanduser()
    return p if p.is_absolute() else Path(cfg["project"]).parent / p


def page_line(page: dict, aud: dict, draft: bool) -> str:
    """The page's one line for link previews and indexes (<meta name="description">): the subhead, who it's
    for, and a draft's stamp. The first paragraph on the page is the funder line, not what the page says."""
    return " ".join(x for x in (page.get("subhead") or "", f"For {aud['name']}." if aud.get("name") else "",
                                "Internal draft." if draft else "") if x)


def page_html(ctx: dict) -> str:
    page, reg, aud, proj, cfg = keep_units(ctx["page"]), ctx["reg"], ctx["aud"], ctx["proj"], ctx["cfg"]
    brand, assets = opt(proj).get("brand") or {}, ctx["assets"]
    pal = brand.get("palette") or {}
    css = ((ASSETS / "page.css").read_text(encoding="utf-8").replace("__BRAND__", pal.get("brand", "#022851"))
           .replace("__ACCENT__", pal.get("accent", "#FFBF00")).replace("__HEAT__", pal.get("heat", "#F18A00"))
           .replace("__FAMILY__", ctx["font_family"]))
    font = (f'@font-face{{font-family:"{ctx["font_family"]}";src:url({ctx["font_uri"]}) format("{ctx["font_format"]}");'
            'font-weight:100 900;font-style:normal;font-display:block}')

    def img(aid, cls=""):
        a = assets[aid]
        return (f'<img class="{cls}" src="{a["uri"]}" alt="{esc(a["alt"])}">'
                + (f'<p class="credit">{esc(a["credit"])}</p>' if a.get("credit") else ""))

    def visual(b):
        if not b or not (b.get("chart") or b.get("image")):
            return ""
        ch = b.get("chart")
        head = "".join(f'<p class="{c}">{esc(t)}</p>' for c, t in (("fig-title", (ch or b).get("title")), ("fig-sub", (ch or b).get("sub"))) if t)
        body = f'<div class="chart">{chart_svg(ch, reg)}</div>' if ch else img(b["image"])
        note = (ch or b).get("note") or b.get("caption")
        return f'<figure class="visual">{head}{body}' + (f'<p class="fig-note">{esc(note)}</p>' if note else "") + "</figure>"

    org = (f'<img class="logo" src="{assets[brand["logo"]]["uri"]}" alt="{esc(assets[brand["logo"]]["alt"])}">'
           if brand.get("logo") else f'<span class="orgname">{esc(brand.get("org_line") or proj["name"])}</span>')
    stamp = f'<p class="stamp">{esc(ctx["stamp"])}</p>' if ctx["stamp"] else ""
    tag = opt(aud).get("tag") or f"For {aud['name']}"
    hero_v = visual(page.get("hero"))
    blocks = {}
    blocks["brand"] = (f'<div class="brandbar" data-block="brand bar"><div class="brandbar-in pad"><p class="org">{org}'
                       + (f'<span class="funder">{esc(brand["funder_line"])}</span>' if brand.get("funder_line") else "")
                       + f'</p>{stamp}</div></div>')
    blocks["hero"] = (f'<header class="hero pad" data-block="headline"><p class="tag">{esc(tag)}</p>'
                      f'<div class="hero-grid{"" if hero_v else " solo"}"><div class="hero-copy"><h1>{esc(page["headline"])}</h1>'
                      f'<p class="sub">{esc(page["subhead"])}</p></div>{hero_v}</div></header>')
    tiles = "".join(
        f'<div class="tile"><div class="big">{esc(t["big"])}' + (f'<small>{esc(t["unit"])}</small>' if t.get("unit") else "")
        + f'</div><p class="cap">{esc(t["line"])}</p>' + (f'<p class="src">{esc(t["source_line"])}</p>' if t.get("source_line") else "")
        + "</div>" for t in page["tiles"])
    blocks["tiles"] = f'<section class="tiles pad" aria-label="Key numbers" data-block="the numbers"><div class="tile-row">{tiles}</div></section>'
    how = page.get("how") or {}
    if how.get("steps"):
        blocks["how"] = (f'<section class="how pad" data-block="how it works"><p class="kicker">How it works</p>'
                         + (f'<h2>{esc(how["title"])}</h2>' if how.get("title") else "")
                         + f'<div class="flow">{flow_svg(how["steps"])}<ol class="steps">'
                         + "".join(f'<li><strong>{esc(s.get("title"))}</strong> {esc(s.get("text"))}</li>' for s in how["steps"])
                         + "</ol></div>"
                         + (f'<figure class="how-img">{img(how["image"])}</figure>' if how.get("image") else "") + "</section>")
    st = page.get("status") or {}
    if any(st.get(k) for k in ("shown", "underway", "next")):
        labels = st.get("labels") or ["Shown so far", "Underway", "Next"]
        cols = "".join(f'<div><h3>{esc(lab)}</h3><p>{esc(st[k])}</p></div>'
                       for lab, k in zip(labels, ("shown", "underway", "next")) if st.get(k))
        blocks["status"] = (f'<section class="status pad" data-block="where it stands"><p class="kicker">Where it stands</p>'
                            f'<div class="ledger">{cols}</div></section>')
    if ctx.get("voice"):
        v = ctx["voice"]
        blocks["voice"] = (f'<section class="voice-wrap pad" data-block="in their words"><blockquote class="voice">'
                           f'<p>“{esc(v["quote"])}”</p><cite>{esc(v["attribution"])}</cite></blockquote></section>')
    cta = page.get("cta") or {}
    if cta.get("ask"):
        blocks["cta"] = (f'<section class="cta-wrap pad" data-block="the ask"><p class="cta"><strong>{esc(cta["ask"])}</strong>'
                         + (f'<span>{esc(cta["detail"])}</span>' if cta.get("detail") else "") + "</p></section>")
    about = page.get("about") or {}
    hero_ch = (page.get("hero") or {}).get("chart")
    if about.get("items") or about.get("chart") or hero_ch:
        blocks["about"] = (f'<section class="data pad" data-block="about the data"><details><summary>'
                           f'{esc(about.get("title") or "About the data")} <span>{esc(about.get("sub") or "Methods and sources")}</span>'
                           f'</summary>{visual(about)}<ul class="method">' + "".join(f"<li>{esc(x)}</li>" for x in about.get("items") or [])
                           + "</ul>" + (chart_table(hero_ch, reg) if hero_ch else "") + "</details></section>")
    foot = " · ".join(x for x in (brand.get("footer"), brand.get("award_line"), brand.get("result_class")) if x)
    blocks["footer"] = (f'<footer class="pad" data-block="footer"><p class="proj">{esc(proj.get("full_name") or proj["name"])}</p>'
                        + (f"<p>{esc(foot)}</p>" if foot else "") + "</footer>")
    order = {"executive": ("brand", "hero", "tiles", "voice", "how", "cta", "status", "about", "footer"),
             "technical": ("brand", "hero", "tiles", "how", "status", "cta", "voice", "about", "footer"),
             "policy": ("brand", "hero", "tiles", "status", "cta", "how", "voice", "about", "footer")}[cfg["layout"]]
    body = "\n".join(blocks[b] for b in order if b in blocks)
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<!-- One-pager {esc(cfg['slug'])}, revision {ctx['revision']} ({esc(ctx['synthesis_id'])}), built by {TOOL['tool']} {TOOL['tool_version']}.
     Every number on it is listed, with its passage and who checked it, in {esc(cfg['slug'])}-figures.md. -->
<title>{esc(page['headline'])}</title>
<meta name="description" content="{esc(page_line(page, aud, ctx['draft']))}">
<style>{font}
{css}</style>
</head>
<body class="layout-{cfg['layout']}">
<div class="page"><article class="sheet">
{body}
</article></div>
</body>
</html>
"""


# ------------------------------------------------------------------ the figures note (the build's own 05a)
def figures_markdown(ctx: dict, e=str) -> str:
    """The figures note. `e` escapes every piece of source or writer text: str for the Markdown file,
    dossier.md for the copy pandoc turns into Word (a passage that is a table row would otherwise
    be read as Markdown and lost)."""
    page, reg, aud, proj, cfg, J = ctx["page"], ctx["reg"], ctx["aud"], ctx["proj"], ctx["cfg"], ctx["J"]
    idx, exts = ctx["idx"], ctx["exts"]
    figs = closure(page_figures(page), reg)
    L = [f"# {e(cfg['title'])}: the figures", "",
         f"Page `{cfg['slug']}`, revision {ctx['revision']}, for {e(aud['name'])}. "
         + ("An internal draft, not for distribution." if ctx["draft"] else f"Distribution: {cfg['distribution']}.")
         + f" {e(ctx['writers'])}", "",
         "Every number on the page is listed below with the passage it comes from, found again in the saved source "
         "when this note was made, and with who has checked it. Fixity is not authenticity: a matching passage shows "
         "which words were read, not that they are right.", ""]
    open_ = {x: w for x, w in ctx["open"].items() if w}
    if open_ or ctx["voice_note"] or ctx["asset_notes"]:
        L += ["## Before this goes outside the team", ""]
        L += [f"- {e(reason[0].upper() + reason[1:])} ({len(fids)}): {e(', '.join(fids))}." for reason, fids in open_groups(open_)]
        L += [f"- {e(n)}" for n in ctx["voice_note"] + ctx["asset_notes"]] + [""]
        who = [c for c in opt(proj).get("clearances") or [] if isinstance(c, dict) and c.get("what")]
        if who:
            L += ["Who may release what (the project profile):", ""]
            L += [f"- {e(c['what'])}: {e(c.get('who') or 'not named')}" for c in who] + [""]
    L += ["## Every number on the page", "", "| Figure | On the page | Class | Checked | Release | For these readers |",
          "|---|---|---|---|---|---|"]
    for x in figs:
        s = status_of(x, reg, J, aud["id"])
        L.append(f"| {e(x)} | {e(' / '.join(shown_forms(reg[x])))} | {s['class'] or '—'} | {s['check'] or 'not yet'} | "
                 f"{s['clearance'] or 'not recorded'} | {s['rel'] or 'not judged'} |")
    L.append("")
    for x in figs:
        f, s = reg[x], status_of(x, reg, J, aud["id"])
        L += [f"### {e(f['label'])} ({e(x)})", "", f"**Value:** {e(f['value'])}{' ' + e(f['unit']) if f.get('unit') else ''}."
              + (f" **Shown as:** {e(' / '.join(shown_forms(f)))}." if shown_forms(f) != [f["value"]] else ""), ""]
        if f.get("rounded"):
            L += [f"The page shows it as {e(', '.join(f['rounded']))}, which is not how the source writes it: check the rounding.", ""]
        if f.get("extraction_id"):
            ext = exts.get(f["extraction_id"]) or {}
            src = idx.get(f["source_id"]) or {}
            pg = D.page_of(ext)
            L += [f"> {e(f['quote'])}", ">", f"> — {e(src.get('title') or f['source_id'])}" + (f", p. {pg}" if pg else "")
                  + f" ({f['source_id']}, passage {f['extraction_id']})", ""]
        for i, r in enumerate(f.get("readings") or [], 1):
            agree = {True: "agrees", False: "DISAGREES", None: "no tolerance set"}[r.get("agrees")]
            L += [f"Reading {i}: {e(r['value'])} ({agree}" + (f", tolerance {f['tolerance']}" if f.get("tolerance") is not None else "")
                  + f"), from {e((idx.get(r['source_id']) or {}).get('title') or r['source_id'])}:", "", f"> {e(r['quote'])}", ""]
        if f.get("derived"):
            L += [f"Worked out from {e(', '.join(f['derived']['from']))}: {e(f['derived']['rule'])}.", ""]
        if f.get("thin"):
            L += [f"**Thin:** {e(f['thin'])}.", ""]
        for book, what in ((BOOKS["class"][0], "Class"), (BOOKS["check"][0], "Checked"), (BOOKS["clearance"][0], "Release"),
                           (rel_slug(aud["id"]), "For these readers")):
            xs = J.get((x, book), [])
            if xs:
                L += [f"- {what}: " + "; ".join(f"{e(a['code'])} ({e(a['coder'])}" + (f": {e(a['rationale'])}" if a.get("rationale") else "") + ")"
                                                for a in sorted(xs, key=lambda a: -KIND_RANK.get(a.get("kind"), 0)))]
        for what in s["disagree"]:
            book = {"relevance": rel_slug(aud["id"]), "class": BOOKS["class"][0], "check": BOOKS["check"][0],
                    "clearance": BOOKS["clearance"][0]}[what]
            xs = J.get((x, book), [])
            top = max(xs, key=lambda a: (KIND_RANK.get(a.get("kind"), 0), a.get("at") or ""))
            L += [f"- Coders disagree on {what}. " + (f"{e(top['coder'])}'s answer ({e(top['code'])}) is used: a person's answer outranks a model's."
                  if KIND_RANK.get(top.get("kind"), 0) >= KIND_RANK["human"] else
                  f"No person has answered, so the newest model answer ({e(top['code'])}, {e(top['coder'])}) is used until one does.")]
        if f.get("note"):
            L += [f"- Note: {e(f['note'])}"]
        L.append("")
    plain = page.get("plain_numbers") or []
    if plain:
        L += ["## Numbers on the page that are not figures", ""] + [f"- {e(p['text'])}: {e(p['why'])}" for p in plain] + [""]
    if ctx["changes"]:
        L += [f"## Wording changed since revision {ctx['revision'] - 1}", ""]
        L += [f"- {e(c['where'])}: {e(c['before'] or '(none)')} → {e(c['after'] or '(removed)')}" for c in ctx["changes"]] + [""]
    if page.get("limits"):
        L += ["## Left off the page", ""] + [f"- {e(x)}" for x in page["limits"]] + [""]
    qs = list(page.get("questions") or [])
    for x in figs:
        s = status_of(x, reg, J, aud["id"])
        if s["check"] == "query":
            qs.append(f"{x}: a reviewer's query is open.")
        if reg[x].get("thin"):
            qs.append(f"{x}: {reg[x]['thin']}. Does the claim still hold?")
        if any(r.get("agrees") is False for r in reg[x].get("readings") or []):
            qs.append(f"{x}: the readings disagree. Which value is right?")
        if reg[x].get("rounded"):
            qs.append(f"{x}: is it fair to show it as {', '.join(reg[x]['rounded'])}?")
    if qs:
        L += ["## Questions for the engineers", ""] + [f"- [ ] {e(q)}" for q in qs] + [""]
    used = sorted({reg[x][k] for x in figs for k in ("source_id",) if reg[x].get(k)}
                  | {r["source_id"] for x in figs for r in reg[x].get("readings") or []}
                  | ({ctx["voice"]["source_id"]} if ctx.get("voice") else set()))
    L += ["## Sources", ""]
    for sid in used:
        src = idx.get(sid) or {}
        L.append(f"- {sid}: {e(src.get('title') or sid)}" + (f" (SHA-256 {src['sha256'][:16]}…)" if src.get("sha256") else ""))
    return "\n".join(L) + "\n"


def make_docx(md_text: str, quotes: list[str], dest: Path, tmp: Path, day: str) -> tuple[str, list[str]]:
    """The figures note as Word (pandoc, the dossier's restyled reference). -> (its text, quotes missing from it).
    SOURCE_DATE_EPOCH, the page's own date, keeps an unchanged note's bytes the same."""
    C.require_tools("pandoc")
    src = tmp / "figures.md"
    src.write_text(md_text, encoding="utf-8")
    ref = D.reference_docx(tmp / "reference.docx")
    p = subprocess.run(["pandoc", str(src), "-f", "markdown-smart", "-t", "docx", "--reference-doc", str(ref), "-o", str(dest)],
                       capture_output=True, text=True, timeout=300,
                       env={**os.environ, "SOURCE_DATE_EPOCH": str(int(dt.datetime.fromisoformat(day[:10]).replace(tzinfo=dt.timezone.utc).timestamp()))})
    if p.returncode != 0:
        raise C.Unavailable(f"pandoc failed: {p.stderr.strip()[:400]}")
    stamp = tuple(int(x) for x in day[:10].split("-")) + (0, 0, 0)
    with zipfile.ZipFile(dest) as zin:  # pandoc stamps the parts it copies from its reference with the time
        parts = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for i, data in parts:
            zi = zipfile.ZipInfo(i.filename, date_time=stamp)
            zi.compress_type, zi.external_attr = i.compress_type, i.external_attr
            zout.writestr(zi, data)
    words = D.docx_text(dest)
    flat = D._norm_ws(words).replace("\n", " ")
    return words, [q for q in quotes if D._norm_ws(q) not in flat]


# ------------------------------------------------------------------ build and check
def playwright_dir(given: str | None) -> Path | None:
    for d in (given, os.environ.get("PLAYWRIGHT_DIR"), C.CONFIG.get("playwright_dir")):
        if d and (Path(d).expanduser() / "node_modules" / "playwright").is_dir():
            return Path(d).expanduser()
    return None


_PDF_DATE = re.compile(rb"/(CreationDate|ModDate) \(D:[^)]*\)")


def steady_pdf(data: bytes, day: str) -> bytes:
    """Chromium stamps the PDF with the time; the page's own date in its place (the same length,
    so the cross-reference offsets hold) makes an unchanged page print to the same bytes."""
    stamp = f"D:{day.replace('-', '')}000000Z".encode()

    def sub(m):
        old = m.group(0)
        new = b"/" + m.group(1) + b" (" + stamp + b")"
        return new.ljust(len(old)) if len(new) <= len(old) else old
    return _PDF_DATE.sub(sub, data)


def prepare(corpus: Path, slug: str, draft: bool) -> dict:
    """Everything a build checks and renders, or SystemExit with what is wrong."""
    rep = validate(corpus)
    if rep.get("errors"):
        raise SystemExit("the corpus does not validate; a one-pager is built only from one that does:\n  "
                         + "\n  ".join(f"{e.get('code')}: {str(e.get('detail'))[:120]}" for e in rep["errors"][:8]))
    cfg = load_cfg(corpus, slug)
    aud, proj = load_profiles(cfg)
    reg, idx, exts = load_reg(corpus), D.source_index(corpus), extractions(corpus)
    syn = current_page(corpus, slug)
    if not syn:
        raise SystemExit(f"no page yet for {slug}: run write, or add a page you wrote")
    o = syn["ext"]["upc_corpus"]["one_pager"]
    page = o["page"]
    stale = [x for x, snap in o["figures"].items() if x not in reg or reg[x].get("value") != snap.get("value")
             or reg[x].get("extraction_id") != snap.get("extraction_id")]
    if stale:
        raise SystemExit(f"figures changed since revision {o['revision']} was written: {', '.join(stale)}; "
                         "add the page again so its words match")
    problems = []

    def again(eid, quote, where):
        try:
            e = check_extraction(corpus, eid, exts)
        except SystemExit as err:
            problems.append(f"{where}: {err}")
            return
        if e["direct_quote"] != quote:
            problems.append(f"{where}: its quotation differs from the stored passage {eid}")
    figs = closure(page_figures(page), reg)
    for x in figs:
        if reg[x].get("extraction_id"):
            again(reg[x]["extraction_id"], reg[x]["quote"], x)
        for i, r in enumerate(reg[x].get("readings") or [], 1):
            again(r["extraction_id"], r["quote"], f"{x} reading {i}")
    if o.get("voice"):
        again(o["voice"]["extraction_id"], o["voice"]["quote"], "voice")
    if problems:
        raise SystemExit("refusing to build: passages not found again in their sources:\n  " + "\n  ".join(problems[:12]))
    J = judgements(corpus, reg)
    dist = cfg["distribution"]
    hard, final = {}, {}
    for x in figs:
        hard[x], final[x] = gate(x, reg, J, aud["id"], dist)
    voice_note, asset_notes = [], []
    consent = {(idx.get(c["source"]) or {}).get("source_id") for c in opt(proj).get("consent") or []}
    if o.get("voice") and o["voice"]["source_id"] not in consent:
        voice_note.append("The quotation in their words has no recorded consent for a page outside the team "
                          "(project one_pager.consent).")
    assets = {}
    wanted = [x for x in ((page.get("hero") or {}).get("image"), (page.get("how") or {}).get("image"),
                          (page.get("about") or {}).get("image"), (opt(proj).get("brand") or {}).get("logo")) if x]
    by_id = {a["id"]: a for a in opt(proj).get("assets") or []}
    for aid in dict.fromkeys(wanted):
        a = by_id.get(aid)
        if not a:
            raise SystemExit(f"the page uses asset {aid!r}, which is not in the project profile")
        p = asset_path(cfg, a)
        if not p.exists():
            raise SystemExit(f"asset {aid}: no file at {p}")
        if not a.get("may_publish") and a.get("kind") != "logo":
            asset_notes.append(f"The image {aid} is not marked may_publish in the project profile.")
        assets[aid] = {**a, "uri": data_uri(p)}
    blocking = [f"{x}: {'; '.join(h)}" for x, h in hard.items() if h]
    if not draft:
        blocking += [f"{x}: {'; '.join(w)}" for x, w in final.items() if w] + voice_note + asset_notes
    if blocking:
        raise SystemExit("refusing to build " + ("this draft" if draft else "a final page") + ":\n  " + "\n  ".join(blocking[:20])
                         + ("" if draft else "\n  (build with --draft for an internal draft that lists these)"))
    nums = number_problems(page, reg, proj)
    if nums:
        raise SystemExit("refusing to build: numbers that are not its figures:\n  " + "\n  ".join(nums[:12]))
    brand = opt(proj).get("brand") or {}
    if brand.get("font"):
        fp = Path(brand["font"]["file"]).expanduser()
        fp = fp if fp.is_absolute() else Path(cfg["project"]).parent / fp
        if not fp.exists():
            raise SystemExit(f"brand.font: no file at {fp}")
        family, font_uri = brand["font"]["family"], data_uri(fp)
    else:
        fp, family = ASSETS / "fonts" / "figtree.woff2", "Figtree"
        font_uri = data_uri(fp)
    shown = cfg.get("writer_labels") or {}
    real = o.get("writer") or ""
    writers = (f"Words by {shown.get(real, real)}" + (f", using {real}" if shown.get(real, real) != real else "") + ".") if real else ""
    stamp = (brand.get("draft_stamp") or DRAFT_STAMP) if draft or dist == "internal" else ""
    return {"cfg": cfg, "aud": aud, "proj": proj, "reg": reg, "idx": idx, "exts": exts, "J": J, "page": page, "syn": syn,
            "revision": o["revision"], "synthesis_id": syn["synthesis_id"], "voice": o.get("voice"), "draft": draft,
            "open": final if draft else {}, "voice_note": voice_note if draft else [], "asset_notes": asset_notes if draft else [],
            "assets": assets, "font_family": family, "font_uri": font_uri,
            "font_format": {".woff2": "woff2", ".woff": "woff", ".ttf": "truetype", ".otf": "opentype"}[fp.suffix.lower()],
            "writers": writers, "stamp": stamp, "changes": o.get("changes") or [], "figures": figs}


def text_checks(ctx: dict, texts: dict[str, str], corpus: Path) -> list[str]:
    aud, proj = ctx["aud"], ctx["proj"]
    plain = {k: re.sub(r"data:[a-z+/.-]+;base64,[A-Za-z0-9+/=]+", "data:", v) for k, v in texts.items()}
    hits = suppressed_in(plain, list(aud.get("never_publish") or []) + list(proj.get("never_publish") or []))
    leaks = D.local_paths_in(plain, corpus)
    alias = str(corpus).removeprefix("/private")  # macOS spells /var and /tmp both ways
    if alias != str(corpus):
        leaks += [(n, alias) for n, t in plain.items() if alias in t and (n, alias) not in leaks]
    return ([f"a never-publish phrase would appear: {n}: {w!r}" for n, w in hits]
            + [f"a local path would be published: {n}: {w!r}" for n, w in leaks])


def banned_in(ctx: dict, html: str = "") -> list[str]:
    """The audience's words to avoid that the page's own words use. The project's name and award lines in
    the footer, and a verbatim quotation, are not the writer's words, so they are not read."""
    seen = " ".join(t for _, t in page_texts(ctx["page"])).replace("\u00a0", " ").lower()
    return [w for w in opt(ctx["aud"]).get("banned_words") or [] if re.search(r"\b" + re.escape(w.lower()) + r"\b", seen)]


def build(corpus: Path, slug: str, out: Path, *, draft: bool = False, pdf: bool = True, docx: bool = True,
          pw: str | None = None) -> dict:
    corpus, out = Path(corpus).resolve(), Path(out).expanduser().resolve()
    before = (C.load_json(D.manifest_path(out), {}) or {}).get("files") or {}
    if out.exists() and any(out.iterdir()) and not before:
        raise SystemExit(f"refusing to build into {out}: it holds files no one-pager build made; give an empty "
                         "folder, or the one an earlier build of this page wrote")
    out.mkdir(parents=True, exist_ok=True)
    ctx = prepare(corpus, slug, draft)
    html = page_html(ctx)
    md = figures_markdown(ctx)
    files = {"index.html": html.encode(), f"{slug}-figures.md": md.encode()}
    texts = {"index.html": html, f"{slug}-figures.md": md}
    render = None
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        if docx:
            quotes = [ctx["reg"][x]["quote"] for x in ctx["figures"] if ctx["reg"][x].get("quote")]
            words, missing = make_docx(figures_markdown(ctx, D.md), quotes, tmp / f"{slug}-figures.docx", tmp, str(ctx["cfg"]["date"]))
            if missing:
                raise SystemExit("refusing to build: passages missing from the Word file:\n  " + "\n  ".join(repr(m[:90]) for m in missing[:8]))
            texts[f"{slug}-figures.docx"] = words
            files[f"{slug}-figures.docx"] = (tmp / f"{slug}-figures.docx").read_bytes()
        bad = text_checks(ctx, texts, corpus)
        if bad:
            raise SystemExit("refusing to build:\n  " + "\n  ".join(bad[:12]))
        if pdf:
            d = playwright_dir(pw)
            if not d:
                raise C.Unavailable("no Playwright to print the PDF: pass --playwright-dir <a folder with "
                                    "node_modules/playwright> (or set PLAYWRIGHT_DIR), or build with --no-pdf")
            C.require_tools("node")
            (tmp / "index.html").write_text(html, encoding="utf-8")
            p = subprocess.run(["node", str(RENDER), str(tmp / "index.html"), str(tmp / "page.pdf"), str(d)],
                               capture_output=True, text=True, timeout=300)
            if p.returncode != 0:
                raise C.Unavailable(f"the browser could not print the page: {(p.stderr or p.stdout).strip()[:300]}")
            render = json.loads(p.stdout.strip().splitlines()[-1])
            probs = []
            if render["pages"] != 1:
                probs.append(f"it prints on {render['pages']} pages, not one: {render['overflow_px']} px too tall"
                             + (f"; these blocks run past the first sheet: {', '.join(render['overflowing'])}" if render["overflowing"] else ""))
            if render["phone_overflow_px"]:
                probs.append(f"on a phone (390 px) it scrolls sideways by {render['phone_overflow_px']} px")
            if render["console_errors"]:
                probs.append("the browser reported errors: " + "; ".join(render["console_errors"][:3]))
            if probs:
                raise SystemExit("refusing to build:\n  " + "\n  ".join(probs))
            files[f"{slug}.pdf"] = steady_pdf((tmp / "page.pdf").read_bytes(), str(ctx["cfg"]["date"])[:10])
    for name, data in files.items():
        (out / name).write_bytes(data)
    for name in before:  # what the last build made that this one did not; nothing else is touched
        if name not in files and (out / name).is_file():
            (out / name).unlink()
    warn = banned_in(ctx, html)
    record = {"built": now(), **TOOL, "corpus": str(corpus), "slug": slug, "revision": ctx["revision"],
              "synthesis_id": ctx["synthesis_id"], "draft": draft, "distribution": ctx["cfg"]["distribution"],
              "figures": ctx["figures"], "codings": ctx["syn"]["ext"]["upc_corpus"]["one_pager"].get("codings") or [],
              "audience": {"path": ctx["cfg"]["audience"], "digest": digest(ctx["aud"])},
              "project": {"path": ctx["cfg"]["project"], "digest": digest(ctx["proj"])},
              "files": {n: hashlib.sha256(b).hexdigest() for n, b in sorted(files.items())},
              "open": {x: w for x, w in ctx["open"].items() if w}, "banned_words": warn, "render": render}
    C.save_json(D.manifest_path(out), record)
    C.upc("event", "--corpus", str(corpus), check=False, stdin=json.dumps({
        "activity_type": "export", "tool": TOOL["tool"], "started_at": now(), "status": "success",
        "inputs": {"synthesis_ids": [ctx["synthesis_id"]]},
        "notes": f"one-pager {slug} r{ctx['revision']}{' (draft)' if draft else ''}: {len(ctx['figures'])} figures -> {out.name}"}))
    return record


def cmd_check(corpus: Path, slug: str, draft: bool) -> int:
    ctx = prepare(Path(corpus).resolve(), slug, draft)
    html, md = page_html(ctx), figures_markdown(ctx)
    bad = text_checks(ctx, {"index.html": html, f"{slug}-figures.md": md}, Path(corpus).resolve())
    open_ = {x: w for x, w in ctx["open"].items() if w}
    C.counts_header("One-pager: check", [("page", f"{slug} r{ctx['revision']} ({ctx['synthesis_id']})"),
                                         ("figures", len(ctx["figures"])), ("passages found again", "all"),
                                         ("open for a final page", len(open_)), ("text problems", len(bad)),
                                         ("words to avoid", ", ".join(banned_in(ctx, html)) or "none")])
    print_open(open_)
    for b in bad:
        print(f"  PROBLEM {b}")
    return 1 if bad else 0


# ------------------------------------------------------------------ main
def cmd_figures_list(corpus: Path, slug: str | None) -> int:
    reg = load_reg(corpus)
    J = judgements(corpus, reg)
    aud = load_profiles(load_cfg(corpus, slug))[0] if slug else None
    C.counts_header("Figures", [("in the register", len(reg)), ("audience", aud["id"] if aud else "— (pass --slug)")])
    for x in sorted(reg):
        f = reg[x]
        s = status_of(x, reg, J, aud["id"] if aud else "")
        print(f"  {x:32} {f['value'][:18]:18} class {s['class'] or '—':10} check {s['check'] or '—':9} "
              f"release {s['clearance'] or '—':13}" + (f" {aud['id']}: {s['rel'] or '—'}" if aud else ""))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("init", help="write or update a page's config")
    p.add_argument("--corpus", required=True)
    p.add_argument("--slug", required=True)
    for f in ("audience", "project", "title", "date"):
        p.add_argument(f"--{f}")
    p.add_argument("--layout", choices=LAYOUTS)
    p.add_argument("--distribution", choices=DISTRIBUTIONS)
    p.add_argument("--show-writer-as", action="append", metavar="WRITER=LABEL")
    p = sub.add_parser("figures", help="the figure register")
    fs = p.add_subparsers(dest="fcmd", required=True)
    q = fs.add_parser("add")
    q.add_argument("--corpus", required=True)
    q.add_argument("--from", dest="src", required=True)
    q.add_argument("--by", default="a person", help="who entered them, for the passages' provenance")
    q.add_argument("--replace", action="store_true")
    q.add_argument("--plan", action="store_true")
    q = fs.add_parser("import")
    q.add_argument("--corpus", required=True)
    q.add_argument("--from-dossier", action="store_true", required=True)
    q.add_argument("--plan", action="store_true")
    q = fs.add_parser("list")
    q.add_argument("--corpus", required=True)
    q.add_argument("--slug")
    p = sub.add_parser("code", help="the local model judges each figure for the page's audience")
    p.add_argument("--corpus", required=True)
    p.add_argument("--slug", required=True)
    p.add_argument("--figure", action="append")
    p.add_argument("--redo", action="store_true")
    p.add_argument("--plan", action="store_true")
    p = sub.add_parser("review", help="a sheet of every figure's answers; apply a person's")
    p.add_argument("--corpus", required=True)
    p.add_argument("--slug", required=True)
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--out")
    g.add_argument("--apply")
    p.add_argument("--coder")
    p = sub.add_parser("reply", help="save a reviewer's reply as a source, for signoff --ref")
    p.add_argument("--corpus", required=True)
    p.add_argument("--from", dest="src", required=True)
    p.add_argument("--by", required=True)
    p.add_argument("--title")
    p.add_argument("--date")
    p.add_argument("--plan", action="store_true")
    p = sub.add_parser("signoff", help="record a reviewer's check, or a release")
    p.add_argument("--corpus", required=True)
    p.add_argument("--figure", action="append", required=True)
    p.add_argument("--by", required=True)
    p.add_argument("--check", choices=CODES["check"])
    p.add_argument("--clearance", choices=CODES["clearance"])
    p.add_argument("--ref")
    p.add_argument("--note")
    p = sub.add_parser("write", help="the local model drafts the page from its eligible figures")
    p.add_argument("--corpus", required=True)
    p.add_argument("--slug", required=True)
    p.add_argument("--plan", action="store_true")
    p = sub.add_parser("add", help="store a page written anywhere")
    p.add_argument("--corpus", required=True)
    p.add_argument("--slug", required=True)
    p.add_argument("--from", dest="src", required=True)
    p.add_argument("--writer", required=True)
    p.add_argument("--model")
    p.add_argument("--method", default="manual")
    p.add_argument("--plan", action="store_true")
    p = sub.add_parser("build", help="the page (HTML and PDF) and its figures note")
    p.add_argument("--corpus", required=True)
    p.add_argument("--slug", required=True)
    p.add_argument("--out", required=True)
    p.add_argument("--draft", action="store_true")
    p.add_argument("--no-pdf", action="store_true")
    p.add_argument("--no-docx", action="store_true")
    p.add_argument("--playwright-dir")
    p = sub.add_parser("check", help="every check a build makes, without the browser; write nothing")
    p.add_argument("--corpus", required=True)
    p.add_argument("--slug", required=True)
    p.add_argument("--draft", action="store_true")
    a = ap.parse_args()
    corpus = Path(a.corpus).expanduser().resolve()
    if a.cmd == "init":
        rc, nxt = cmd_init(corpus, a), "figures add, then code or review"
    elif a.cmd == "figures":
        if a.fcmd == "add":
            items = C.load_json(a.src, None)
            items = items.get("figures") if isinstance(items, dict) else items
            if not isinstance(items, list):
                raise SystemExit(f"{a.src}: a JSON list of figures (or {{\"figures\": [...]}})")
            rc, nxt = cmd_figures_add(corpus, items, by=a.by, replace=a.replace, plan=a.plan), "code or review"
        elif a.fcmd == "import":
            rc, nxt = cmd_figures_import(corpus, a.plan), "code or review"
        else:
            rc, nxt = cmd_figures_list(corpus, a.slug), None
    elif a.cmd == "code":
        rc, nxt = cmd_code(corpus, a.slug, a.figure, a.redo, a.plan), "review --out"
    elif a.cmd == "review":
        if a.out:
            rc, nxt = cmd_review_out(corpus, a.slug, Path(a.out)), "edit the sheet, then review --apply"
        else:
            if not a.coder:
                raise SystemExit("review --apply needs --coder (whose answers these are)")
            rc, nxt = cmd_review_apply(corpus, a.slug, Path(a.apply), a.coder), "signoff, then write or add"
    elif a.cmd == "reply":
        sid = cmd_reply(corpus, Path(a.src), a.by, a.title, a.date, a.plan)
        rc, nxt = 0, f"signoff --figure <id> --by {a.by} --check confirmed --ref {sid}"
    elif a.cmd == "signoff":
        rc, nxt = cmd_signoff(corpus, a.figure, a.by, a.check, a.clearance, a.ref, a.note), "build"
    elif a.cmd == "write":
        rc, nxt = cmd_write(corpus, a.slug, a.plan), "build --draft"
    elif a.cmd == "add":
        page = C.load_json(a.src, None)
        rep = add_page(corpus, a.slug, page if isinstance(page, dict) else {}, writer=a.writer, model=a.model,
                       method=a.method, plan=a.plan)
        print_page_report(rep)
        rc, nxt = 0, "build --draft" if rep.get("open") else "build"
    elif a.cmd == "build":
        t0 = time.time()
        r = build(corpus, a.slug, Path(a.out), draft=a.draft, pdf=not a.no_pdf, docx=not a.no_docx, pw=a.playwright_dir)
        C.counts_header("One-pager", [("page", f"{a.slug} r{r['revision']} ({r['synthesis_id']})"),
                                      ("kind", "internal draft" if r["draft"] else f"final, {r['distribution']}"),
                                      ("figures", len(r["figures"])), ("files", ", ".join(r["files"])),
                                      ("open for a final page", len(r["open"])),
                                      ("words to avoid found", ", ".join(r["banned_words"]) or "none"),
                                      ("out", Path(a.out).expanduser()), ("seconds", round(time.time() - t0, 1))])
        print_open(r["open"])
        rc, nxt = 0, "look at the page and its figures note"
    else:
        rc, nxt = cmd_check(corpus, a.slug, a.draft), "build"
    if rc == 0 and nxt:
        print(json.dumps({"ok": True, "next": nxt}))
    return rc


if __name__ == "__main__":
    try:
        sys.exit(main())
    except C.Unavailable as e:
        C.die(f"UNAVAILABLE: {e}")
