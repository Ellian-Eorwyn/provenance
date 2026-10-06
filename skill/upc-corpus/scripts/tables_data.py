"""Data layer of the tables route: loading, profiling, links, codebooks, unit
specs, and the rendering of one unit (a row, a row with its linked rows, or a
line of a row) into the text a model reads.

No model is called here and nothing here writes outside a run folder. Three
rules hold throughout:

  * **Whitelist rendering.** A unit's text contains only the columns its spec
    names. A blind part is blind because nothing else can reach the prompt.
  * **Evidence by sentence id.** The author's own words are split into
    sentences with ids (M1.3, S2.1, E1.4 …) that map back to exact characters of
    a cell. A model cites ids; it never retypes a quotation.
  * **Other people's words are marked.** Quote blocks inside a cell (a forum
    reply quoting an earlier post) are not the author's words: they get no ids,
    and when the quoted row is already in the case they collapse to a marker.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import unicodedata
from bisect import bisect_right
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

csv.field_size_limit(10**9)

TABLE_SUFFIXES = {".csv", ".tsv", ".jsonl", ".ndjson", ".xlsx", ".xlsm"}
LIST_SEPARATOR = ";"
LINK_MIN_CONTAINMENT = 0.9
PARTIAL_LINK_CONTAINMENT = 0.5
LINK_SAMPLE = 20000
SAMPLES_PER_COLUMN = 3
SAMPLE_CHARS = 80
ROLES = ["id", "link", "text", "short_text", "speaker", "datetime", "url", "category",
         "number", "flag", "derived", "code_slot", "other"]


# --- small helpers ----------------------------------------------------------

def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def canonical_json(obj) -> str:
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def slug_token(value: str) -> str:
    """A UPC code/slug token: lowercase [a-z0-9-] (codes appear in site URLs)."""
    token = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode().lower()
    token = re.sub(r"[^a-z0-9]+", "-", token).strip("-")
    return token or "x"


def list_values(value: str) -> list[str]:
    return [v.strip() for v in str(value or "").split(LIST_SEPARATOR) if v.strip()]


def percentile(sorted_values: list[int], p: float) -> int:
    if not sorted_values:
        return 0
    k = min(len(sorted_values) - 1, max(0, int(round(p * (len(sorted_values) - 1)))))
    return sorted_values[k]


# --- loading ------------------------------------------------------------------

@dataclass
class Table:
    name: str
    path: Path
    columns: list[str]
    rows: list[dict]
    fmt: str
    sha256: str = ""
    sheet: str | None = None
    _index: dict = field(default_factory=dict, repr=False)

    def index(self, column: str, *, as_list: bool = False) -> dict[str, list[int]]:
        """value -> row positions (list-aware when as_list). Cached."""
        key = (column, as_list)
        if key not in self._index:
            idx: dict[str, list[int]] = defaultdict(list)
            for n, row in enumerate(self.rows):
                value = row.get(column, "")
                for v in (list_values(value) if as_list else ([value.strip()] if value and value.strip() else [])):
                    idx[v].append(n)
            self._index[key] = idx
        return self._index[key]


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, (list, tuple)):
        if all(not isinstance(v, (dict, list)) for v in value):
            return LIST_SEPARATOR.join(_cell(v) for v in value)
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def load_table(path: str | Path, name: str | None = None) -> list[Table]:
    """Load one file as one table (or one table per sheet for xlsx)."""
    path = Path(path)
    suffix = path.suffix.lower()
    stem = name or path.stem
    digest = sha256_file(path)
    if suffix in (".csv", ".tsv"):
        with open(path, encoding="utf-8-sig", newline="") as fh:
            head = fh.read(65536)
            fh.seek(0)
            delimiter = "\t" if suffix == ".tsv" else ","
            if suffix == ".csv":
                try:
                    delimiter = csv.Sniffer().sniff(head, delimiters=",\t;|").delimiter
                except csv.Error:
                    delimiter = ","
            reader = csv.DictReader(fh, delimiter=delimiter)
            columns = list(reader.fieldnames or [])
            rows = [{c: (r.get(c) if r.get(c) is not None else "") for c in columns} for r in reader]
        return [Table(stem, path, columns, rows, suffix[1:], digest)]
    if suffix in (".jsonl", ".ndjson"):
        columns: list[str] = []
        seen = set()
        raw = []
        with open(path, encoding="utf-8-sig") as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                obj = json.loads(line)
                if not isinstance(obj, dict):
                    continue
                for k in obj:
                    if k not in seen:
                        seen.add(k)
                        columns.append(k)
                raw.append(obj)
        rows = [{c: _cell(obj.get(c)) for c in columns} for obj in raw]
        return [Table(stem, path, columns, rows, "jsonl", digest)]
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl  # noqa: PLC0415
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        tables = []
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if not header:
                continue
            columns = [str(h).strip() if h is not None else f"column_{i + 1}" for i, h in enumerate(header)]
            rows = []
            for values in it:
                if values is None or all(v is None for v in values):
                    continue
                rows.append({c: _cell(values[i] if i < len(values) else None) for i, c in enumerate(columns)})
            sheet_name = stem if len(wb.worksheets) == 1 else f"{stem}.{slug_token(ws.title)}"
            tables.append(Table(sheet_name, path, columns, rows, "xlsx", digest, sheet=ws.title))
        return tables
    raise ValueError(f"unsupported table format: {path.name}")


def discover_inputs(inputs: list[str | Path]) -> list[Path]:
    """Files given, plus the table files directly inside folders given."""
    found = []
    for item in inputs:
        p = Path(item).expanduser()
        if p.is_dir():
            found += sorted(q for q in p.iterdir() if q.is_file() and q.suffix.lower() in TABLE_SUFFIXES
                            and not q.name.startswith("."))
        elif p.is_file():
            found.append(p)
        else:
            raise FileNotFoundError(str(p))
    return found


def unique_names(paths: list[Path]) -> dict[Path, str]:
    """A short, unique table name per file (its stem, or parent_stem on a clash)."""
    counts = Counter(p.stem for p in paths)
    return {p: (p.stem if counts[p.stem] == 1 else f"{slug_token(p.parent.name)}_{p.stem}") for p in paths}


# --- profiling ----------------------------------------------------------------

_INT = re.compile(r"^-?\d+$")
_FLOAT = re.compile(r"^-?\d+(\.\d+)?([eE]-?\d+)?$")
_BOOL = {"true", "false", "yes", "no", "0", "1", "y", "n", "t", "f"}
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")
_URL = re.compile(r"^https?://", re.I)
_TOKEN = re.compile(r"^[A-Za-z]{0,3}[\w.-]{1,40}$")
QUOTE_MARKERS = {
    "xenforo": re.compile(r"(?m)^.{1,80} said:\s*$|Click to expand\.\.\."),
    "markdown": re.compile(r"(?m)^\s*>"),
    "email": re.compile(r"(?m)^On .{5,120} wrote:\s*$"),
}


def column_facts(table: Table, column: str) -> dict:
    values = [r.get(column, "") for r in table.rows]
    filled = [v for v in values if v.strip()]
    lengths = sorted(len(v) for v in filled)
    distinct = len(set(filled))

    def share(pred) -> float:
        return round(sum(1 for v in filled if pred(v)) / len(filled), 3) if filled else 0.0

    listy = share(lambda v: LIST_SEPARATOR in v and all(_TOKEN.match(t) for t in list_values(v)))
    facts = {
        "rows": len(values), "filled": len(filled), "fill": round(len(filled) / len(values), 3) if values else 0.0,
        "distinct": distinct, "unique": bool(filled) and distinct == len(filled) and len(filled) == len(values),
        "len_p50": percentile(lengths, 0.5), "len_p90": percentile(lengths, 0.9),
        "len_max": lengths[-1] if lengths else 0,
        "newlines": share(lambda v: "\n" in v),
        "int": share(lambda v: bool(_INT.match(v.strip()))),
        "number": share(lambda v: bool(_FLOAT.match(v.strip()))),
        "bool": share(lambda v: v.strip().lower() in _BOOL),
        "date": share(lambda v: bool(_DATE.match(v.strip()))),
        "url": share(lambda v: bool(_URL.match(v.strip()))),
        "list": listy,
        "samples": [],
        "quote_blocks": {},
    }
    seen = set()
    for v in filled:
        s = v if len(v) <= SAMPLE_CHARS else v[:SAMPLE_CHARS] + "…"
        s = s.replace("\n", " / ")
        if s not in seen:
            seen.add(s)
            facts["samples"].append(s)
        if len(facts["samples"]) >= SAMPLES_PER_COLUMN:
            break
    if facts["len_p90"] >= 40:
        for style, rx in QUOTE_MARKERS.items():
            hits = sum(1 for v in filled if rx.search(v))
            if hits and hits / max(1, len(filled)) >= 0.02:
                facts["quote_blocks"][style] = hits
    return facts


def is_key(facts: dict) -> bool:
    return (facts["unique"] and facts["len_max"] <= 80 and facts["newlines"] == 0
            and facts["date"] < 0.9 and facts["rows"] >= 3)


def _strip_prefix(value: str) -> tuple[str, str]:
    m = re.match(r"^([A-Za-z]{1,3})(\d.*)$", value)
    return (m.group(1), m.group(2)) if m else ("", value)


def find_links(tables: dict[str, Table], facts: dict[str, dict[str, dict]]) -> list[dict]:
    """Links by value containment: column values (list-aware) found among another
    column's unique values. Tolerates a short letter prefix on either side
    (e.g. action ids `a22149108` vs post ids `22149108`)."""
    keys = {(t, c) for t in tables for c, f in facts[t].items() if is_key(f)}
    key_values: dict[tuple[str, str], set[str]] = {}
    key_stripped: dict[tuple[str, str], dict[str, str]] = {}
    for t, c in keys:
        vals = {r[c].strip() for r in tables[t].rows if r.get(c, "").strip()}
        key_values[(t, c)] = vals
        stripped = {}
        for v in vals:
            prefix, rest = _strip_prefix(v)
            stripped.setdefault(rest, prefix)
        key_stripped[(t, c)] = stripped
    links = []
    for t, table in tables.items():
        for c, f in facts[t].items():
            if f["filled"] == 0 or f["len_max"] > 400:
                continue
            as_list = f["list"] > 0
            tokens: list[str] = []
            for r in table.rows:
                v = r.get(c, "")
                tokens += list_values(v) if as_list else ([v.strip()] if v.strip() else [])
                if len(tokens) >= LINK_SAMPLE:
                    break
            if not tokens or not all(_TOKEN.match(tok) for tok in tokens[:200]):
                continue
            distinct_tokens = set(tokens)
            if len(distinct_tokens) < min(5, len(table.rows)):
                continue
            for (kt, kc) in keys:
                if (kt, kc) == (t, c):
                    continue
                kv = key_values[(kt, kc)]
                hit = sum(1 for tok in distinct_tokens if tok in kv)
                transform = None
                if hit / len(distinct_tokens) < LINK_MIN_CONTAINMENT:
                    # try a prefix difference on either side
                    stripped = key_stripped[(kt, kc)]
                    prefixes = Counter()
                    alt = 0
                    for tok in distinct_tokens:
                        p_tok, rest = _strip_prefix(tok)
                        if rest in stripped:
                            alt += 1
                            prefixes[(p_tok, stripped[rest])] += 1
                    if prefixes and alt > hit:
                        (from_prefix, to_prefix), _ = prefixes.most_common(1)[0]
                        transform = {"strip": from_prefix, "add": to_prefix}
                        hit = alt
                containment = hit / len(distinct_tokens)
                if containment >= PARTIAL_LINK_CONTAINMENT:
                    links.append({"from": f"{t}.{c}", "to": f"{kt}.{kc}", "list": as_list,
                                  "transform": transform, "containment": round(containment, 3),
                                  "partial": containment < LINK_MIN_CONTAINMENT,
                                  "distinct_values": len(distinct_tokens)})
    links.sort(key=lambda link: (link["partial"], -link["containment"], link["from"]))
    return links


def apply_transform(value: str, transform: dict | None) -> str:
    if not transform:
        return value
    prefix, rest = _strip_prefix(value)
    if transform.get("strip") and prefix == transform["strip"]:
        value = rest
    elif transform.get("strip") and prefix and prefix != transform["strip"]:
        return value
    return f"{transform.get('add', '')}{value}" if transform.get("add") and not value.startswith(transform["add"]) else value


def profile_tables(tables: dict[str, Table]) -> dict:
    facts = {t: {c: column_facts(tb, c) for c in tb.columns} for t, tb in tables.items()}
    links = find_links(tables, facts)
    out = {"tables": {}, "links": links}
    for t, tb in tables.items():
        out["tables"][t] = {
            "path": str(tb.path), "format": tb.fmt, "sheet": tb.sheet, "sha256": tb.sha256,
            "rows": len(tb.rows), "keys": [c for c, f in facts[t].items() if is_key(f)],
            "columns": {c: facts[t][c] for c in tb.columns},
        }
    return out


def check_roles(table_profile: dict, roles: dict[str, str], links: list[dict], table: str) -> list[str]:
    """Contradictions between a model's column roles and the measured facts."""
    warnings = []
    linked_from = {link["from"].split(".", 1)[1] for link in links if link["from"].startswith(table + ".")}
    for column, role in roles.items():
        f = table_profile["columns"].get(column)
        if f is None:
            warnings.append(f"{column}: labelled {role} but there is no such column")
            continue
        if role == "id" and not f["unique"]:
            warnings.append(f"{column}: labelled id but its values are not unique")
        if role == "text" and f["len_p50"] < 20:
            warnings.append(f"{column}: labelled quotable text but its median length is {f['len_p50']}")
        if role == "code_slot" and f["fill"] > 0.05:
            warnings.append(f"{column}: labelled an empty coding slot but {f['fill']:.0%} of rows are filled")
        if role == "link" and column not in linked_from:
            warnings.append(f"{column}: labelled a link but no link from it was measured")
        if role == "datetime" and f["date"] < 0.8 and f["filled"]:
            warnings.append(f"{column}: labelled a date but only {f['date']:.0%} of values look like dates")
    return warnings


# --- codebooks ----------------------------------------------------------------

FIELD_TYPES = {"single", "multi", "int", "text", "multi_with_quotes", "multi_with_polarity"}
DEFAULT_POLARITY = ["enforced", "violated", "contested"]


@dataclass
class Codebook:
    version: str
    fields: dict[str, dict]
    parts: dict[str, dict]
    sha256: str
    path: Path


def load_codebook(path: str | Path) -> Codebook:
    """A field codebook ({version, fields, parts}) or upc-corpus's list
    format ([{slug, question, multi_label, codes[...]}]); normalised to fields."""
    path = Path(path)
    raw = json.loads(path.read_text(encoding="utf-8"))
    digest = sha256_file(path)
    if isinstance(raw, dict) and isinstance(raw.get("fields"), dict):
        fields = {}
        for name, f in raw["fields"].items():
            ftype = f.get("type", "single")
            if ftype not in FIELD_TYPES:
                raise ValueError(f"field {name}: unknown type {ftype!r}")
            values = f.get("values")
            fields[name] = {"type": ftype, "def": f.get("def", ""),
                            "values": dict(values) if isinstance(values, dict) else None,
                            "lookup": values if isinstance(values, str) else None}
        parts = raw.get("parts") or {"main": {"unit": raw.get("unit", ""), "fields": list(fields)}}
        for p, spec in parts.items():
            missing = [x for x in spec.get("fields", []) if x not in fields]
            if missing:
                raise ValueError(f"part {p}: unknown fields {missing}")
        return Codebook(str(raw.get("version", "")), fields, parts, digest, path)
    if isinstance(raw, list):
        fields = {}
        for b in raw:
            name = b["slug"]
            fields[name] = {"type": "multi" if b.get("multi_label") else "single",
                            "def": b.get("question") or b.get("title") or "",
                            "values": {c["code"]: c.get("definition") or c.get("label") or "" for c in b.get("codes", [])},
                            "lookup": None}
        return Codebook("", fields, {"main": {"unit": "", "fields": list(fields)}}, digest, path)
    raise ValueError("codebook must be a field codebook object or a list of codebooks")


# --- unit specs -----------------------------------------------------------------

UNITS_FORMAT = "upc-corpus-units/1"


class SpecError(ValueError):
    pass


def resolve_path(tables: dict[str, Table], start_rows: list[tuple[str, dict]], hops: list[dict]) -> list[tuple[str, dict]]:
    """Follow link hops from rows. A hop is {"from": "t.col", "to": "u.col"}; values
    in `from` may be lists; `transform` maps a value's prefix. Order is kept and
    duplicates dropped."""
    current = start_rows
    for hop in hops:
        ft, fc = hop["from"].split(".", 1)
        tt, tc = hop["to"].split(".", 1)
        target = tables[tt]
        idx = target.index(tc, as_list=bool(hop.get("to_list")))
        nxt, seen = [], set()
        for t, row in current:
            if t != ft:
                raise SpecError(f"hop {hop['from']} applied to a row of {t}")
            raw = row.get(fc, "")
            values = list_values(raw) if hop.get("list") else ([raw.strip()] if raw.strip() else [])
            for v in values:
                v = apply_transform(v, hop.get("transform"))
                for n in idx.get(v, []):
                    if (tt, n) not in seen:
                        seen.add((tt, n))
                        nxt.append((tt, target.rows[n]))
        current = nxt
    return current


def _cond_ok(tables: dict[str, Table], table: str, row: dict, cond: dict, cache: dict) -> bool:
    column = cond.get("column")
    value = (row.get(column, "") if column else "").strip()
    if "not_empty" in cond:
        return bool(value) == bool(cond["not_empty"])
    if "equals" in cond:
        return value == str(cond["equals"])
    if "in" in cond or "not_in" in cond:
        ref = cond.get("in") or cond.get("not_in")
        key = canonical_json(ref)
        if key not in cache:
            if "table" in ref:
                t2, c2 = ref["table"], ref["column"]
                vals = set()
                for r in tables[t2].rows:
                    for v in (list_values(r.get(c2, "")) if ref.get("list") else [r.get(c2, "").strip()]):
                        if v:
                            vals.add(apply_transform(v, ref.get("transform")))
                cache[key] = vals
            elif "values" in ref:
                cache[key] = {str(v) for v in ref["values"]}
            else:
                raise SpecError(f"condition needs a table or values: {ref}")
        hit = value in cache[key]
        return hit if "in" in cond else not hit
    raise SpecError(f"unknown condition {cond}")


def unit_rows(tables: dict[str, Table], part: dict, screened: dict[str, dict] | None = None) -> list[tuple[str, dict]]:
    """The base rows of a part: a table or a union of tables (a later table wins
    for the same id), filtered by `where`, optionally by screening routes."""
    names = part["table"] if isinstance(part["table"], list) else [part["table"]]
    id_col = part["id"]
    by_id: dict[str, tuple[str, dict]] = {}
    for name in names:
        if name not in tables:
            raise SpecError(f"unknown table {name}")
        for row in tables[name].rows:
            rid = row.get(id_col, "").strip()
            if rid:
                by_id[rid] = (name, row)
    cache: dict = {}
    out = []
    for rid, (name, row) in by_id.items():
        if all(_cond_ok(tables, name, row, cond, cache) for cond in part.get("where", [])):
            out.append((name, row))
    route = part.get("screened")
    if route:
        screened = screened or {}
        keep = []
        for name, row in out:
            s = screened.get(row[id_col].strip())
            if s and s.get("r", 0) >= route.get("min_relevance", 2) and route["part"] in s.get("parts", []):
                keep.append((name, row))
        out = keep
    return out


# --- quote blocks and sentences --------------------------------------------------

@dataclass
class Segment:
    start: int
    end: int
    kind: str          # "own" or "quote"
    author: str = ""   # for quotes: whose words (may be "")
    source: str = ""   # for quotes: the quoted row id when known


_XF_SAID = re.compile(r"(?m)^(?P<name>[^\n]{1,80}?) said:[ \t]*$")
_XF_END = "Click to expand..."


def split_quote_blocks(text: str, style: str | None, sources: list[str] | None = None) -> list[Segment]:
    """Split a cell into own words and quoted blocks. `sources` (e.g. a post's
    quoted_post_ids) are paired with the blocks in order when their counts match."""
    if not text:
        return []
    if style == "xenforo":
        segs, pos = [], 0
        blocks = []
        search = 0
        while True:
            end = text.find(_XF_END, search)
            if end < 0:
                break
            end_full = end + len(_XF_END)
            said = None
            for m in _XF_SAID.finditer(text, pos, end):
                said = m
            start = said.start() if said else pos
            blocks.append((start, end_full, said.group("name").strip() if said else ""))
            pos = end_full
            search = end_full
        if sources and len(sources) != len(blocks):
            sources = None
        for i, (start, end_full, name) in enumerate(blocks):
            if start > (segs[-1].end if segs else 0):
                segs.append(Segment(segs[-1].end if segs else 0, start, "own"))
            segs.append(Segment(start, end_full, "quote", name, sources[i] if sources else ""))
        last = segs[-1].end if segs else 0
        if last < len(text):
            segs.append(Segment(last, len(text), "own"))
        return _merge_own(segs)
    if style == "markdown":
        segs = []
        for m in re.finditer(r"[^\n]*(\n|$)", text):
            if m.start() == m.end():
                continue
            kind = "quote" if re.match(r"\s*>", m.group()) else "own"
            if segs and segs[-1].kind == kind:
                segs[-1].end = m.end()
            else:
                segs.append(Segment(m.start(), m.end(), kind))
        return segs
    if style == "email":
        m = re.search(r"(?m)^On .{5,120} wrote:\s*$", text)
        if m:
            return _merge_own([Segment(0, m.start(), "own"), Segment(m.start(), len(text), "quote")])
    return [Segment(0, len(text), "own")]


def _merge_own(segs: list[Segment]) -> list[Segment]:
    out = []
    for s in segs:
        if s.end <= s.start:
            continue
        if out and s.kind == "own" and out[-1].kind == "own":
            out[-1].end = s.end
        else:
            out.append(s)
    return out


def quote_body(text: str, seg: Segment) -> str:
    """The quoted words of a xenforo block, without its 'X said:' line and the
    'Click to expand...' marker."""
    body = text[seg.start:seg.end]
    body = re.sub(r"^[^\n]{1,80}? said:[ \t]*\n", "", body, count=1)
    if body.endswith(_XF_END):
        body = body[: -len(_XF_END)]
    return body.strip()


_SENT_END = re.compile(r"(?<=[.!?…])[\"'”’)\]]*\s+(?=[\"'“‘(\[]?[A-Z0-9])")


def sentences(text: str, start: int, end: int, *, max_chars: int = 600) -> list[tuple[int, int]]:
    """Sentence spans within text[start:end]: line breaks end a sentence (forum
    posts are full of lists), as do . ! ? before a capital. Leading/trailing
    whitespace trimmed; empties dropped; very long runs split at ; or , ."""
    spans = []
    for m in re.finditer(r"[^\n]+", text[start:end]):
        line_start = start + m.start()
        line = m.group()
        cuts = [0] + [x.end() for x in _SENT_END.finditer(line)] + [len(line)]
        for a, b in zip(cuts, cuts[1:]):
            while b - a > max_chars:
                window = line[a:a + max_chars]
                cut = max(window.rfind("; "), window.rfind(", "))
                cut = cut + 2 if cut > max_chars // 3 else max_chars
                spans.append((line_start + a, line_start + a + cut))
                a += cut
            spans.append((line_start + a, line_start + b))
    out = []
    for a, b in spans:
        while a < b and text[a].isspace():
            a += 1
        while b > a and text[b - 1].isspace():
            b -= 1
        if b - a >= 2 and re.search(r"\w", text[a:b]):
            out.append((a, b))
    return out


# --- rendering one unit -------------------------------------------------------

@dataclass
class Evidence:
    eid: str
    table: str
    row: str
    column: str
    start: int
    end: int
    speaker: str
    code: str
    text: str


@dataclass
class Rendered:
    unit_id: str
    part: str
    text: str                                  # what the model reads (with [ids])
    clean: str                                 # the same without ids (a reading copy)
    evidence: dict[str, Evidence]              # id -> cell span
    clean_spans: dict[str, tuple[int, int]]    # id -> span in `clean`
    blocks: list[dict]                         # [{code, table, row, label}]
    speakers: dict[str, list[str]]             # speaker code -> ids
    header: dict                               # unit metadata
    hidden_hits: list[dict] = field(default_factory=list)
    sha256: str = ""


class _Builder:
    def __init__(self, show_ids: bool):
        self.parts: list[str] = []
        self.clean: list[str] = []
        self.len_clean = 0
        self.show_ids = show_ids
        self.clean_spans: dict[str, tuple[int, int]] = {}

    def add(self, text: str, eid: str | None = None):
        if eid and self.show_ids:
            self.parts.append(f"[{eid}] ")
        self.parts.append(text)
        if eid:
            self.clean_spans[eid] = (self.len_clean, self.len_clean + len(text))
        self.clean.append(text)
        self.len_clean += len(text)

    def text(self) -> str:
        return "".join(self.parts)

    def clean_text(self) -> str:
        return "".join(self.clean)


def _fmt_value(value: str, how: str | None) -> str:
    value = (value or "").strip()
    if how == "date":
        return value[:10]
    return value


def _neighbors(tables: dict[str, Table], spec: dict, anchor: tuple[str, dict]) -> tuple[list[dict], list[dict]]:
    """All rows before/after an anchor row within its group, in order, minus
    excluded rows (e.g. moderator posts) and, if asked, rows without text.
    The caller trims to the wanted counts after dropping rows already shown."""
    t = spec["table"]
    table = tables[t]
    group_col, order_col, id_col = spec["group"], spec["order"], spec["id"]
    anchor_row = anchor[1]
    members = [table.rows[n] for n in table.index(group_col).get((anchor_row.get(group_col) or "").strip(), [])]
    members.sort(key=lambda r: (r.get(order_col, ""), r.get(id_col, "")))
    ids = [r.get(id_col, "") for r in members]
    aid = anchor_row.get(id_col, "") if anchor[0] == t else None
    if aid is not None and aid in ids:
        pos = ids.index(aid)
        before, after = members[:pos], members[pos + 1:]
    else:
        key = anchor_row.get(spec.get("anchor_order", order_col), "")
        pos = bisect_right([r.get(order_col, "") for r in members], key)
        before, after = members[:pos], members[pos:]
    cache: dict = {}
    adjacent = spec.get("adjacent")  # e.g. {"column": "post_number", "max_gap": 8}

    def number(r: dict) -> int | None:
        m = re.search(r"\d+", (r.get(adjacent["column"]) or "") if adjacent else "")
        return int(m.group()) if m else None

    anchor_number = number(anchor_row) if adjacent else None

    def keep(r: dict) -> bool:
        if any(_cond_ok(tables, t, r, cond, cache) for cond in spec.get("exclude", [])):
            return False
        if spec.get("require_text") and not (r.get(spec["text"]) or "").strip():
            return False
        if adjacent and anchor_number is not None:
            n = number(r)
            if n is None or abs(n - anchor_number) > adjacent.get("max_gap", 8):
                return False
        return True

    return [r for r in before if keep(r)], [r for r in after if keep(r)]


def _segment_rows(text: str, by: str, min_chars: int, keep_heading: str | None = None,
                  style: str | None = None, include_quotes: bool = False) -> list[tuple[int, int, str]]:
    """Units within a cell: (start, end, section heading so far). A heading is
    a short capitalised line without end punctuation; it is context, not a
    unit, unless it matches `keep_heading` (e.g. '^(Rule \\d+|Policy:)').
    Lines inside quote blocks (someone else's words) never become units."""
    out = []
    heading = ""
    keep = re.compile(keep_heading) if keep_heading else None
    quoted = [(s.start, s.end) for s in split_quote_blocks(text, style) if s.kind == "quote"] if style else []
    in_quote = lambda a, b: any(qa <= a and b <= qb for qa, qb in quoted)  # noqa: E731
    if by == "line":
        for m in re.finditer(r"[^\n]+", text):
            a, b = m.start(), m.end()
            if in_quote(a, b) and (not include_quotes or _XF_SAID.fullmatch(m.group().strip() + "")
                                   or m.group().strip() == _XF_END or m.group().strip().endswith(" said:")):
                continue
            while a < b and (text[a].isspace() or text[a] in "-–•*\u200b"):
                a += 1
            while b > a and text[b - 1].isspace():
                b -= 1
            line = text[a:b]
            if not line:
                continue
            is_heading = (len(line) < 80 and not line.endswith((".", ":", "?", "!", ";", ","))
                          and line[:1].isupper() and len(line.split()) <= 10)
            if is_heading or (keep and keep.search(line)):
                heading = line
                if not (keep and keep.search(line)):
                    continue
            if len(line) >= min_chars:
                out.append((a, b, heading if heading != line else ""))
    elif by == "paragraph":
        for m in re.finditer(r"(?s)(?:[^\n]|\n(?!\s*\n))+", text):
            a, b = m.start(), m.end()
            while b > a and text[b - 1].isspace():
                b -= 1
            if b - a >= min_chars:
                out.append((a, b, ""))
    else:
        raise SpecError(f"unknown segment mode {by}")
    return out


def _norm_ws(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


class Renderer:
    """Builds the text of one unit from a validated part spec."""

    def __init__(self, tables: dict[str, Table], spec: dict, part_name: str, *, coded: dict | None = None,
                 name_map=None):
        self.tables = tables
        self.spec = spec
        self.part_name = part_name
        self.part = spec["parts"][part_name]
        self.quote_style = (spec.get("quote_blocks") or {}).get("style")
        self.quote_sources = (spec.get("quote_blocks") or {}).get("sources_column")
        self.coded = coded or {}
        self.name_map = name_map or (lambda s: s)

    def _emit_row(self, b: _Builder, ev: dict, block_code: str, table: str, row: dict, row_id: str,
                  text_col: str, speaker: str, in_case: dict[str, str], texts: dict[str, str], *,
                  span: tuple[int, int] | None = None, quoted_chars: int = 800, cite_quotes: bool = False) -> None:
        """One cell into the builder: own words as id'd sentences, quote blocks
        as markers (quoted row in the case) or as marked, id-less quotations."""
        text = row.get(text_col, "") or ""
        lo, hi = span if span else (0, len(text))
        sources = list_values(row.get(self.quote_sources, "")) if self.quote_sources else None
        segs = [s for s in split_quote_blocks(text, self.quote_style, sources) if s.end > lo and s.start < hi]
        if not text.strip() or not segs:
            b.add("(no text)\n")
            return
        n = 0
        for s in segs:
            a, z = max(s.start, lo), min(s.end, hi)
            if s.kind == "quote" and span and cite_quotes:
                b.add(f"‹quoted from {self.name_map(s.author) if s.author else 'someone'} in this post›\n")
                for (x, y) in sentences(text, a, z):
                    n += 1
                    eid = f"{block_code}.{n}"
                    b.add(self.name_map(text[x:y]), eid)
                    b.add("\n")
                    ev[eid] = Evidence(eid, table, row_id, text_col, x, y, s.author or speaker,
                                       block_code.rstrip("0123456789"), text[x:y])
                continue
            if s.kind == "quote":
                body = quote_body(text, Segment(a, z, "quote"))
                marker = None
                if s.source and s.source in in_case and s.source != row_id:
                    marker = in_case[s.source]
                else:
                    probe = _norm_ws(body)[:80]
                    if len(probe) >= 20:
                        marker = next((code for rid, code in in_case.items()
                                       if rid != row_id and probe in texts.get(rid, "")), None)
                if marker:
                    b.add(f"‹quotes {marker}›\n")
                else:
                    who = self.name_map(s.author) if s.author else "someone"
                    shown = body if len(body) <= quoted_chars else body[:quoted_chars].rstrip() + " …"
                    shown = self.name_map(shown)
                    b.add(f"‹quoted from {who}›\n" + "\n".join("> " + x for x in shown.splitlines())
                          + "\n‹end quote›\n")
                continue
            for (x, y) in sentences(text, a, z):
                n += 1
                eid = f"{block_code}.{n}"
                b.add(self.name_map(text[x:y]), eid)
                b.add("\n")
                ev[eid] = Evidence(eid, table, row_id, text_col, x, y, speaker,
                                   block_code.rstrip("0123456789"), text[x:y])

    def _title(self, label: str, code: str, row: dict, speaker_col: str | None, date_col: str | None) -> str:
        bits = [f"== {label} {code}"]
        if speaker_col and (row.get(speaker_col) or "").strip():
            bits.append(f"· {self.name_map(row[speaker_col].strip())}")
        if date_col and (row.get(date_col) or "").strip():
            bits.append(f"· {row[date_col].strip()[:10]}")
        return " ".join(bits) + " ==\n"

    def render(self, unit: tuple[str, dict], *, show_ids: bool = True, segment: tuple[int, int, str] | None = None,
               unit_id: str | None = None, show_codes: bool = True) -> Rendered:
        part = self.part
        table, row = unit
        uid = unit_id or row.get(part["id"], "").strip()
        base_id = row.get(part["id"], "").strip()
        b = _Builder(show_ids)
        ev: dict[str, Evidence] = {}
        blocks: list[dict] = []
        header_vals = {}
        lines = [f"UNIT {self.part_name} {uid}"]
        for h in part.get("header", []):
            v = _fmt_value(row.get(h["column"], ""), h.get("format"))
            if v:
                header_vals[h["column"]] = v
                lines.append(f"{h.get('label', h['column'])}: {self.name_map(v) if h.get('speaker') else v}")
        for hl in part.get("header_lists", []):
            vals = [(r.get(hl["column"]) or "").strip() for _t, r in resolve_path(self.tables, [unit], hl["path"])]
            vals = list(dict.fromkeys(v for v in vals if v))
            if vals:
                lines.append(f"{hl.get('label', hl['column'])}: {', '.join(vals)}")
        b.add("\n".join(lines) + "\n\n")
        unit_code = part.get("code", "U")
        in_case: dict[str, str] = {}
        texts: dict[str, str] = {}
        if part.get("text"):
            in_case[base_id] = f"{unit_code}1"
            texts[base_id] = _norm_ws(row.get(part["text"]["column"], ""))
        related = []
        for rel in part.get("related", []):
            rows = resolve_path(self.tables, [unit], rel["path"])[: rel.get("max", 50)]
            related.append((rel, rows))
            for i, (t, r) in enumerate(rows, 1):
                rid = r.get(rel["id"], "").strip()
                in_case.setdefault(rid, f"{rel['code']}{i}")
                texts.setdefault(rid, _norm_ws(r.get(rel["text"], "")))
        context = []
        for ctx in part.get("context", []):
            anchors = [unit] if not ctx.get("anchor_path") else resolve_path(self.tables, [unit], ctx["anchor_path"])
            if not anchors:
                context.append((ctx, [], []))
                continue
            before, after = _neighbors(self.tables, ctx, anchors[0])
            before = [r for r in before if r.get(ctx["id"], "").strip() not in in_case]
            after = [r for r in after if r.get(ctx["id"], "").strip() not in in_case]
            before = before[-ctx.get("before", 0):] if ctx.get("before") else []
            after = after[: ctx.get("after", 0)] if ctx.get("after") else []
            context.append((ctx, before, after))
            for i, r in enumerate(before, 1):
                rid = r.get(ctx["id"], "").strip()
                in_case.setdefault(rid, f"{ctx.get('code_before', 'E')}{i}")
                texts.setdefault(rid, _norm_ws(r.get(ctx["text"], "")))
            for i, r in enumerate(after, 1):
                rid = r.get(ctx["id"], "").strip()
                in_case.setdefault(rid, f"{ctx.get('code_after', 'L')}{i}")
                texts.setdefault(rid, _norm_ws(r.get(ctx["text"], "")))

        def emit_context(which: str):
            for ctx, before, after in context:
                rows = before if which == "before" else after
                for i, r in enumerate(rows, 1):
                    code = f"{ctx.get('code_' + which, 'E' if which == 'before' else 'L')}{i}"
                    label = ctx.get("label_" + which, "EARLIER POST" if which == "before" else "LATER POST")
                    rid = r.get(ctx["id"], "").strip()
                    b.add(self._title(label, code, r, ctx.get("speaker"), ctx.get("date")))
                    self._emit_row(b, ev, code, ctx["table"], r, rid, ctx["text"], (r.get(ctx.get("speaker") or "", "") or "").strip(),
                                   in_case, texts, quoted_chars=ctx.get("quoted_chars", 800))
                    b.add("\n")
                    blocks.append({"code": code, "table": ctx["table"], "row": rid})

        emit_context("before")
        for rel, rows in related:
            if not rows and rel.get("empty_note"):
                b.add(f"== {rel['label']} ==\n{rel['empty_note']}\n\n")
            for i, (t, r) in enumerate(rows, 1):
                code = f"{rel['code']}{i}"
                rid = r.get(rel["id"], "").strip()
                b.add(self._title(rel["label"], code, r, rel.get("speaker"), rel.get("date")))
                for extra in rel.get("fields", []):
                    v = _fmt_value(r.get(extra["column"], ""), extra.get("format"))
                    if v:
                        b.add(f"{extra.get('label', extra['column'])}: {v}\n")
                self._emit_row(b, ev, code, t, r, rid, rel["text"], (r.get(rel.get("speaker") or "", "") or "").strip(),
                               in_case, texts, quoted_chars=rel.get("quoted_chars", 800))
                b.add("\n")
                blocks.append({"code": code, "table": t, "row": rid})
                inc = rel.get("include_codes")
                if inc and show_codes:
                    prior = (self.coded.get(inc) or {}).get(rid)
                    shown = {k: v for k, v in prior.items() if not k.startswith("_")} if prior else None
                    b.add(f"CODES FROM {inc} FOR {code}: "
                          f"{json.dumps(shown, ensure_ascii=False) if shown else '(not coded)'}\n\n")
        if part.get("text"):
            t_spec = part["text"]
            code = f"{unit_code}1"
            b.add(self._title(t_spec.get("label", "TEXT"), code, row, t_spec.get("speaker"), t_spec.get("date")))
            if segment and segment[2]:
                b.add(f"(section: {segment[2]})\n")
            self._emit_row(b, ev, code, table, row, base_id, t_spec["column"],
                           (row.get(t_spec.get("speaker") or "", "") or "").strip(), in_case, texts,
                           span=(segment[0], segment[1]) if segment else None,
                           quoted_chars=t_spec.get("quoted_chars", 800),
                           cite_quotes=bool((part.get("segment") or {}).get("include_quotes")))
            b.add("\n")
            blocks.append({"code": code, "table": table, "row": base_id})
            seg_spec = part.get("segment") or {}
            if segment and seg_spec.get("neighbors"):
                text = row.get(t_spec["column"], "")
                units = _segment_rows(text, seg_spec.get("by", "line"), seg_spec.get("min_chars", 12),
                                      seg_spec.get("keep_heading"), self.quote_style, seg_spec.get("include_quotes", False))
                pos = next((i for i, u in enumerate(units) if u[0] == segment[0]), None)
                if pos is not None:
                    k = seg_spec["neighbors"]
                    around = units[max(0, pos - k):pos] + units[pos + 1:pos + 1 + k]
                    if around:
                        b.add("NEIGHBOURING STATEMENTS (context only):\n" +
                              "\n".join("  " + self.name_map(text[u[0]:u[1]][:300]) for u in around) + "\n\n")
        emit_context("after")
        text = b.text().rstrip() + "\n"
        clean = b.clean_text().rstrip() + "\n"
        speakers: dict[str, list[str]] = defaultdict(list)
        for eid, e in ev.items():
            speakers[e.code].append(eid)
        rendered = Rendered(uid, self.part_name, text, clean, ev, b.clean_spans, blocks, dict(speakers),
                            {"table": table, "row": base_id, **header_vals})
        rendered.sha256 = sha256_text(text)
        return rendered

    def units(self, screened: dict | None = None) -> list[tuple[str, tuple[str, dict], tuple | None]]:
        """(unit id, base row, segment or None) for every unit of this part."""
        base = unit_rows(self.tables, self.part, screened)
        seg = self.part.get("segment")
        if not seg:
            return [(r.get(self.part["id"], "").strip(), (t, r), None) for t, r in base]
        out = []
        for t, r in base:
            text = r.get(self.part["text"]["column"], "") or ""
            for n, (a, z, heading) in enumerate(_segment_rows(text, seg.get("by", "line"), seg.get("min_chars", 12),
                                                              seg.get("keep_heading"), self.quote_style,
                                                              seg.get("include_quotes", False)), 1):
                out.append((f"{r.get(self.part['id'], '').strip()}-{n:03d}", (t, r), (a, z, heading)))
        return out


# --- blind leak check --------------------------------------------------------------

def own_words(text: str, style: str | None) -> str:
    return "".join(text[s.start:s.end] for s in split_quote_blocks(text, style) if s.kind == "own")


LEAK_COVERAGE = 0.6


def leak_check(rendered_text: str, hidden_texts: list[str], style: str | None, run: int = 40,
               coverage: float = LEAK_COVERAGE) -> list[str]:
    """Hidden texts (their author's own words only) that largely reappear in a
    blind unit's text: at least `coverage` of their 40-character windows are
    found. A moderator quoting a phrase of the post inline covers little of the
    moderator's text and is not flagged; moderator words edited into a post
    (legacy red text) cover most of it and are. Returns the first run of each."""
    body = rendered_text.split("\n\n", 1)[1] if "\n\n" in rendered_text else rendered_text  # skip the header
    hay = re.sub(r"\s+", " ", body)
    hits = []
    for hidden in hidden_texts:
        words = re.sub(r"\s+", " ", own_words(hidden or "", style)).strip()
        if len(words) < 20:
            continue
        if len(words) < run:
            if words in hay:
                hits.append(words)
            continue
        step = max(10, run // 2)
        windows = [words[i:i + run] for i in range(0, len(words) - run + 1, step)]
        found = [w for w in windows if w in hay]
        if windows and len(found) / len(windows) >= coverage:
            hits.append(found[0])
    return hits


def validate_spec(spec: dict, tables: dict[str, Table], codebook: Codebook) -> list[str]:
    """Problems that must be fixed before coding (empty list = fine)."""
    problems = []
    if spec.get("format") != UNITS_FORMAT:
        problems.append(f"units.json format must be {UNITS_FORMAT}")
    for name, part in (spec.get("parts") or {}).items():
        if name not in codebook.parts:
            problems.append(f"{name}: not a part of the codebook ({', '.join(codebook.parts)})")
        names = part.get("table") if isinstance(part.get("table"), list) else [part.get("table")]
        for t in names:
            if t not in tables:
                problems.append(f"{name}: unknown table {t}")
                continue
            if part.get("id") not in tables[t].columns:
                problems.append(f"{name}: table {t} has no id column {part.get('id')}")
        known = [t for t in names if t in tables]
        if known:
            cols = set().union(*(tables[t].columns for t in known))
            for h in part.get("header", []):
                if h["column"] not in cols:
                    problems.append(f"{name}: header column {h['column']} not in {', '.join(known)}")
            if part.get("text") and part["text"]["column"] not in cols:
                problems.append(f"{name}: text column {part['text']['column']} not in {', '.join(known)}")
        for rel in part.get("related", []):
            for hop in rel.get("path", []):
                for side in ("from", "to"):
                    t, _, c = hop[side].partition(".")
                    if t not in tables or c not in tables[t].columns:
                        problems.append(f"{name}: link {hop[side]} does not exist")
        for ctx in part.get("context", []):
            t = ctx.get("table")
            if t not in tables:
                problems.append(f"{name}: context table {t} unknown")
                continue
            for key in ("group", "order", "id", "text"):
                if ctx.get(key) not in tables[t].columns:
                    problems.append(f"{name}: context {key} column {ctx.get(key)} not in {t}")
        for f_name, speakers in (part.get("evidence") or {}).items():
            if f_name not in codebook.fields:
                problems.append(f"{name}: evidence for unknown field {f_name}")
        for f_name in (part.get("int_ranges") or {}):
            if f_name not in codebook.fields or codebook.fields[f_name]["type"] != "int":
                problems.append(f"{name}: int range for non-int field {f_name}")
        for f_name in codebook.parts.get(name, {}).get("fields", []):
            f = codebook.fields[f_name]
            if f["lookup"] and f["lookup"] not in (spec.get("lookups") or {}):
                problems.append(f"{name}: field {f_name} takes values from {f['lookup']}, which units.json does not define")
        inc_parts = {rel.get("include_codes") for rel in part.get("related", []) if rel.get("include_codes")}
        for p in inc_parts:
            if p not in (spec.get("parts") or {}):
                problems.append(f"{name}: includes codes of unknown part {p}")
    for lk, l in (spec.get("lookups") or {}).items():
        if l.get("table") not in tables:
            problems.append(f"lookup {lk}: unknown table {l.get('table')}")
        elif l.get("code") not in tables[l["table"]].columns:
            problems.append(f"lookup {lk}: no column {l.get('code')}")
    return problems


def lookup_values(spec: dict, tables: dict[str, Table], name: str, part: str | None = None) -> dict[str, dict]:
    """code -> {label, definition, date} for a lookup (e.g. CORE_RULES)."""
    lk = spec["lookups"][name]
    out = {}
    for r in tables[lk["table"]].rows:
        code = r.get(lk["code"], "").strip()
        if code:
            out[code] = {"label": r.get(lk.get("label", ""), "").strip(),
                         "definition": r.get(lk.get("definition", ""), "").strip(),
                         "date": r.get(lk.get("date", ""), "").strip()}
    for code, label in ((lk.get("add_by_part") or {}).get(part or "", {}) or {}).items():
        out[code] = {"label": label, "definition": "", "date": ""}
    return out


def field_values(spec: dict, tables: dict[str, Table], codebook: Codebook, field_name: str, part: str) -> dict[str, str] | None:
    """Allowed codes -> meaning for a field (None for int/text)."""
    f = codebook.fields[field_name]
    if f["lookup"]:
        return {c: (v["label"] + (f". {v['definition']}" if v["definition"] else "")) for c, v in
                lookup_values(spec, tables, f["lookup"], part).items()}
    return f["values"]


def neutral_int_range(f: dict) -> tuple[int, int] | None:
    nums = [int(x) for x in re.findall(r"(?<![\w.])(\d{1,2})(?=\s+[a-z])", f.get("def", ""))]
    return (min(nums), max(nums)) if len(nums) >= 2 else None


# --- a run folder ------------------------------------------------------------------

RUN_DIRS = ("inputs", "coded", "screen", "drafts", "exports", "agree", "state")
DEFAULT_READONLY = [p for p in os.environ.get("UPC_TABLES_READONLY", "").split(":") if p]


class GuardError(RuntimeError):
    pass


def _inside(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


class Run:
    """One coding project: a folder holding snapshots of its inputs, the
    profile, the unit spec, codings, screening, drafts and exports."""

    def __init__(self, root: str | Path):
        self.root = Path(root).expanduser().resolve()
        self.config_path = self.root / "config.json"
        self._tables: dict[str, Table] | None = None

    # -- files
    def path(self, *parts: str) -> Path:
        return self.root.joinpath(*parts)

    @property
    def config(self) -> dict:
        return json.loads(self.config_path.read_text(encoding="utf-8")) if self.config_path.exists() else {}

    def save_config(self, cfg: dict) -> None:
        self.write_json(self.config_path, cfg)

    @property
    def name(self) -> str:
        return self.config.get("name") or self.root.name

    @property
    def slug(self) -> str:
        return slug_token(self.name)[:40]

    def readonly_roots(self) -> list[Path]:
        cfg = self.config
        roots = [Path(p).expanduser() for p in cfg.get("readonly", []) + DEFAULT_READONLY]
        roots += [Path(i["source_path"]).expanduser().parent for i in cfg.get("inputs", [])]
        if cfg.get("codebook", {}).get("source_path"):
            roots.append(Path(cfg["codebook"]["source_path"]).expanduser().parent)
        return roots

    def guard(self, path: str | Path) -> Path:
        """Refuse any write inside a read-only root (the inputs' folders, the
        codebook's folder and anything listed as read-only)."""
        p = Path(path).expanduser()
        for root in self.readonly_roots():
            if _inside(p, root):
                raise GuardError(f"refusing to write {p}: it is inside a read-only folder ({root})")
        return p

    def write_text(self, path: str | Path, text: str) -> None:
        p = self.guard(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        tmp = p.with_suffix(p.suffix + ".tmp")
        tmp.write_text(text, encoding="utf-8")
        os.replace(tmp, p)

    def write_json(self, path: str | Path, obj) -> None:
        self.write_text(path, json.dumps(obj, ensure_ascii=False, indent=1) + "\n")

    def append_jsonl(self, path: str | Path, record: dict) -> None:
        p = self.guard(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        with open(p, "a", encoding="utf-8") as fh:
            fh.write(json.dumps(record, ensure_ascii=False) + "\n")

    # -- inputs
    def tables(self) -> dict[str, Table]:
        if self._tables is None:
            tables = {}
            for item in self.config.get("inputs", []):
                snap = self.path("inputs", item["snapshot"])
                for tb in load_table(snap, item["name"]):
                    if tb.sha256 != item["sha256"]:
                        raise GuardError(f"the snapshot {snap.name} changed after init (sha256 differs)")
                    tables[tb.name] = tb
            self._tables = tables
        return self._tables

    def source_drift(self) -> list[str]:
        """Inputs whose original file changed since it was snapshotted."""
        out = []
        for item in self.config.get("inputs", []) + ([self.config["codebook"]] if self.config.get("codebook") else []):
            src = Path(item["source_path"]).expanduser()
            if not src.exists():
                out.append(f"{src} (missing)")
            elif sha256_file(src) != item["sha256"]:
                out.append(str(src))
        return out

    def codebook(self) -> Codebook:
        cb = self.config.get("codebook")
        if not cb:
            raise SpecError("this run has no codebook (init --codebook …)")
        return load_codebook(self.path("inputs", cb["snapshot"]))

    def spec(self) -> dict:
        p = self.path("units.json")
        if not p.exists():
            raise SpecError("no units.json yet: run `units --draft` first")
        return json.loads(p.read_text(encoding="utf-8"))

    def spec_sha(self) -> str:
        spec = self.spec()
        return sha256_text(canonical_json({k: v for k, v in spec.items() if k not in ("approved",)}))[:12]

    # -- codings
    def coded_files(self, part: str) -> list[Path]:
        return sorted(self.path("coded").glob(f"{part}.*.jsonl"))

    def coded(self, part: str, coder: str | None = None) -> dict[str, dict]:
        """unit id -> latest record, for one coder (default: the newest file)."""
        files = [f for f in self.coded_files(part) if coder is None or f.name == f"{part}.{coder}.jsonl"]
        if coder is None:  # the model's codings: not human sheets, reviews or repeat checks
            files = [f for f in files if not f.name.endswith(".repeat.jsonl")
                     and not f.name[len(part) + 1:].startswith(("human-", "review-"))]
        if not files:
            return {}
        path = max(files, key=lambda f: f.stat().st_mtime) if coder is None else files[0]
        out = {}
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                rec = json.loads(line)
                out[rec["unit"]] = rec
        return out

    def coded_values(self, part: str, coder: str | None = None) -> dict[str, dict]:
        """unit id -> field values only (what a dependent part is shown)."""
        return {u: rec.get("fields", {}) for u, rec in self.coded(part, coder).items()}


def snapshot_inputs(run: Run, sources: list[Path], codebook: Path | None) -> dict:
    """Copy inputs into run/inputs (byte-identical) and record their sha256."""
    import shutil  # noqa: PLC0415
    names = unique_names(sources)
    items = []
    run.path("inputs").mkdir(parents=True, exist_ok=True)
    for src in sources:
        name = names[src]
        dest = run.path("inputs", name + src.suffix.lower())
        run.guard(dest)
        shutil.copyfile(src, dest)
        items.append({"name": name, "source_path": str(src), "snapshot": dest.name, "sha256": sha256_file(dest),
                      "bytes": dest.stat().st_size})
    cb = None
    if codebook:
        dest = run.path("inputs", "codebook" + codebook.suffix.lower())
        run.guard(dest)
        shutil.copyfile(codebook, dest)
        cb = {"source_path": str(codebook), "snapshot": dest.name, "sha256": sha256_file(dest)}
    sums = "".join(f"{i['sha256']}  {i['snapshot']}\n" for i in items)
    if cb:
        sums += f"{cb['sha256']}  {cb['snapshot']}\n"
    run.write_text(run.path("inputs", "SHA256SUMS"), sums)
    return {"inputs": items, "codebook": cb}
