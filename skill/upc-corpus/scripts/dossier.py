#!/usr/bin/env python3
"""A dossier: what a set of documents says that matters to one project, written for one audience.

    python3 dossier.py init  --corpus <dir> --audience <a.yaml> --project <p.yaml> --title T [options]
    python3 dossier.py add   --corpus <dir> --from <entry.json> --writer "<who>" [--model M] [--replace] [--plan]
    python3 dossier.py write --corpus <dir> [--source <id>] [--redo] [--plan]
    python3 dossier.py brief --corpus <dir> --from <brief.json> --writer "<who>" [--model M] [--plan]
    python3 dossier.py build --corpus <dir> --out <dir> [--copies] [--no-docx] [--no-xlsx] [--no-images]
    python3 dossier.py check --corpus <dir>

Two profiles steer it. Both are kept outside the corpus and reused from one dossier
to the next (shapes in references/profiles/):

  * an AUDIENCE: who reads it, what they already know, what they look for, how
    they like to read, and phrases that must never be published;
  * a PROJECT: what it is, and the lenses (usually its tasks) that relevance is
    judged against.

The objects are UPC's own, so a dossier is a projection (spec/09):

  * one `x-dossier-entry` generation per document. It holds the writer's summary,
    notes and questions (labelled with the writer's name) and the document's key
    points, figures, dates, people and positions, terms and the works it cites,
    each with a quotation that `upc anchor` found in the stored source. A
    quotation it cannot find is dropped and reported, never stored.
  * one `x-dossier-brief` synthesis: the cover memo across the documents, every
    theme backed by quotations that hop C checks.

`add` takes an entry written anywhere (a person, Claude, any model); `write` has
the configured local model write it. Both go through the same gate.

`build` refuses a corpus that does not validate and finds every quotation again
in the stored source. It then writes the site (a cover that fits the point on one
screen, the memo, one folder per document, a timeline, who's who, terms, leads and
method), a Word file and a spreadsheet. Every page puts what matters at a glance and
the detail one click away. It reads the Word file back and refuses if a quotation is missing
from it. It also refuses if a never-publish phrase (from either profile or the
config) or a local path would appear in anything it writes. The build record is
kept beside the output folder, never in it.

The entry and brief shapes are in references/dossier.md.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import html as htmlmod
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import common as C  # noqa: E402
from corpus_view import (check_extraction, esc, extractions, find_span, generations, latest_per_source,  # noqa: E402
                         meta_content, suppressed_in, thumbnail, validate)

TOOL = {"tool": "upc-corpus/dossier", "tool_version": "1.1.0"}
ENTRY, BRIEF = "x-dossier-entry", "x-dossier-brief"
SCHEMAS = Path(__file__).resolve().parents[1] / "references" / "profiles"
MAX_COPY = 25 * 1024 * 1024  # a static host's per-file limit (Cloudflare Pages)
ITEMS = ("key_points", "figures", "dates", "actors", "terms", "cites")
ITEM_LABEL = {"key_points": "key point", "figures": "figure", "dates": "date", "actors": "person or organization",
              "terms": "term", "cites": "cited work"}
KIND_LABEL = {"paper": "Journal article", "report": "Report", "white-paper": "White paper", "filing": "Regulatory filing",
              "comments": "Comments", "decision": "Decision", "web-page": "Web page", "slides": "Slides",
              "blog": "Blog post", "docket": "Docket document", "other": "Document"}


# ------------------------------------------------------------------ profiles
_TYPES = {"string": str, "array": list, "object": dict, "integer": int, "number": (int, float), "boolean": bool}


def schema_errors(schema: dict, value, where: str = "", root: dict | None = None) -> list[str]:
    """The few JSON Schema rules the profile schemas use (type, required, properties,
    additionalProperties, items, enum, pattern, minItems, minLength, minimum, $ref)."""
    root = root or schema
    if "$ref" in schema:
        node = root
        for part in schema["$ref"].removeprefix("#/").split("/"):
            node = node[part]
        return schema_errors(node, value, where, root)
    here = where or "profile"
    t = schema.get("type")
    if t:
        ts = t if isinstance(t, list) else [t]
        if not any(isinstance(value, _TYPES[x]) and not (x in ("integer", "number") and isinstance(value, bool))
                   for x in ts):
            return [f"{here}: should be {' or '.join(ts)}, not {type(value).__name__}"]
    errs = []
    if "enum" in schema and value not in schema["enum"]:
        errs.append(f"{here}: should be one of {schema['enum']}")
    if isinstance(value, str):
        if "pattern" in schema and not re.search(schema["pattern"], value):
            errs.append(f"{here}: {value!r} should match {schema['pattern']}")
        if len(value) < schema.get("minLength", 0):
            errs.append(f"{here}: {value!r} is shorter than {schema['minLength']} characters")
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value < schema.get("minimum", value):
        errs.append(f"{here}: should be at least {schema['minimum']}")
    if isinstance(value, list):
        if len(value) < schema.get("minItems", 0):
            errs.append(f"{here}: needs at least {schema['minItems']}")
        if isinstance(schema.get("items"), dict):
            for i, v in enumerate(value):
                errs += schema_errors(schema["items"], v, f"{here}[{i}]", root)
    if isinstance(value, dict):
        dot = f"{where}." if where else ""
        errs += [f"{dot}{k}: missing" for k in schema.get("required", []) if value.get(k) in (None, "", [], {})]
        props, extra = schema.get("properties", {}), schema.get("additionalProperties")
        for k, v in value.items():
            if k in props:
                errs += schema_errors(props[k], v, f"{dot}{k}", root)
            elif isinstance(extra, dict):
                errs += schema_errors(extra, v, f"{dot}{k}", root)
    return errs


def _plain_dates(x):
    """YAML reads an unquoted 2026-10-19 as a date; a profile means the text."""
    if isinstance(x, (dt.date, dt.datetime)):
        return x.isoformat()
    if isinstance(x, list):
        return [_plain_dates(v) for v in x]
    if isinstance(x, dict):
        return {k: _plain_dates(v) for k, v in x.items()}
    return x


def load_profile(path: str | Path, kind: str) -> dict:
    """A profile (YAML or JSON), checked against references/profiles/<kind>.schema.json."""
    p = Path(path).expanduser()
    if not p.exists():
        raise SystemExit(f"no such {kind} profile: {p}")
    text = p.read_text(encoding="utf-8")
    if p.suffix in (".yaml", ".yml"):
        try:
            import yaml  # noqa: PLC0415
        except ImportError:
            raise C.Unavailable("reading a YAML profile needs PyYAML (pip install pyyaml); or give it as JSON") from None
        data = _plain_dates(yaml.safe_load(text))
    else:
        data = json.loads(text)
    schema = C.load_json(SCHEMAS / f"{kind}.schema.json", {})
    errs = schema_errors(schema, data)
    if kind == "project" and not errs:
        slugs = [lens["slug"] for lens in data["lenses"]]
        errs += [f"lenses: {s!r} is used twice" for s in sorted({s for s in slugs if slugs.count(s) > 1})]
    if errs:
        raise SystemExit(f"the {kind} profile {p.name} does not fit references/profiles/{kind}.schema.json:\n  "
                         + "\n  ".join(errs[:12]))
    return data


def lens_titles(project: dict) -> dict[str, str]:
    return {lens["slug"]: lens.get("short") or lens["title"] for lens in project["lenses"]}


# ------------------------------------------------------------------ config and sources
def cfg_path(corpus: Path) -> Path:
    return C.state_dir(corpus, "dossier") / "dossier.json"


def load_cfg(corpus: Path) -> dict:
    cfg = C.load_json(cfg_path(corpus), None)
    if not cfg:
        raise SystemExit(f"no dossier here yet: run `dossier.py init --corpus {corpus} …` first")
    return cfg


def source_index(corpus: Path) -> dict[str, dict]:
    """Every source with a reading copy, by source_id and by alias (the same dict)."""
    out: dict[str, dict] = {}
    for alias, info in C.corpus_reps(corpus).items():
        sj = C.load_json(corpus / "sources" / info["slug"] / "source.json", {})
        reps = sj.get("representations", [])
        pdf = next((r for r in reps if str(r.get("media_type", "")).endswith("pdf") and r.get("path")), None)
        raw = next((r for r in reps if r.get("role") == "raw_html" and r.get("path")), None)
        ret = sj.get("retrieval") or {}
        full = {**info, "alias": alias, "record": sj,
                "pdf_path": corpus / pdf["path"] if pdf else None,
                "raw_path": corpus / raw["path"] if raw else None,
                "url": ret.get("final_url") or ret.get("original_url"),
                "fetched_at": ret.get("fetched_at"),
                "sha256": str(ret.get("sha256") or (pdf or raw or {}).get("sha256") or "").removeprefix("sha256:")}
        out[alias] = full
        out[info["source_id"]] = full
    return out


def exs_id(corpus: Path) -> str:
    """The dossier's passage set. Its namespace is recorded in the config the first time, so a corpus that is
    moved or copied keeps the same set instead of starting a new one under its new folder name."""
    cfg = C.load_json(cfg_path(corpus), None) or {}
    ns = cfg.get("namespace") or C.corpus_settings(corpus).get("namespace") \
        or re.sub(r"[^a-z0-9]+", "-", corpus.name.lower()).strip("-") or "dossier"
    if cfg and not cfg.get("namespace"):
        C.save_json(cfg_path(corpus), {**cfg, "namespace": ns})
    return f"exs-{ns}-dossier"


def anchor_quotes(corpus: Path, info: dict, quotes: list[str | None], *, model: str, prompt: str,
                  dry: bool = False) -> list[dict]:
    """Each proposed quotation, found in the source's reading copy by `upc anchor`:
    {quote (the stored characters), extraction_id}, or {quote, status} where it was not.

    Before asking, the same two repairs deep-research's anchor() makes: a quote that
    differs from the page only in case, curly marks, dashes or spacing becomes the
    page's own characters, and one that occurs twice is widened to the shortest
    unique span. Neither can invent text: both return the document's characters."""
    text = Path(info["text_path"]).read_text(encoding="utf-8")
    fixed = []
    for q in quotes:
        q = re.sub(r"[ \t]+", " ", str(q or "")).strip()
        if q and q not in text:
            q = find_span(q, text) or q
        if q and C.occurrences(text, q) > 1:
            q = C.widen_to_unique(text, q) or q
        fixed.append(q)
    out = [{"quote": q, "status": "no quotation given"} for q in fixed]
    todo = [i for i, q in enumerate(fixed) if q]
    if not todo:
        return out
    # --source: a reading copy shared by two sources (one paper saved twice) is one
    # representation; the passages belong to the source this entry is about.
    args = ["anchor", "--corpus", str(corpus), "--rep", info["text_rep"], "--source", info["source_id"],
            "--normalize", "--set", exs_id(corpus),
            "--type", "evidence", "--tool", TOOL["tool"], "--model", model or "unknown", "--prompt-version", prompt]
    res = C.upc_jsonl(*args, *(["--dry-run"] if dry else []), records=[{"quote": fixed[i]} for i in todo])
    if not isinstance(res, dict) or "results" not in res:
        raise C.Unavailable(f"upc anchor gave no results: {str(res)[:200]}")
    by_i = {r["i"]: r for r in res["results"]}
    exts = {} if dry else extractions(corpus)
    for j, i in enumerate(todo):
        r = by_i.get(j) or {}
        if r.get("status") in ("anchored", "refound", "exists"):
            eid = r.get("extraction_id")
            stored = (exts.get(eid) or {}).get("direct_quote") or r.get("stored_quote") or fixed[i]
            out[i] = {"quote": stored, "extraction_id": eid}
        else:
            out[i] = {"quote": fixed[i], "status": r.get("status") or "no result"}
    return out


def pdf_pages(path: Path | None) -> int | None:
    if not path or not path.exists():
        return None
    p = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True)
    m = re.search(r"^Pages:\s+(\d+)", p.stdout, re.M)
    return int(m.group(1)) if m else None


# ------------------------------------------------------------------ init
def cmd_init(corpus: Path, a) -> int:
    cfg = C.load_json(cfg_path(corpus), {}) or {}
    for key, val in (("title", a.title), ("subtitle", a.subtitle), ("prepared_for", a.prepared_for),
                     ("prepared_by", a.prepared_by), ("date", a.date), ("stamp", a.stamp), ("file_no", a.file_no),
                     ("file_name", a.file_name)):
        if val is not None:
            cfg[key] = val
    if a.audience:
        cfg["audience"] = str(Path(a.audience).expanduser().resolve())
    if a.project:
        cfg["project"] = str(Path(a.project).expanduser().resolve())
    if a.order:
        cfg["order"] = [x.strip() for x in a.order.split(",") if x.strip()]
    if a.suppress:
        cfg["suppress"] = sorted(set(cfg.get("suppress", [])) | set(a.suppress))
    for pair in getattr(a, "show_writer_as", None) or []:
        real, sep, label = pair.partition("=")
        if not sep or not real.strip():
            raise SystemExit("--show-writer-as: WRITER=LABEL, e.g. 'Claude (Opus 5.5)=Hermes'")
        labels = cfg.setdefault("writer_labels", {})
        if label.strip():
            labels[real.strip()] = label.strip()
        else:
            labels.pop(real.strip(), None)
    if a.sections:
        secs = C.load_json(a.sections, None)
        if not isinstance(secs, list) or not all(isinstance(s, dict) and s.get("slug") and s.get("title") for s in secs):
            raise SystemExit("--sections: a JSON list of {slug, title, description}")
        cfg["sections"] = secs
    for need in ("title", "audience", "project"):
        if not cfg.get(need):
            raise SystemExit(f"init needs --{need} the first time")
    aud, proj = load_profile(cfg["audience"], "audience"), load_profile(cfg["project"], "project")
    if a.request_source:
        idx = source_index(corpus)
        info = idx.get(a.request_source)
        if not info or not info.get("text_rep"):
            raise SystemExit(f"--request-source: no source {a.request_source!r} with a reading copy in this corpus")
        req = {"source_id": info["source_id"], "by": a.request_by or (cfg.get("request") or {}).get("by"),
               "date": a.request_date or (cfg.get("request") or {}).get("date")}
        if a.request_quote:
            got = anchor_quotes(corpus, info, [a.request_quote], model="none (a person chose it)",
                                prompt="dossier-request@1")[0]
            if not got.get("extraction_id"):
                raise SystemExit(f"--request-quote is not in the request source ({got.get('status')})")
            req.update(extraction_id=got["extraction_id"], quote=got["quote"])
        cfg["request"] = req
    cfg.setdefault("date", dt.date.today().isoformat())
    cfg.setdefault("prepared_for", f"{aud['readers'][0]['name']}, {aud['name']}")
    cfg.setdefault("stamp", f"For {aud['name']} · {proj['name']}")
    C.save_json(cfg_path(corpus), cfg)
    C.counts_header("Dossier: config", [("config", cfg_path(corpus)), ("title", cfg["title"]),
                                        ("audience", f"{aud['name']} ({len(aud['readers'])} readers)"),
                                        ("project", f"{proj['name']} ({len(proj['lenses'])} lenses)"),
                                        ("request", (cfg.get("request") or {}).get("source_id") or "—")])
    return 0


# ------------------------------------------------------------------ add (an entry)
# The cover has to fit on one screen, so the lines it shows have word limits; `add` and `brief` refuse
# anything longer and say by how much. The memo and the document pages hold the rest.
LIMITS = {"for_us": 45, "read_this": 14, "bottom_line": 90, "theme title": 14, "next step": 22, "headline": 24}
MAX_THEMES, MAX_STEPS = 6, 6
# Where a next step can point: a page of the site, or a document (its source id or alias).
PAGE_OF = {"leads": "leads.html", "timeline": "timeline.html", "people": "people.html", "terms": "terms.html",
           "memo": "memo.html", "method": "method.html"}
SEE_LABEL = {"leads": "leads", "timeline": "timeline", "people": "who's who", "terms": "terms", "memo": "memo",
             "method": "method"}


def nwords(s) -> int:
    return len(str(s or "").split())


def over(field: str, text, where: str | None = None) -> str | None:
    n = nwords(text)
    return f"{where or field}: {n} words; at most {LIMITS[field]}" if n > LIMITS[field] else None


def _prose(s: str) -> str:
    """The writer's words in the generation's Markdown: a straight double quotation
    mark there would read as an unchecked quotation, so it becomes a single one."""
    return re.sub(r"\s+", " ", str(s or "")).replace('"', "'").strip()


def marker(quote: str, eid: str) -> str:
    """A UPC quote marker (spec/04): inline, or the block form when the quotation
    holds a double quotation mark or a line break."""
    if '"' in quote or "\n" in quote:
        return f"```quote {eid}\n{quote}\n```"
    return f'"{quote}" [{eid}]'


def entry_markdown(e: dict, project: dict) -> str:
    lt = lens_titles(project)
    md = [f"# {_prose(e['title'])}", "", f"Written by {_prose(e['writer'])}.", "", _prose(e["summary"]), ""]
    if e.get("for_us"):
        md += [f"What matters for {_prose(project['name'])}: {_prose(e['for_us'])}", ""]
    if e.get("read_this"):
        md += [f"Read: {_prose(e['read_this'])}", ""]
    if e.get("why_it_matters"):
        md += [f"## Why it matters for {_prose(project['name'])}", ""]
        md += [f"- {lt.get(w.get('lens'), 'General')}: {_prose(w.get('text'))}" for w in e["why_it_matters"]] + [""]
    for sec in ITEMS:
        if not e.get(sec):
            continue
        md += [f"## {sec.replace('_', ' ').capitalize()}", ""]
        for it in e[sec]:
            head = {"key_points": it.get("text"), "figures": f"{it.get('label')}: {it.get('value')}",
                    "dates": f"{it.get('date')}: {it.get('event')}", "actors": f"{it.get('name')}: {it.get('stance')}",
                    "terms": f"{it.get('term')}: {it.get('definition')}",
                    "cites": f"{it.get('ref')}" + (f" ({it.get('note')})" if it.get("note") else "")}[sec]
            md += [_prose(head), "", marker(it["quote"], it["extraction_id"]), ""]
    if e.get("questions"):
        md += ["## Questions for the team", ""] + [f"- {_prose(q)}" for q in e["questions"]] + [""]
    return "\n".join(md)


def _check_entry_shape(entry: dict, project: dict, sections: list[dict] | None = None) -> list[str]:
    lenses = {lens["slug"] for lens in project["lenses"]}
    errs = []
    if entry.get("section") and entry["section"] not in {s["slug"] for s in sections or []}:
        errs.append(f"section: {entry['section']!r} is not one of the config's sections "
                    f"({', '.join(s['slug'] for s in sections or []) or 'none set: init --sections'})")
    if not str(entry.get("summary") or "").strip():
        errs.append("summary: missing")
    for f in ("for_us", "read_this"):
        if entry.get(f) and not isinstance(entry[f], str):
            errs.append(f"{f}: should be a sentence of text")
        elif (msg := over(f, entry.get(f))):
            errs.append(msg + (" (the document's line on the cover)" if f == "for_us" else ""))
    for i, w in enumerate(entry.get("why_it_matters") or []):
        if w.get("lens") not in lenses:
            errs.append(f"why_it_matters[{i}].lens: {w.get('lens')!r} is not a lens of {project['name']}")
    for k, v in (entry.get("relevance") or {}).items():
        if k not in lenses:
            errs.append(f"relevance: {k!r} is not a lens of {project['name']}")
        elif not (isinstance(v, int) and 0 <= v <= 3):
            errs.append(f"relevance.{k}: should be 0, 1, 2 or 3")
    for i, kp in enumerate(entry.get("key_points") or []):
        if kp.get("lens") and kp["lens"] not in lenses:
            errs.append(f"key_points[{i}].lens: {kp['lens']!r} is not a lens of {project['name']}")
    if errs:
        errs.append(f"(the lenses are: {', '.join(sorted(lenses))})")
    return errs


def add_entry(corpus: Path, entry: dict, *, writer: str, model: str | None = None, method: str = "model",
              replace: bool = False, plan: bool = False, prompt: str = "dossier-entry@1") -> dict:
    """Anchor an entry's quotations and store it as an x-dossier-entry generation."""
    cfg = load_cfg(corpus)
    project = load_profile(cfg["project"], "project")
    idx = source_index(corpus)
    info = idx.get(entry.get("source"))
    if not info:
        raise SystemExit(f"no source {entry.get('source')!r} with a reading copy in this corpus "
                         "(give its source_id or alias; `upc validate` lists them)")
    errs = _check_entry_shape(entry, project, cfg.get("sections"))
    if errs:
        raise SystemExit("the entry does not fit (references/dossier.md, 'An entry'):\n  " + "\n  ".join(errs))
    model_label = model or writer

    flat = [(sec, i, it) for sec in ITEMS for i, it in enumerate(entry.get(sec) or []) if isinstance(it, dict)]
    got = anchor_quotes(corpus, info, [it.get("quote") for _, _, it in flat], model=model_label, prompt=prompt, dry=plan)
    kept: dict[str, list] = {sec: [] for sec in ITEMS}
    refused = []
    for (sec, i, it), g in zip(flat, got):
        if g.get("extraction_id") or (plan and "status" not in g):
            kept[sec].append({**{k: v for k, v in it.items() if k != "quote"}, "quote": g["quote"],
                              "extraction_id": g.get("extraction_id")})
        else:
            refused.append({"section": sec, "item": i, "quote": str(it.get("quote") or "")[:200], "status": g.get("status")})

    sug = None
    if entry.get("suggested_by"):
        s = entry["suggested_by"]
        req = cfg.get("request") or {}
        sinfo = idx.get(s.get("source") or req.get("source_id"))
        if not sinfo:
            refused.append({"section": "suggested_by", "item": 0, "quote": str(s.get("quote"))[:200],
                            "status": "no request source (init --request-source, or give suggested_by.source)"})
        else:
            g = anchor_quotes(corpus, sinfo, [s.get("quote")], model=model_label, prompt=prompt, dry=plan)[0]
            if g.get("extraction_id") or (plan and "status" not in g):
                sug = {"by": s.get("by") or req.get("by"), "source_id": sinfo["source_id"],
                       "extraction_id": g.get("extraction_id"), "quote": g["quote"]}
            else:
                refused.append({"section": "suggested_by", "item": 0, "quote": str(s.get("quote"))[:200],
                                "status": g.get("status")})

    pages = pdf_pages(info.get("pdf_path"))
    exhibits = []
    for i, ex in enumerate(entry.get("exhibits") or []):
        pg = ex.get("page")
        if not pages or not isinstance(pg, int) or not 1 <= pg <= pages:
            refused.append({"section": "exhibits", "item": i, "quote": str(ex.get("caption"))[:200],
                            "status": f"page {pg!r} is not a page of this source's PDF ({pages or 'no PDF'})"})
        else:
            exhibits.append({"page": pg, "caption": str(ex.get("caption") or "").strip()})

    keep = ("title", "short_title", "section", "kind", "authors", "author_org", "date", "identifier", "license", "summary",
            "for_us", "read_this", "why_it_matters", "relevance", "questions", "read_chars", "total_chars")
    e = {**{k: entry[k] for k in keep if entry.get(k) not in (None, "", [], {})}, "writer": writer,
         "title": entry.get("title") or info["title"], "pages": pages or entry.get("pages"),
         **kept, "suggested_by": sug, "exhibits": exhibits, "refused": refused}
    report = {"source_id": info["source_id"], "title": e["title"], "kept": {s: len(v) for s, v in kept.items()},
              "refused": refused, "suggested_by": bool(sug), "exhibits": len(exhibits)}
    if plan:
        report["status"] = "plan"
        return report
    md_path = C.state_dir(corpus, "dossier", "entries") / f"{info['slug']}.md"
    md_path.write_text(entry_markdown(e, project), encoding="utf-8")
    produced = {**TOOL, "method": method, "prompt_version": prompt, **({"model": model} if model else {})}
    res = C.upc("add", "generation", "--corpus", str(corpus), "--output", str(md_path), "--tool", TOOL["tool"],
                *(["--replace"] if replace else []), check=False, stdin=json.dumps({
                    "type": ENTRY, "title": e["title"], "source_id": info["source_id"],
                    "representation_ref": info["text_rep"], "provenance": {"produced_by": produced},
                    "ext": {"upc_corpus": {"dossier_entry": e}}}))
    if not (isinstance(res, dict) and res.get("status") in ("ok", "exists", "replaced")):
        fails = (res or {}).get("failures", []) if isinstance(res, dict) else []
        detail = (res or {}).get("detail") if isinstance(res, dict) else res
        raise SystemExit(f"the entry was refused: {detail}\n  " + "\n  ".join(
            f"{f.get('marker')}: {f.get('reason')}" for f in fails[:8]))
    report.update(status=res["status"], generation_id=res["generation_id"])
    return report


def print_entry_report(rep: dict) -> None:
    kept = ", ".join(f"{n} {ITEM_LABEL[s]}{'s' if n != 1 else ''}" for s, n in rep["kept"].items() if n) or "no quotations"
    print(f"  {rep['title'][:70]:72} {rep.get('status', '')}: {kept}"
          + (f"; {rep['exhibits']} exhibit(s)" if rep.get("exhibits") else "")
          + ("; suggested-by quote" if rep.get("suggested_by") else ""))
    for r in rep["refused"]:
        print(f"    NOT FOUND {r['section']}[{r['item']}] ({r['status']}): {r['quote'][:110]!r}")


# ------------------------------------------------------------------ write (the local model)
def entry_schema(project: dict) -> dict:
    lenses = [lens["slug"] for lens in project["lenses"]]
    s, q = {"type": "string"}, {"type": "string"}
    obj = lambda **p: {"type": "object", "properties": p, "required": list(p)}  # noqa: E731
    arr = lambda item: {"type": "array", "items": item}  # noqa: E731
    return obj(kind={"type": "string", "enum": list(KIND_LABEL)}, author_org=s, date=s, summary=s, for_us=s, read_this=s,
               why_it_matters=arr(obj(lens={"type": "string", "enum": lenses}, text=s)),
               relevance={"type": "object", "properties": {k: {"type": "integer", "minimum": 0, "maximum": 3} for k in lenses}},
               key_points=arr(obj(text=s, lens={"type": "string", "enum": lenses}, quote=q)),
               figures=arr(obj(label=s, value=s, quote=q)), dates=arr(obj(date=s, event=s, quote=q)),
               actors=arr(obj(name=s, type=s, stance=s, quote=q)), terms=arr(obj(term=s, definition=s, quote=q)),
               cites=arr(obj(ref=s, note=s, quote=q)), questions=arr(s))


def style_text(aud: dict) -> str:
    s = aud.get("style") or ["plain, short, specific"]
    return "; ".join(f"{k}: {v}" for k, v in s.items()) if isinstance(s, dict) else "; ".join(s)


def write_prompt(aud: dict, proj: dict) -> str:
    readers = "; ".join(f"{r['name']} ({r['role']})" for r in aud["readers"])
    lenses = "\n".join(f"- {lens['slug']}: {lens['title']}: {lens['description']}" for lens in proj["lenses"])
    return f"""You write one entry of a dossier: what one document says that matters to a project, for its readers.

The readers: {aud['name']}{' (' + aud['organization'] + ')' if aud.get('organization') else ''}: {readers}.
They already know (do not explain): {'; '.join(aud.get('knows') or ['—'])}.
They look for: {'; '.join(aud['wants'])}.
How they like to read: {style_text(aud)}.

The project: {proj['name']}{' (' + proj['full_name'] + ')' if proj.get('full_name') else ''}.
Judge relevance against these lenses:
{lenses}
Project rules: {'; '.join(proj.get('rules') or ['—'])}.

Output JSON only:
- kind, author_org (who wrote or issued it), date (as the document gives it).
- summary: about {aud.get('summary_words', 120)} words in your own words: what the document is, what it argues or
  finds, and whose view it represents. Only what the document says; no quotation marks; nothing from your own knowledge.
- for_us: at most {LIMITS['for_us'] - 5} words, plain: what in this document matters most for {proj['name']}. It is the
  document's one line on the dossier's cover, so make it specific (a finding, a case, a number, a gap), not a description.
- read_this: at most {LIMITS['read_this'] - 2} words: which part to read (say, "the executive summary and Figure 2"),
  or "the whole document" when it is short.
- why_it_matters: for each lens it really bears on, one or two sentences on why, for this project.
- relevance: 0-3 for every lens (0 none, 1 background, 2 useful, 3 central).
- key_points (3-8), figures, dates, actors (people or organizations and the position the document gives them),
  terms (defined in the document), cites (earlier works it leans on most), questions (for the team, at most 3).
- Every quote: copied character for character from the document, one sentence or less (at most 60 words),
  from a single passage. A quotation not in the document is thrown away, and so is its item."""


def cmd_write(corpus: Path, source: str | None, redo: bool, plan: bool, max_chars: int) -> int:
    cfg = load_cfg(corpus)
    aud, proj = load_profile(cfg["audience"], "audience"), load_profile(cfg["project"], "project")
    idx = source_index(corpus)
    have = {e["_source_id"] for e in entries(corpus)}
    req = (cfg.get("request") or {}).get("source_id")
    seen, todo = set(), []
    for info in idx.values():
        sid = info["source_id"]
        if sid in seen or sid == req or not info.get("text_rep"):
            continue
        seen.add(sid)
        if source and source not in (sid, info["alias"]):
            continue
        if redo or sid not in have:
            todo.append(info)
    C.counts_header("Write: plan", [("documents", len(seen)), ("with an entry", len(have)), ("to write now", len(todo)),
                                    ("model", f"{C.DEFAULT_MODEL} at {C.CHAT_URL}"), ("characters read, at most", max_chars)])
    if plan or not todo:
        return 0
    system, schema = write_prompt(aud, proj), entry_schema(proj)
    t0, failed = time.time(), 0
    for info in todo:
        text = Path(info["text_path"]).read_text(encoding="utf-8")
        try:
            out = C.llm_chat(system, f"Title: {info['title']}\nURL: {info.get('url') or '—'}\n\n{text[:max_chars]}",
                             schema=schema, max_tokens=6000, temperature=0.2)
        except C.Unavailable as e:
            print(f"  UNAVAILABLE {info['title'][:60]}: {e}")
            failed += 1
            continue
        d = C.parse_json_loose(out["content"]) or {}
        if not d.get("summary"):
            print(f"  NO ENTRY {info['title'][:60]}: the model's reply had no summary")
            failed += 1
            continue
        d.update(source=info["source_id"], read_chars=min(len(text), max_chars), total_chars=len(text))
        rep = add_entry(corpus, d, writer=f"{C.AGENT_NAME} ({C.DEFAULT_MODEL})", model=C.DEFAULT_MODEL,
                        replace=redo, prompt="dossier-entry@1")
        print_entry_report(rep)
    C.upc("regen", str(corpus), check=False)
    C.counts_header("Write: result", [("written", len(todo) - failed), ("failed", failed),
                                      ("seconds", round(time.time() - t0, 1))])
    return 1 if failed else 0


# ------------------------------------------------------------------ brief (the cover memo)
BRIEF_SECTIONS = (("themes", "Themes"), ("tensions", "Where the documents pull apart"), ("for_project", None))


def add_brief(corpus: Path, brief: dict, *, writer: str, model: str | None = None, method: str = "model",
              plan: bool = False, prompt: str = "dossier-brief@1") -> dict:
    cfg = load_cfg(corpus)
    proj = load_profile(cfg["project"], "project")
    lenses, lt = {lens["slug"] for lens in proj["lenses"]}, lens_titles(proj)
    idx = source_index(corpus)
    exts = extractions(corpus)
    refused, model_label = [], model or writer
    problems = brief_shape(brief, proj, idx)
    if problems:
        raise SystemExit("the brief does not fit (references/dossier.md, 'The cover memo'):\n  " + "\n  ".join(problems))

    # Anchor each quotation in its own source, one CLI call per source.
    asks: dict[str, list] = {}
    for sec, _ in BRIEF_SECTIONS:
        for i, item in enumerate(brief.get(sec) or []):
            if sec == "for_project" and item.get("lens") not in lenses:
                raise SystemExit(f"for_project[{i}].lens: {item.get('lens')!r} is not a lens of {proj['name']}")
            for j, ev in enumerate(item.get("evidence") or []):
                if ev.get("extraction_id"):
                    continue
                info = idx.get(ev.get("source"))
                if not info:
                    refused.append({"where": f"{sec}[{i}].evidence[{j}]", "status": f"no source {ev.get('source')!r}"})
                    continue
                asks.setdefault(info["source_id"], []).append((sec, i, j, ev.get("quote")))
    found: dict[tuple, dict] = {}
    for sid, items in asks.items():
        got = anchor_quotes(corpus, idx[sid], [q for *_, q in items], model=model_label, prompt=prompt, dry=plan)
        for (sec, i, j, _), g in zip(items, got):
            found[(sec, i, j)] = g
    if not plan:
        exts = extractions(corpus)

    steps = []
    for st in steps_of(brief):
        see = st.get("see")
        if see and see not in PAGE_OF:
            see = idx[see]["source_id"]
        steps.append({k: v for k, v in {"step": st["step"], "detail": st.get("detail"), "see": see}.items() if v})
    out: dict = {"writer": writer, "title": brief.get("title") or "Cover memo", "lede": brief.get("lede", ""),
                 "bottom_line": brief.get("bottom_line", ""), "next_steps": steps, "reading_order": []}
    used: set[str] = set()
    for sec, _ in BRIEF_SECTIONS:
        out[sec] = []
        for i, item in enumerate(brief.get(sec) or []):
            evs = []
            for j, ev in enumerate(item.get("evidence") or []):
                if ev.get("extraction_id"):
                    e = exts.get(ev["extraction_id"])
                    if not e or not e.get("direct_quote"):
                        refused.append({"where": f"{sec}[{i}].evidence[{j}]", "status": f"no passage {ev['extraction_id']}"})
                        continue
                    evs.append({"extraction_id": ev["extraction_id"], "quote": e["direct_quote"], "source_id": e["source_id"]})
                else:
                    g = found.get((sec, i, j)) or {}
                    if g.get("extraction_id") or (plan and "status" not in g and g):
                        evs.append({"extraction_id": g.get("extraction_id"), "quote": g["quote"],
                                    "source_id": idx[ev["source"]]["source_id"]})
                    elif g:
                        refused.append({"where": f"{sec}[{i}].evidence[{j}]", "status": g.get("status"),
                                        "quote": str(ev.get("quote"))[:200]})
            if sec != "for_project" and not evs:
                refused.append({"where": f"{sec}[{i}]", "status": "no quotation was found for it, so it was left out",
                                "quote": str(item.get("title"))[:120]})
                continue
            used.update(e["source_id"] for e in evs)
            out[sec].append({**{k: v for k, v in item.items() if k != "evidence"}, "evidence": evs})
    for i, r in enumerate(brief.get("reading_order") or []):
        info = idx.get(r.get("source"))
        if not info:
            refused.append({"where": f"reading_order[{i}]", "status": f"no source {r.get('source')!r}"})
            continue
        out["reading_order"].append({"source_id": info["source_id"], "why": r.get("why", "")})
    out["refused"] = refused
    report = {"themes": len(out["themes"]), "tensions": len(out["tensions"]), "for_project": len(out["for_project"]),
              "quotations": sum(len(x["evidence"]) for s, _ in BRIEF_SECTIONS for x in out[s]), "refused": refused}
    if plan:
        report["status"] = "plan"
        return report
    if not out["themes"]:
        raise SystemExit("the brief has no theme with a quotation found in its source; nothing was written")

    titles = {sid: idx[sid]["title"] for sid in used}
    md = [f"# {_prose(out['title'])}", "", f"Written by {_prose(writer)}.", ""]
    md += [_prose(x) for x in (out["bottom_line"], out["lede"]) if x] + [""]
    claims, n = [], 0
    for sec, head in BRIEF_SECTIONS:
        if not out[sec]:
            continue
        md += [f"## {head or 'For ' + _prose(proj['name'])}", ""]
        for item in out[sec]:
            title = item.get("title") or item.get("headline") or lt.get(item.get("lens"), "")
            md += [f"### {_prose(title)}", "", _prose(item.get("text")), ""]
            for ev in item["evidence"]:
                md += [marker(ev["quote"], ev["extraction_id"]), f"({_prose(titles[ev['source_id']])} · {ev['source_id']})", ""]
            if item["evidence"]:
                n += 1
                claims.append({"claim_id": f"cl-{n:04d}", "text": f"{title}: {_prose(item.get('text'))}"[:600],
                               "evidence_ids": [ev["extraction_id"] for ev in item["evidence"]],
                               "source_ids": sorted({ev["source_id"] for ev in item["evidence"]}), "confidence": "medium"})
    if out["next_steps"]:
        md += ["## Next steps", ""] + [f"- [ ] {_prose(s['step'])}" + (f" ({_prose(s['detail'])})" if s.get("detail") else "")
                                       for s in out["next_steps"]] + [""]
    md += ["## Sources", ""] + [f"- {sid}: {_prose(titles[sid])}" for sid in sorted(used)] + [""]
    md_path = C.state_dir(corpus, "dossier") / "brief.md"
    md_path.write_text("\n".join(md), encoding="utf-8")
    prior = current_brief(corpus)
    produced = {**TOOL, "method": method, "prompt_version": prompt, **({"model": model} if model else {})}
    res = C.upc("add", "synthesis", "--corpus", str(corpus), "--output", str(md_path), "--tool", TOOL["tool"],
                *(["--supersedes", prior["synthesis_id"]] if prior else []), check=False, stdin=json.dumps({
                    "type": BRIEF, "title": out["title"], "claims": claims,
                    "question": f"What do these documents say that matters to {proj['name']}?",
                    "provenance": {"produced_by": produced}, "ext": {"upc_corpus": {"dossier_brief": out}}}))
    if not (isinstance(res, dict) and res.get("status") == "ok"):
        fails = (res or {}).get("failures", []) if isinstance(res, dict) else []
        raise SystemExit(f"the brief was refused: {(res or {}).get('detail') if isinstance(res, dict) else res}\n  "
                         + "\n  ".join(f"{f.get('marker')}: {f.get('reason')}" for f in fails[:8]))
    report.update(status="ok", synthesis_id=res["synthesis_id"])
    return report


def brief_shape(brief: dict, proj: dict, idx: dict) -> list[str]:
    """What keeps the cover on one screen: a bottom line, a few findings and next steps, each short."""
    errs = []
    if not str(brief.get("bottom_line") or brief.get("lede") or "").strip():
        errs.append("bottom_line: missing (one paragraph: the point of the whole file)")
    if (msg := over("bottom_line", brief.get("bottom_line"))):
        errs.append(msg)
    themes = brief.get("themes") or []
    if len(themes) > MAX_THEMES:
        errs.append(f"themes: {len(themes)}; at most {MAX_THEMES} (the cover lists each; move the rest into tensions or the text)")
    for i, t in enumerate(themes):
        if (msg := over("theme title", t.get("title"), f"themes[{i}].title")):
            errs.append(msg + " (a finding the cover lists, so a short statement)")
    for i, t in enumerate(brief.get("for_project") or []):
        if (msg := over("headline", t.get("headline"), f"for_project[{i}].headline")):
            errs.append(msg)
    raw = brief.get("next_steps") or []
    if len(raw) > MAX_STEPS:
        errs.append(f"next_steps: {len(raw)}; at most {MAX_STEPS}")
    for i, st in enumerate(raw):
        if not isinstance(st, (str, dict)) or (isinstance(st, dict) and not str(st.get("step") or "").strip()):
            errs.append(f"next_steps[{i}]: a string, or {{step, detail, see}} with a step")
            continue
        text = st if isinstance(st, str) else st["step"]
        if (msg := over("next step", text, f"next_steps[{i}]" + ("" if isinstance(st, str) else ".step"))):
            errs.append(msg + " (put the rest in detail)")
        see = st.get("see") if isinstance(st, dict) else None
        if see and see not in PAGE_OF and see not in idx:
            errs.append(f"next_steps[{i}].see: {see!r} is not a page ({', '.join(PAGE_OF)}) or a source of this corpus")
    return errs


# ------------------------------------------------------------------ reading the dossier back
def entries(corpus: Path) -> list[dict]:
    """The newest entry for each document (an older one stays in the corpus as history)."""
    out = []
    for g in latest_per_source(generations(corpus, ENTRY)):
        e = ((g.get("ext") or {}).get("upc_corpus") or {}).get("dossier_entry")
        if e:
            out.append({**e, "_gen": g["generation_id"], "_source_id": g["_source_id"], "_created": g["_created"],
                        "_model": ((g.get("provenance") or {}).get("produced_by") or {}).get("model")})
    return out


def current_brief(corpus: Path) -> dict | None:
    best = None
    for sj in corpus.glob("syntheses/*/synthesis.json"):
        y = C.load_json(sj, {})
        if y.get("type") != BRIEF or y.get("superseded_by"):
            continue
        if best is None or (y.get("provenance") or {}).get("created_at", "") > (best.get("provenance") or {}).get("created_at", ""):
            best = y
    return best


def page_of(e: dict) -> int | None:
    return next((s.get("value") for s in e.get("secondary_locators") or [] if s.get("type") == "page"), None)


def recheck(corpus: Path, es: list[dict], br: dict | None, cfg: dict) -> tuple[list[str], int]:
    """Find every quotation again in its stored source; attach its page. -> (problems, count)."""
    exts = extractions(corpus)
    problems, n = [], 0

    def one(item: dict, where: str):
        nonlocal n
        eid = item.get("extraction_id")
        if not eid:
            problems.append(f"{where}: no passage id")
            return
        try:
            e = check_extraction(corpus, eid, exts)
        except SystemExit as err:
            problems.append(f"{where}: {err}")
            return
        if e["direct_quote"] != item.get("quote"):
            problems.append(f"{where}: its quotation differs from the stored passage {eid}")
            return
        item["page"] = page_of(e)
        n += 1

    for en in es:
        for sec in ITEMS:
            for i, it in enumerate(en.get(sec) or []):
                one(it, f"{en['title'][:50]} · {sec}[{i}]")
        if en.get("suggested_by"):
            one(en["suggested_by"], f"{en['title'][:50]} · suggested_by")
    if br:
        b = br["ext"]["upc_corpus"]["dossier_brief"]
        for sec, _ in BRIEF_SECTIONS:
            for i, item in enumerate(b.get(sec) or []):
                for j, ev in enumerate(item.get("evidence") or []):
                    one(ev, f"brief · {sec}[{i}].evidence[{j}]")
    if (cfg.get("request") or {}).get("extraction_id"):
        one(cfg["request"], "config · request")
    return problems, n


def cmd_check(corpus: Path) -> int:
    rep = validate(corpus)
    cfg = load_cfg(corpus)
    aud, proj = load_profile(cfg["audience"], "audience"), load_profile(cfg["project"], "project")
    es, br = entries(corpus), current_brief(corpus)
    problems, n = recheck(corpus, es, br, cfg)
    errors = rep.get("errors") or []
    C.counts_header("Dossier: check", [("corpus errors (upc validate)", len(errors)), ("entries", len(es)),
                                       ("cover memo", br["synthesis_id"] if br else "none yet"),
                                       ("quotations found again", n), ("quotations NOT found", len(problems)),
                                       ("audience", aud["name"]), ("project", proj["name"])])
    for e in errors[:8]:
        print(f"  ERROR {e.get('code')}: {str(e.get('detail'))[:140]}")
    for p in problems[:20]:
        print(f"  NOT FOUND {p}")
    return 1 if errors or problems else 0


# ------------------------------------------------------------------ the site
# Every page follows one rule: what matters at a glance, the detail one click away. Lists become cards
# (a face to scan, the passage behind it in a panel), filters and a search narrow them, and nothing a
# reader needs is hidden without a visible way in. Without JavaScript the cards open in place.

# One quiet palette, light and dark. --k1..--k6 tell groups apart (timeline threads, kinds of actor).
LIGHT = ("--desk:#efece6;--folder:#f7f3ea;--folder-2:#ebe4d5;--folder-edge:#d9cfbb;--paper:#fffdf9;--ink:#2a2723;"
         "--muted:#6f685d;--rule:#e8e1d4;--stamp:#a0533f;--link:#2f5f80;--postit:#fbf4d6;--postit-ink:#4a4128;"
         "--chip:#f3eee4;--tape:rgba(226,214,186,.75);--shadow:0 1px 2px rgba(50,40,25,.06),0 4px 14px rgba(50,40,25,.06);"
         "--k1:#6b8f71;--k2:#5c7a99;--k3:#b08a3e;--k4:#8b6589;--k5:#4f8a8b;--k6:#b0705a;color-scheme:light")
DARK = ("--desk:#161513;--folder:#201e1a;--folder-2:#2a2722;--folder-edge:#3d382f;--paper:#1c1a17;--ink:#ece7de;"
        "--muted:#a69e91;--rule:#322e28;--stamp:#d99a86;--link:#93bedf;--postit:#36321f;--postit-ink:#efe6c6;"
        "--chip:#26231e;--tape:rgba(110,100,75,.45);--shadow:0 1px 2px rgba(0,0,0,.3),0 4px 14px rgba(0,0,0,.25);"
        "--k1:#93b598;--k2:#8ba7c4;--k3:#d2ae66;--k4:#b893b6;--k5:#7db5b6;--k6:#d69a85;color-scheme:dark")
FONTS = ('--mono:"American Typewriter","Courier Prime","Courier New",ui-monospace,Menlo,monospace;'
         '--serif:"Iowan Old Style","Palatino Linotype",Palatino,Georgia,serif;'
         '--sans:-apple-system,BlinkMacSystemFont,"Segoe UI",system-ui,sans-serif')

CSS = (":root{" + LIGHT + ";" + FONTS + "}\n@media (prefers-color-scheme:dark){:root:not([data-theme=light]){" + DARK + "}}\n"
       ":root[data-theme=dark]{" + DARK + "}\n" + """
*{box-sizing:border-box}
html{scroll-padding-top:12px}
body{margin:0;background:var(--desk);color:var(--ink);font:16px/1.6 var(--sans);-webkit-text-size-adjust:100%;overflow-wrap:break-word}
a{color:var(--link);text-underline-offset:2px}
h1,h2,h3{font-family:var(--serif);font-weight:600;line-height:1.22;margin:0 0 .5em}
h1{font-size:clamp(1.5rem,2.5vw,2.15rem);letter-spacing:-.01em}
h2{font-size:1.28rem} h3{font-size:1.06rem}
p{margin:0 0 .8em}
.wrap{max-width:1600px;margin:0 auto;padding:0 clamp(12px,2.4vw,32px) 40px}
nav.tabs{display:flex;gap:4px;overflow-x:auto;padding:20px 0 0;scrollbar-width:thin}
nav.tabs a{flex:0 0 auto;background:var(--folder-2);color:var(--ink);text-decoration:none;font:600 .72rem/1 var(--mono);
letter-spacing:.07em;text-transform:uppercase;padding:11px 15px 10px;border:1px solid var(--folder-edge);border-bottom:none;
border-radius:9px 9px 0 0}
nav.tabs a:hover{background:var(--folder)}
nav.tabs a.here{background:var(--folder);position:relative;top:1px}
.topbar{display:flex;align-items:flex-end;gap:8px}
.topbar nav.tabs{flex:1 1 auto;min-width:0}
button.theme{flex:none;margin:0 0 7px;display:inline-flex;align-items:center;gap:6px;background:none;color:var(--muted);
border:1px solid var(--folder-edge);border-radius:999px;padding:6px 11px;font:600 .68rem/1 var(--mono);letter-spacing:.07em;
text-transform:uppercase;cursor:pointer}
button.theme:hover,button.theme:focus-visible{color:var(--ink);background:var(--chip)}
button.theme svg{width:14px;height:14px;fill:none;stroke:currentColor;stroke-width:2;stroke-linecap:round}
button.theme[hidden],button.theme[data-to=dark] .sun,button.theme[data-to=light] .moon{display:none}
@media (max-width:560px){button.theme span{display:none}button.theme{padding:7px}}
main.folder{background:var(--folder);border:1px solid var(--folder-edge);border-radius:0 12px 12px 12px;
padding:clamp(14px,2.2vw,30px);box-shadow:var(--shadow)}
.sheet{background:var(--paper);border:1px solid var(--rule);border-radius:4px;padding:clamp(14px,2vw,26px);
box-shadow:var(--shadow);position:relative;margin:0 0 20px}
.prose{max-width:74ch}
.label{font:600 .7rem/1.3 var(--mono);letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}
.byline,.src{font-size:.84rem;color:var(--muted)}
.lede{font-size:1.03rem;color:var(--muted);max-width:76ch;margin:0}
.note{font-size:.84rem;color:var(--muted)}
.stamp{display:inline-block;font:700 .72rem/1 var(--mono);letter-spacing:.12em;text-transform:uppercase;color:var(--stamp);
border:2px solid var(--stamp);border-radius:5px;padding:6px 9px;transform:rotate(-2deg);opacity:.8;white-space:nowrap}
.btn{display:inline-flex;align-items:center;gap:6px;background:var(--paper);border:1px solid var(--folder-edge);
color:var(--ink);text-decoration:none;font-size:.86rem;padding:6px 12px;border-radius:8px;white-space:nowrap}
.btn:hover{border-color:var(--ink)}
.crumbs{font-size:.86rem;margin:0 0 10px}
.page-head{display:flex;flex-wrap:wrap;justify-content:space-between;align-items:flex-end;gap:12px 28px;margin:0 0 16px}
.page-head h1{margin:0 0 .2em}
.stats{display:flex;flex-wrap:wrap;gap:8px}
.stat{background:var(--paper);border:1px solid var(--rule);border-radius:4px;padding:7px 14px 8px;min-width:96px}
.stat b{display:block;font:600 1.45rem/1.15 var(--serif)}
.stat span{font-size:.76rem;color:var(--muted)}
.toolbar{display:flex;flex-wrap:wrap;align-items:center;gap:8px 14px;margin:0 0 14px}
.chips{display:flex;gap:6px;align-items:center;overflow-x:auto;padding-bottom:3px;max-width:100%;scrollbar-width:thin}
.chips>.label{flex:none;margin-right:2px}
button.chip{flex:none;font:inherit;font-size:.84rem;border:1px solid var(--folder-edge);background:var(--chip);color:var(--ink);
padding:4px 11px;border-radius:999px;cursor:pointer;white-space:nowrap;display:inline-flex;align-items:center;gap:6px}
button.chip:hover{border-color:var(--ink)}
button.chip[aria-pressed=true]{background:var(--ink);color:var(--paper);border-color:var(--ink)}
button.chip b{font:700 .72rem var(--mono);color:var(--stamp)} button.chip[aria-pressed=true] b{color:inherit}
button.chip i{font-style:normal;font-size:.72rem;opacity:.7}
input.find{font:inherit;font-size:.9rem;padding:6px 13px;border:1px solid var(--folder-edge);border-radius:999px;
background:var(--paper);color:var(--ink);width:min(260px,100%)}
.shown{font-size:.8rem;color:var(--muted)}
.k1{--k:var(--k1)}.k2{--k:var(--k2)}.k3{--k:var(--k3)}.k4{--k:var(--k4)}.k5{--k:var(--k5)}.k6{--k:var(--k6)}
.sw{display:inline-block;width:10px;height:10px;border-radius:50%;flex:none;background:var(--k,var(--muted))}
.cards{display:grid;grid-template-columns:repeat(auto-fill,minmax(250px,1fr));gap:12px}
details.card{background:var(--paper);border:1px solid var(--rule);border-top:3px solid var(--k,var(--folder-edge));
border-radius:4px;box-shadow:var(--shadow);min-width:0}
details.card>summary{list-style:none;cursor:pointer;padding:10px 13px 12px;display:flex;flex-direction:column;gap:4px;height:100%}
details.card>summary::-webkit-details-marker{display:none}
details.card:hover{border-color:var(--folder-edge);border-top-color:var(--k,var(--stamp))}
details.card>summary:focus-visible{outline:2px solid var(--link);outline-offset:2px}
details.card[open]>summary{height:auto}
details.card .more{padding:6px 13px 13px;border-top:1px dashed var(--rule)}
.clamp2,.clamp3{display:-webkit-box;-webkit-box-orient:vertical;overflow:hidden}
.clamp2{-webkit-line-clamp:2}.clamp3{-webkit-line-clamp:3}
.meta{font:600 .66rem/1.35 var(--mono);letter-spacing:.07em;text-transform:uppercase;color:var(--muted)}
.t{font-weight:600;line-height:1.35}
.from{display:flex;align-items:center;gap:6px;margin-top:auto;padding-top:4px;font-size:.78rem;color:var(--muted)}
dialog.drawer{border:none;padding:0;background:var(--paper);color:var(--ink);border-radius:10px;
box-shadow:0 24px 70px rgba(0,0,0,.35);width:min(760px,calc(100vw - 24px));max-height:min(86vh,940px);overflow:auto}
dialog.drawer::backdrop{background:rgba(28,24,18,.45)}
.dr-bar{position:sticky;top:0;z-index:1;display:flex;align-items:center;gap:6px;padding:8px 10px 8px 18px;
background:var(--paper);border-bottom:1px solid var(--rule)}
.dr-bar .pos{margin-right:auto;font:600 .68rem var(--mono);letter-spacing:.07em;text-transform:uppercase;color:var(--muted)}
.dr-bar button{font:inherit;font-size:.82rem;border:1px solid var(--rule);background:var(--chip);color:var(--ink);
border-radius:6px;padding:4px 10px;cursor:pointer}
.dr-body{padding:16px 22px 24px}
.dr-body .face{display:flex;flex-direction:column;gap:4px;margin:0 0 14px;padding-bottom:12px;border-bottom:1px solid var(--rule)}
.dr-body .face .t{font:600 1.22rem/1.3 var(--serif)}
.dr-body .face .clamp2,.dr-body .face .clamp3{display:block;-webkit-line-clamp:unset;overflow:visible}
.dr-body .face .from{margin-top:2px}
@media (max-width:640px){dialog.drawer{width:100vw;max-width:100vw;margin:auto 0 0;border-radius:14px 14px 0 0;max-height:88vh}}
blockquote.q{margin:8px 0 4px;padding:2px 0 2px 14px;border-left:3px solid var(--folder-edge);font-family:var(--serif);font-size:1.02rem}
ul.ev,ol.points,ul.plain{list-style:none;margin:0;padding:0}
ul.ev>li{margin:0 0 16px}
details.ev{margin:4px 0 6px} details.ev>summary{cursor:pointer;width:max-content;max-width:100%;font:600 .7rem/1.3 var(--mono);
letter-spacing:.08em;text-transform:uppercase;color:var(--muted);padding:4px 0} details.ev[open]>summary{margin-bottom:6px}
button.copy{font:600 .66rem/1 var(--mono);letter-spacing:.06em;text-transform:uppercase;border:1px solid var(--rule);
background:transparent;color:var(--muted);border-radius:5px;padding:3px 7px;margin-left:8px;cursor:pointer}
.postit{background:var(--postit);color:var(--postit-ink);padding:14px 16px 16px;box-shadow:0 2px 7px rgba(60,45,20,.14);
transform:rotate(-.5deg);border-bottom-right-radius:22px 7px;margin:0 0 18px}
.postit h2,.postit h3,.postit .label{color:var(--postit-ink)}
.postit .who{display:block;font-size:.72rem;opacity:.75;margin-top:6px}
.postits{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:16px;margin:6px 0 22px}
.postits .postit:nth-child(even){transform:rotate(.6deg)}
.slip{background:var(--paper);border:1px dashed var(--folder-edge);padding:12px 14px;margin:0 0 18px;font-size:.9rem}
.slip p{margin:4px 0 0} .slip blockquote{margin:6px 0 0;font-family:var(--serif)}
ul.check li{padding-left:28px;position:relative;margin:0 0 9px;list-style:none}
ul.check li::before{content:"";position:absolute;left:2px;top:4px;width:14px;height:14px;border:1.5px solid currentColor;
border-radius:3px;opacity:.6}
ul.check{padding:0;margin:0}
.lens-tag{font:600 .66rem/1.2 var(--mono);letter-spacing:.07em;text-transform:uppercase;color:var(--stamp)}
.scroll{overflow-x:auto}
table.ledger{width:100%;border-collapse:collapse;font-size:.9rem}
table.ledger th{font:600 .68rem/1.2 var(--mono);letter-spacing:.07em;text-transform:uppercase;color:var(--muted);
text-align:left;border-bottom:2px solid var(--rule);padding:8px}
table.ledger td{border-bottom:1px solid var(--rule);padding:9px 8px;vertical-align:top}
code{font-size:.8rem;overflow-wrap:anywhere}
.exhibits{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:26px;margin-top:10px}
figure.exhibit{margin:0;background:var(--paper);padding:10px 10px 8px;box-shadow:var(--shadow);position:relative}
figure.exhibit::before{content:"";position:absolute;top:-11px;left:50%;width:96px;height:24px;margin-left:-48px;
background:var(--tape);transform:rotate(-2deg)}
figure.exhibit img{width:100%;display:block;border:1px solid var(--rule)}
figure.exhibit figcaption{font-size:.82rem;color:var(--muted);margin-top:6px}
.prevnext{display:flex;justify-content:space-between;gap:12px;margin-top:18px;font-size:.9rem}
footer.site{margin:26px 0 0;color:var(--muted);font-size:.82rem;max-width:90ch}
.hide{display:none!important}
.sec{margin:26px 0 0}
.sec-head{display:flex;flex-wrap:wrap;align-items:baseline;gap:4px 14px;margin:0 0 10px}
.sec-head h2{margin:0;font-size:1.25rem}
.d{display:block;border-radius:50%}
.d3{width:15px;height:15px;background:var(--stamp)}
.d2{width:11px;height:11px;margin-top:2px;background:var(--muted);opacity:.8}
.d1{width:8px;height:8px;margin-top:3px;border:1.5px solid var(--muted);opacity:.7}
.d0{width:4px;height:4px;margin-top:5px;background:var(--rule)}
.d.il{display:inline-block;vertical-align:middle;margin:0 3px 0 8px}

/* the cover */
.cv-head{display:grid;grid-template-columns:minmax(0,1fr) auto;gap:4px 24px;align-items:start;margin-bottom:16px}
.cv-head h1{margin:.15em 0 .15em}
.cv-meta{margin:0;font-size:.9rem;color:var(--muted)} .cv-meta b{color:var(--ink);font-weight:600}
.cv-actions{display:flex;gap:8px;justify-content:flex-end;align-items:center}
.cv-actions .stamp{margin-right:6px}
.glance{display:grid;grid-template-columns:minmax(0,1.65fr) minmax(300px,1fr);gap:clamp(16px,2vw,28px);align-items:start}
.bl{margin:0;padding:clamp(16px,2vw,26px) clamp(16px,2.2vw,30px)}
.bl-top{display:flex;justify-content:space-between;align-items:baseline;gap:12px}
.bl h2{font-size:1.2rem;margin:0}
p.bottom{font:1.12rem/1.58 var(--serif);max-width:74ch;margin:.5em 0 1em}
ol.find{list-style:none;margin:2px 0 0;padding:0;border-top:1px solid var(--rule)}
ol.find a{display:flex;gap:10px;align-items:baseline;padding:7px 2px;border-bottom:1px solid var(--rule);color:var(--ink);text-decoration:none}
ol.find a:hover .ft,ol.find a:focus-visible .ft{text-decoration:underline}
ol.find .n{font:700 .74rem/1 var(--mono);color:var(--stamp);min-width:1.1em}
ol.find .ft{font-weight:600;font-size:.98rem;line-height:1.35}
ol.find .go{margin-left:auto;padding-left:10px;font-size:.78rem;color:var(--muted);white-space:nowrap}
p.more{margin:10px 0 0;font-size:.9rem}
.side-col .steps{margin:6px 0 0;padding:16px 18px 14px}
.steps h2{font-size:1.2rem;margin:0 0 10px}
.steps ul.check li{font-size:.95rem;line-height:1.4;margin:0 0 10px}
.steps li a{font-size:.8rem;white-space:nowrap;color:inherit;opacity:.85}
.side-col .slip{margin:18px 0 0;font-size:.86rem}
.reg{background:var(--paper);border:1px solid var(--rule);border-radius:4px;box-shadow:var(--shadow)}
.rh,a.r{display:grid;grid-template-columns:60px minmax(210px,1.05fr) minmax(0,2.1fr) repeat(var(--ncol),40px);gap:0 14px;align-items:start}
.rh{position:sticky;top:0;z-index:2;background:var(--paper);border-bottom:2px solid var(--rule);padding:9px 14px 7px;
font:600 .68rem/1.2 var(--mono);letter-spacing:.07em;text-transform:uppercase;color:var(--muted);align-items:end}
.rh i{font-style:normal;opacity:.75}
.rh button.c{all:unset;cursor:pointer;text-align:center;padding:3px 0;border-radius:4px;color:var(--ink)}
.rh button.c:hover,.rh button.c.on{background:var(--ink);color:var(--paper)}
.grp h3{margin:0;padding:12px 14px 6px;font:600 .72rem/1.3 var(--mono);letter-spacing:.08em;text-transform:uppercase;
color:var(--stamp);border-top:1px solid var(--rule)}
.reg .grp:first-of-type h3{border-top:none}
a.r{padding:9px 14px;color:var(--ink);text-decoration:none;border-top:1px solid var(--rule)}
.grp h3+a.r{border-top:none}
a.r:hover{background:var(--chip)} a.r:hover b{text-decoration:underline}
a.r .th{width:60px;height:76px;overflow:hidden;border:1px solid var(--rule);background:var(--chip);display:block}
a.r .th img{width:100%;display:block}
a.r .th .ph{display:flex;height:100%;align-items:flex-end;padding:5px;font:600 .58rem/1.1 var(--mono);text-transform:uppercase;color:var(--muted)}
a.r .ti{display:flex;flex-direction:column;gap:2px;min-width:0}
a.r .ti b{font-size:.98rem;line-height:1.3}
a.r .ft{font-size:.8rem;line-height:1.35;color:var(--muted)}
a.r .mt{font:600 .64rem/1.3 var(--mono);letter-spacing:.07em;text-transform:uppercase;color:var(--stamp)}
a.r .fu{font-size:.92rem;line-height:1.45;max-width:80ch}
a.r .read{display:block;margin-top:3px;font:600 .66rem/1.3 var(--mono);letter-spacing:.05em;text-transform:uppercase;color:var(--muted)}
a.r .c{display:flex;justify-content:center;padding-top:6px}
a.r .chips-sm{display:none}
.filterable.f [data-col]{opacity:.35} .filterable.f [data-col].on{opacity:1}
.lensnote{background:var(--paper);border:1px solid var(--rule);border-left:4px solid var(--stamp);padding:9px 14px;margin:0 0 12px;font-size:.94rem}
.lensnote[hidden]{display:none}

/* the memo */
.memo-grid{display:grid;grid-template-columns:230px minmax(0,1fr);gap:28px;align-items:start}
.toc{position:sticky;top:14px;font-size:.88rem}
.toc ol{list-style:none;margin:6px 0 16px;padding:0}
.toc li{margin:0 0 5px} .toc a{color:var(--ink);text-decoration:none} .toc a:hover{text-decoration:underline}
.toc .sub{padding-left:12px;font-size:.84rem}
.memo{max-width:900px;padding-top:34px}
.clip{position:absolute;top:-24px;left:26px;width:22px;height:64px;color:var(--muted);transform:rotate(7deg)}
.memo-head{display:grid;grid-template-columns:auto 1fr;gap:3px 14px;font:.86rem/1.45 var(--mono);
border-bottom:2px solid var(--rule);padding-bottom:12px;margin-bottom:16px}
.memo-head b{letter-spacing:.06em}
.memo h2{margin-top:1.2em}
.point{border-top:1px solid var(--rule);padding-top:14px;margin-top:16px}
.point h3 .n{font:700 .8rem var(--mono);color:var(--stamp);margin-right:6px}
ol.steps-full{list-style:none;margin:0;padding:0}
ol.steps-full>li{padding:10px 0 10px 30px;border-top:1px solid var(--rule);position:relative}
ol.steps-full>li::before{content:"";position:absolute;left:2px;top:14px;width:15px;height:15px;border:1.5px solid var(--muted);border-radius:3px}
ol.steps-full b{display:block}

/* the timeline */
.ribbon{display:flex;align-items:flex-end;gap:5px;padding:12px 14px 0;background:var(--paper);border:1px solid var(--rule);
border-radius:4px;box-shadow:var(--shadow);overflow-x:auto;margin:0 0 16px;scrollbar-width:thin}
.ribbon a.bar{flex:1 0 38px;display:flex;flex-direction:column;align-items:stretch;justify-content:flex-end;text-decoration:none;color:var(--muted)}
.ribbon .stack{display:flex;flex-direction:column-reverse;gap:1px;margin:0 6px}
.ribbon .stack i{display:block;background:var(--k);border-radius:2px;min-height:3px}
.ribbon .n{font:600 .66rem var(--mono);text-align:center;color:var(--muted);padding-bottom:2px}
.ribbon .yr{font:600 .7rem var(--mono);text-align:center;padding:5px 0 8px;border-top:1px solid var(--rule);color:var(--ink)}
.ribbon a.bar:hover .yr{color:var(--stamp)}
.ribbon a.bar.ahead .stack i{opacity:.45}
.ribbon a.bar.ahead .yr{color:var(--muted);font-style:italic}
.ribbon .gap{flex:0 0 18px;align-self:flex-end;text-align:center;color:var(--muted);padding-bottom:8px}
.year{display:grid;grid-template-columns:92px minmax(0,1fr);gap:0 18px;position:relative;padding:0 0 22px}
.year::before{content:"";position:absolute;left:45px;top:6px;bottom:0;width:2px;background:var(--folder-edge)}
.year h2.yl{position:sticky;top:10px;align-self:start;margin:0;text-align:center;font-size:1.2rem;background:var(--folder);
z-index:1;padding:4px 0 6px;border-radius:6px}
.year h2.yl span{display:block;font-size:.8em;line-height:1.1}
.year h2.yl small{display:block;font:600 .62rem var(--mono);color:var(--muted);letter-spacing:.06em;text-transform:uppercase}
.year .cards{grid-template-columns:repeat(auto-fill,minmax(250px,1fr))}
details.ev-card.ahead{background:repeating-linear-gradient(135deg,var(--paper) 0 12px,var(--chip) 12px 14px)}
.tl-note{font-size:.82rem;color:var(--muted);margin:-6px 0 14px}

/* who's who, terms, leads */
.who{display:flex;gap:10px;align-items:center}
.av{flex:none;display:grid;place-items:center;width:38px;height:38px;border-radius:50%;background:var(--k,var(--muted));
color:var(--paper);font:700 .7rem var(--mono);letter-spacing:.02em}
.grp-sec{margin:0 0 22px}
.grp-sec>h2{display:flex;align-items:center;gap:8px;font-size:1.12rem;margin:0 0 10px}
.grp-sec>h2 .sw{width:12px;height:12px}
.grp-sec>h2 small{font:600 .7rem var(--mono);color:var(--muted)}
.letters{display:flex;flex-wrap:wrap;gap:4px}
.letters a{font:600 .8rem var(--mono);text-decoration:none;color:var(--ink);background:var(--paper);border:1px solid var(--rule);
border-radius:5px;padding:3px 8px}
.letters a:hover{border-color:var(--ink)}
.lsec{display:contents}
.terms-flow .letter{margin:0;padding:8px 4px 0;font:600 2.6rem/1 var(--serif);color:var(--stamp);border-top:3px solid var(--stamp)}
.badge{align-self:flex-start;font:600 .62rem/1.2 var(--mono);letter-spacing:.07em;text-transform:uppercase;border-radius:999px;
padding:2px 8px;border:1px solid currentColor}
.badge.here{color:var(--k1)} .badge.pull{color:var(--stamp)}

/* a document's folder */
.docgrid{display:grid;grid-template-columns:minmax(0,1fr) 300px;gap:28px;align-items:start}
.matters{background:var(--paper);border:1px solid var(--rule);border-left:4px solid var(--stamp);border-radius:4px;
padding:14px 18px;box-shadow:var(--shadow);margin:0 0 18px}
.matters p.big{font:1.12rem/1.55 var(--serif);margin:.3em 0 .4em;max-width:74ch}
.side .shot{background:var(--paper);padding:8px;box-shadow:var(--shadow);margin:0 0 16px}
.side .shot img{width:100%;display:block}
dl.facts{margin:0;font-size:.88rem} dl.facts dt{font:600 .66rem/1.2 var(--mono);letter-spacing:.08em;text-transform:uppercase;
color:var(--muted);margin-top:10px} dl.facts dd{margin:2px 0 0;overflow-wrap:anywhere}
.relgrid{display:grid;grid-template-columns:auto auto;justify-content:start;gap:4px 10px;font-size:.84rem;margin-top:12px;align-items:center}
.relgrid .d{display:inline-block;vertical-align:middle}
ol.points>li{padding:12px 0;border-top:1px solid var(--rule)} ol.points>li:first-child{border-top:none;padding-top:0}
.jump{display:flex;flex-wrap:wrap;gap:6px;margin:0 0 18px}
.jump a{font-size:.82rem;background:var(--paper);border:1px solid var(--rule);border-radius:999px;padding:3px 11px;
text-decoration:none;color:var(--ink)} .jump a:hover{border-color:var(--ink)}
.figs{grid-template-columns:repeat(auto-fill,minmax(190px,1fr))}
.fig .v{font:600 1.3rem/1.2 var(--serif)}

@media (max-width:1280px){.memo-grid{grid-template-columns:200px minmax(0,1fr)}}
@media (max-width:1000px){.glance{grid-template-columns:minmax(0,1fr)}
.rh,a.r{grid-template-columns:52px minmax(0,1fr) minmax(0,1.6fr) repeat(var(--ncol),30px);gap:0 10px}
.docgrid,.memo-grid{grid-template-columns:minmax(0,1fr)}.side{order:-1}.side .shot{max-width:220px}.toc{position:static}}
@media (max-width:760px){.cv-head{grid-template-columns:minmax(0,1fr)}.cv-actions{justify-content:flex-start;flex-wrap:wrap}
.rh{display:none}
a.r{grid-template-columns:52px minmax(0,1fr);gap:4px 12px}
a.r .th{grid-row:span 2;width:52px;height:66px}
a.r .fu,a.r .chips-sm{grid-column:2}
a.r .c{display:none}
a.r .chips-sm{display:flex;flex-wrap:wrap;gap:5px}
a.r .chips-sm span{font-size:.7rem;background:var(--chip);border:1px solid var(--rule);border-radius:999px;padding:1px 8px;white-space:nowrap}
a.r .chips-sm b{letter-spacing:1px;color:var(--stamp);font-weight:400}
.year{grid-template-columns:58px minmax(0,1fr);gap:0 10px}.year::before{left:28px}
.year .cards,.cards{grid-template-columns:minmax(0,1fr)}}
@media print{:root:is([data-theme],:not([data-theme])){""" + LIGHT + """;--paper:#fff;--folder:#fff;--chip:#fff;--postit:#fff}
body{background:#fff}nav.tabs,.toolbar,button.copy,button.theme,.btn,.ribbon,.toc,dialog{display:none}
main.folder,.sheet,details.card,figure.exhibit,.reg{box-shadow:none;border-color:#ccc;background:#fff}
.postit,.stamp{transform:none}.postit{border:1px solid #bbb}a{color:inherit}.rh,.year h2.yl{position:static}
details.card .more{display:block}.clamp2,.clamp3{display:block;-webkit-line-clamp:unset}}
""")

JS = """
function upcApply(root){var g=root.dataset.g||'',w=root.dataset.w||'',q=(root.dataset.q||'').trim().toLowerCase();
root.classList.toggle('f',!!g);
root.querySelectorAll('[data-item]').forEach(function(it){var ok=(!g||(' '+(it.dataset.g||'')+' ').indexOf(' '+g+' ')>=0)&&
(!w||it.dataset.w===w)&&(!q||it.textContent.toLowerCase().indexOf(q)>=0);it.classList.toggle('hide',!ok)});
root.querySelectorAll('[data-section]').forEach(function(s){s.classList.toggle('hide',!s.querySelector('[data-item]:not(.hide)'))});
root.querySelectorAll('[data-for]').forEach(function(n){n.hidden=n.dataset.for!==g});
root.querySelectorAll('[data-col]').forEach(function(c){c.classList.toggle('on',!!g&&c.dataset.col===g)});
root.querySelectorAll('button[data-set]').forEach(function(b){b.setAttribute('aria-pressed',String((root.dataset[b.dataset.set]||'')===b.dataset.v))});
var n=root.querySelector('.shown');if(n)n.textContent=root.querySelectorAll('[data-item]:not(.hide)').length+' shown'}
document.addEventListener('click',function(e){var b=e.target.closest('button[data-set]');if(!b)return;var root=b.closest('.filterable');
if(!root)return;var s=b.dataset.set;root.dataset[s]=(root.dataset[s]===b.dataset.v)?'':b.dataset.v;upcApply(root)});
document.querySelectorAll('.filterable').forEach(function(root){var i=root.querySelector('input.find');
if(i)i.addEventListener('input',function(){root.dataset.q=i.value;upcApply(root)});upcApply(root)});
document.addEventListener('click',function(e){var b=e.target.closest('button.copy');if(!b)return;var t=b.getAttribute('data-cite');
function done(){b.textContent='Copied';setTimeout(function(){b.textContent='Copy'},1400)}
if(navigator.clipboard){navigator.clipboard.writeText(t).then(done,function(){})}});
window.addEventListener('beforeprint',function(){document.querySelectorAll('details').forEach(function(d){d.open=true})});
(function(){var d=document.getElementById('drawer');if(!d||!d.showModal)return;
var body=d.querySelector('.dr-body'),pos=d.querySelector('.pos'),cur=null;
function list(){return Array.prototype.filter.call(document.querySelectorAll('details.card'),function(c){return c.offsetParent!==null})}
function show(c){cur=c;var s=c.querySelector('summary'),m=c.querySelector('.more');
var k=(c.className.match(/\bk\d\b/)||[''])[0];body.innerHTML='<div class="face '+k+'">'+s.innerHTML+'</div>'+(m?m.innerHTML:'');
var L=list(),i=L.indexOf(c);pos.textContent=i<0?'':(i+1)+' of '+L.length;if(!d.open)d.showModal();d.scrollTop=0}
function step(k){var L=list(),i=L.indexOf(cur);if(i<0||!L.length)return;show(L[(i+k+L.length)%L.length])}
document.addEventListener('click',function(e){var s=e.target.closest('details.card>summary');if(!s||e.target.closest('a'))return;
e.preventDefault();show(s.parentElement)});
d.querySelector('.prev').addEventListener('click',function(){step(-1)});
d.querySelector('.next').addEventListener('click',function(){step(1)});
d.querySelector('.x').addEventListener('click',function(){d.close()});
d.addEventListener('click',function(e){if(e.target===d)d.close()});
d.addEventListener('keydown',function(e){if(e.key==='ArrowRight'){step(1);e.preventDefault()}else if(e.key==='ArrowLeft'){step(-1);e.preventDefault()}});
d.addEventListener('close',function(){var s=cur&&cur.querySelector('summary');if(s)s.focus()})})();
(function(){var r=document.documentElement,b=document.querySelector('button.theme');if(!b)return;
var mq=window.matchMedia?window.matchMedia('(prefers-color-scheme: dark)'):null;
function cur(){return r.getAttribute('data-theme')||(mq&&mq.matches?'dark':'light')}
function show(){var to=cur()==='dark'?'light':'dark';b.setAttribute('data-to',to);b.title='Switch to '+to+' mode';
b.setAttribute('aria-label',b.title);b.querySelector('span').textContent=to==='dark'?'Dark':'Light'}
b.hidden=false;show();
b.addEventListener('click',function(){var t=cur()==='dark'?'light':'dark';r.setAttribute('data-theme',t);
try{localStorage.setItem('upc-dossier-theme',t)}catch(e){}show()});
if(mq&&mq.addEventListener)mq.addEventListener('change',show)})();
"""

# Before the stylesheet, so a reader's choice of light or dark applies without a flash.
THEME_HEAD = ("<script>try{var t=localStorage.getItem('upc-dossier-theme');"
              "if(t==='light'||t==='dark')document.documentElement.setAttribute('data-theme',t)}catch(e){}</script>")

THEME_BUTTON = ('<button class="theme" type="button" hidden><svg class="moon" viewBox="0 0 24 24" aria-hidden="true">'
                '<path d="M20 14.5A8 8 0 1 1 9.5 4a6.5 6.5 0 0 0 10.5 10.5z"/></svg><svg class="sun" viewBox="0 0 24 24" '
                'aria-hidden="true"><circle cx="12" cy="12" r="4"/><path d="M12 2v2M12 20v2M4.9 4.9l1.4 1.4M17.7 17.7l1.4 1.4'
                'M2 12h2M20 12h2M4.9 19.1l1.4-1.4M17.7 6.3l1.4-1.4"/></svg><span>Dark</span></button>')

# The one panel every card opens into.
DRAWER = ('<dialog class="drawer" id="drawer" aria-label="Details"><div class="dr"><div class="dr-bar"><span class="pos"></span>'
          '<button class="prev" type="button" aria-label="Previous">←</button><button class="next" type="button" '
          'aria-label="Next">→</button><button class="x" type="button">Close</button></div><div class="dr-body"></div></div></dialog>')

CLIP = ('<svg class="clip" viewBox="0 0 22 64" aria-hidden="true"><path d="M6 14v36a5 5 0 0 0 10 0V9a7.5 7.5 0 0 0-15 0v42" '
        'fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round"/></svg>')


def fmt_day(s: str | None) -> str:
    if not s:
        return ""
    try:
        return dt.date.fromisoformat(str(s)[:10]).strftime("%-d %b %Y")
    except ValueError:
        return str(s)


def fmt_when(s) -> str:
    """A date as the documents give it, made readable: 9 Mar 2022, Mar 2022, 2022, or as written."""
    s = str(s or "").strip()
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
            return dt.date.fromisoformat(s).strftime("%-d %b %Y")
        if re.fullmatch(r"\d{4}-\d{2}", s):
            return dt.date(int(s[:4]), int(s[5:7]), 1).strftime("%b %Y")
    except ValueError:
        pass
    return s


def slug(s: str, n: int = 48) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")[:n].strip("-") or "document"


def short_title(d: dict) -> str:
    if d.get("short_title"):
        return d["short_title"]
    year = re.search(r"(19|20)\d\d", str(d.get("date") or ""))
    if d.get("author_org") and year:
        return f"{d['author_org']} {year.group(0)}"
    t = d["title"]
    return t if len(t) <= 60 else t[:57].rsplit(" ", 1)[0] + "…"


def first_sentence(s) -> str:
    return re.split(r"(?<=[.!?])\s+", str(s or "").strip(), maxsplit=1)[0]


def cite_text(quote: str, d: dict, page: int | None) -> str:
    return f"“{quote}” ({short_title(d)}{f', p. {page}' if page else ''})"


def quote_html(item: dict, d: dict, link: str | None = None) -> str:
    src = f'<a href="{esc(link)}">{esc(short_title(d))}</a>' if link else esc(short_title(d))
    page = f", p. {item['page']}" if item.get("page") else ""
    return (f'<blockquote class="q">“{esc(item["quote"])}”</blockquote><div class="src">{src}{page}'
            f'<button class="copy" type="button" data-cite="{esc(cite_text(item["quote"], d, item.get("page")))}">Copy</button></div>')


def dots(n: int) -> str:
    return f'<b aria-label="{n} of 3">{"●" * n}{"○" * (3 - n)}</b>'


def dot(n: int) -> str:
    return f'<span class="d d{n}" aria-label="{n} of 3"></span>' if n else '<span class="d d0" aria-hidden="true"></span>'


def lens_code(lens: dict) -> str:
    """A column label of a few characters: "2.1" for "2.1 Transitions", "H&E" for "Households & equity"."""
    short = str(lens.get("short") or lens["title"])
    first = short.split()[0]
    if re.fullmatch(r"[\d.]+", first):
        return first
    words = [w for w in re.split(r"[\s/,]+", short) if w[:1].isalnum() and w.lower() not in ("and", "the", "of", "for")]
    return ("&" if "&" in short else "").join(w[0].upper() for w in words[:2]) or short[:3]


def lens_chip(lens: dict) -> str:
    short, code = str(lens.get("short") or lens["title"]), lens_code(lens)
    if short.startswith(code + " "):
        return f"<b>{esc(code)}</b> {esc(short[len(code) + 1:])}"
    return esc(short)


def rel_html(d: dict, ctx: dict) -> str:
    rel = d.get("relevance") or {}
    rows = [f'{dot(rel.get(lens["slug"], 0))}<span>{esc(lens.get("short") or lens["title"])}</span>'
            for lens in ctx["proj"]["lenses"]]
    return f'<div class="relgrid" aria-label="How much it bears on each lens">{"".join(rows)}</div>' if rel else ""


def lenses_of(d: dict) -> str:
    return " ".join(k for k, v in (d.get("relevance") or {}).items() if v >= 2)


def card(face: str, more: str, *, cls: str = "", attrs: str = "") -> str:
    """A face to scan and the detail behind it: opens in the panel, or in place without JavaScript."""
    return f'<details class="card {cls}" data-item {attrs}><summary>{face}</summary><div class="more">{more}</div></details>'


def chips(group: str, label: str, options: list[tuple[str, str]]) -> str:
    """Filter buttons: (value, label HTML); the first is usually ("", "All")."""
    return (f'<div class="chips" role="group" aria-label="{esc(label)}"><span class="label">{esc(label)}</span>'
            + "".join(f'<button type="button" class="chip" data-set="{group}" data-v="{esc(v)}" aria-pressed="false">{lab}</button>'
                      for v, lab in options) + "</div>")


def sw(cls: str) -> str:
    return f'<span class="sw {cls}"></span>'


def find_box(what: str) -> str:
    return f'<input class="find" type="search" placeholder="Search {esc(what)}" aria-label="Search {esc(what)}"><span class="shown"></span>'


def page_html(title: str, body: str, ctx: dict, here: str, depth: int = 0) -> str:
    up = "../" * depth
    tabs = [("index.html", "Cover"), ("memo.html", "Memo"), ("timeline.html", "Timeline"), ("people.html", "Who's who"),
            ("terms.html", "Terms"), ("leads.html", "Leads"), ("method.html", "Method")]
    nav = "".join(f'<a href="{up}{h}" class="{"here" if h == here else ""}">{esc(t)}</a>'
                  for h, t in tabs if h in ctx["pages"])
    cfg = ctx["cfg"]
    foot = (f'Prepared for {esc(cfg["prepared_for"])}{" by " + esc(cfg["prepared_by"]) if cfg.get("prepared_by") else ""}, '
            f'{esc(fmt_day(cfg.get("date")))}. Summaries, notes and the memo are their writers\' own words, labelled as such. '
            "Every quotation matches the saved copy of its source, character for character. Fixity is not authenticity: "
            f'a match shows which words were read, not that they are right. <a href="{up}method.html">How this was made</a>.')
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow">
<meta name="build" content="{esc(ctx['build'])}"><title>{esc(title)}</title>{THEME_HEAD}<style>{CSS}</style></head><body>
<div class="wrap"><div class="topbar"><nav class="tabs" aria-label="Sections">{nav}</nav>{THEME_BUTTON}</div><main class="folder">{body}</main>
<footer class="site">{foot}</footer></div>{DRAWER}<script>{JS}</script></body></html>
"""


def steps_of(b: dict | None) -> list[dict]:
    """Next steps as {step, detail, see}; a plain string is a step with no detail."""
    return [s if isinstance(s, dict) else {"step": str(s)} for s in (b or {}).get("next_steps") or []]


def see_link(see: str | None, ctx: dict, up: str = "") -> str | None:
    if not see:
        return None
    if see in PAGE_OF and PAGE_OF[see] in ctx["pages"]:
        return up + PAGE_OF[see]
    d = next((d for d in ctx["docs"] if d["_source_id"] == see), None)
    return f"{up}docs/{d['_file']}" if d else None


def ev_list(evs: list[dict], ctx: dict, up: str = "") -> str:
    # Folded, so a point reads first and its passages are one click away.
    by_sid = {d["_source_id"]: d for d in ctx["docs"]}
    n = len(evs)
    return (f'<details class="ev"><summary>The passage{"" if n == 1 else "s"} ({n})</summary><ul class="ev">' + "".join(
        f'<li>{quote_html(ev, by_sid[ev["source_id"]], up + "docs/" + by_sid[ev["source_id"]]["_file"]) if ev["source_id"] in by_sid else quote_html(ev, ctx["others"][ev["source_id"]])}</li>'
        for ev in evs) + "</ul></details>") if evs else ""


def cover_groups(ctx: dict) -> list[tuple[str, list[dict]]]:
    """The documents in reading order: the memo's first three, then the rest of its order, then the others by section."""
    by = {d["_source_id"]: d for d in ctx["docs"]}
    order = [r["source_id"] for r in (ctx["brief"] or {}).get("reading_order") or [] if r["source_id"] in by]
    groups = []
    if order:
        groups.append(("Read first", [by[s] for s in order[:3]]))
        if order[3:]:
            groups.append(("Then", [by[s] for s in order[3:]]))
    rest = [d for d in ctx["docs"] if d["_source_id"] not in order]
    for s, ds in doc_groups({**ctx, "docs": rest}):
        groups.append((s["title"] if s else ("Also in the file" if order else "The documents"), ds))
    return [(t, ds) for t, ds in groups if ds]


def reg_row(d: dict, ctx: dict) -> str:
    rel = d.get("relevance") or {}
    kind = KIND_LABEL.get(d.get("kind"), d.get("kind") or "Document")
    thumb = (f'<img src="{esc(d["_cover"])}" alt="" loading="lazy">' if d.get("_cover") else f'<span class="ph">{esc(kind)}</span>')
    year = re.search(r"(19|20)\d\d", str(d.get("date") or ""))
    meta = " · ".join(x for x in (kind, year.group(0) if year else "") if x)
    read = f'<span class="read">Read: {esc(d["read_this"])}</span>' if d.get("read_this") else ""
    cells = "".join(f'<span class="c" data-col="{esc(lens["slug"])}" title="{esc(lens["title"])}: {rel.get(lens["slug"], 0)} of 3">'
                    f'{dot(rel.get(lens["slug"], 0))}</span>' for lens in ctx["proj"]["lenses"])
    small = "".join(f'<span>{esc(lens_code(lens))} <b>{"●" * rel[lens["slug"]]}{"○" * (3 - rel[lens["slug"]])}</b></span>'
                    for lens in ctx["proj"]["lenses"] if rel.get(lens["slug"], 0) >= 2)
    return (f'<a class="r" href="docs/{esc(d["_file"])}" data-item data-g="{esc(lenses_of(d))}">'
            f'<span class="th">{thumb}</span><span class="ti"><b>{esc(short_title(d))}</b>'
            f'<span class="ft clamp2">{esc(d["title"])}</span><span class="mt">{esc(meta)}</span></span>'
            f'<span class="fu">{esc(d.get("for_us") or first_sentence(d.get("summary")))}{read}</span>{cells}'
            f'<span class="chips-sm">{small}</span></a>')


def cover_html(ctx: dict) -> str:
    cfg, proj, docs, b = ctx["cfg"], ctx["proj"], ctx["docs"], ctx["brief"]
    lenses = proj["lenses"]
    dl = "".join(f'<a class="btn" href="{esc(f)}" download>{esc(label)}</a>' for f, label in ctx["downloads"])
    memo_btn = '<a class="btn" href="memo.html">Full memo</a>' if b else ""
    who = [f"For <b>{esc(cfg['prepared_for'])}</b>"] + ([f"Prepared by {esc(cfg['prepared_by'])}"] if cfg.get("prepared_by") else []) \
        + ([esc(fmt_day(cfg.get("date")))] if cfg.get("date") else [])
    head = (f'<header class="cv-head"><div><span class="label">{esc(cfg.get("file_no") or "Dossier")} · {len(docs)} documents · '
            f'{ctx["quotes"]} passages, each checked against its source</span><h1>{esc(cfg["title"])}</h1>'
            f'<p class="cv-meta">{" · ".join(who)}</p></div><div class="cv-actions">'
            f'{"<span class=stamp>" + esc(cfg["stamp"]) + "</span>" if cfg.get("stamp") else ""}{memo_btn}{dl}</div></header>')

    if b:
        bottom = b.get("bottom_line") or b.get("lede") or ""
        finds = "".join(f'<li><a href="memo.html#theme-{i}"><span class="n">{i}</span><span class="ft">{esc(t.get("title"))}</span>'
                        f'<span class="go">passages →</span></a></li>' for i, t in enumerate(b.get("themes") or [], 1))
        tens = len(b.get("tensions") or [])
        left = (f'<article class="sheet bl"><div class="bl-top"><h2>The bottom line</h2><span class="label">Written by '
                f'{esc(b["writer"])}</span></div><p class="bottom">{esc(bottom)}</p>'
                + (f'<h3 class="label">Findings · each opens its passages</h3><ol class="find">{finds}</ol>' if finds else "")
                + (f'<p class="more"><a href="memo.html#tensions">Where the documents pull apart ({tens}) →</a></p>' if tens else "")
                + "</article>")
        steps = steps_of(b)
        items = "".join(f'<li>{esc(s["step"])} '
                        + (f'<a href="{esc(see_link(s.get("see"), ctx))}">{esc(SEE_LABEL.get(s.get("see"), "open"))} →</a>'
                           if see_link(s.get("see"), ctx) else f'<a href="memo.html#step-{i}">details →</a>' if s.get("detail") else "")
                        + "</li>" for i, s in enumerate(steps, 1))
        right = (f'<aside class="postit steps"><h2>Next steps</h2><ul class="check">{items}</ul>'
                 f'<span class="who">— {esc(b["writer"])}</span></aside>') if steps else ""
    else:
        left = '<article class="sheet bl"><p class="note">No cover memo yet (dossier.py brief).</p></article>'
        right = ""
    req = cfg.get("request") or {}
    slip = (f'<div class="slip"><span class="label">Routing slip</span><p>Resources suggested by <b>{esc(req.get("by") or "")}</b>'
            f'{", " + esc(fmt_day(req.get("date"))) if req.get("date") else ""}:</p><blockquote>“{esc(req["quote"])}”</blockquote></div>'
            if req.get("quote") else "")
    glance = f'<section class="glance" aria-label="At a glance">{left}<div class="side-col">{right}{slip}</div></section>'

    used = [lens for lens in lenses if any(lens["slug"] in lenses_of(d).split() for d in docs)]
    notes = {t.get("lens"): t for t in (b or {}).get("for_project") or []}
    lensbar = chips("g", "By task", [("", f"All {len(docs)}")] + [
        (lens["slug"], f'{lens_chip(lens)} <i>{sum(lens["slug"] in lenses_of(d).split() for d in docs)}</i>') for lens in used])
    lnotes = "".join(
        f'<p class="lensnote" data-for="{esc(lens["slug"])}" hidden><b>{esc(lens["title"])}.</b> '
        f'{esc(notes[lens["slug"]].get("headline") or first_sentence(notes[lens["slug"]].get("text")))} '
        f'<a href="memo.html#{esc(lens["slug"])}">The paragraph in the memo, with passages →</a></p>'
        for lens in used if lens["slug"] in notes)
    head_row = (f'<div class="rh" aria-hidden="true"><span></span><span>Document</span><span>What matters for {esc(proj["name"])} '
                f'<i>({esc(docs[0]["writer"]) if docs else ""})</i></span>'
                + "".join(f'<button type="button" class="c" data-set="g" data-v="{esc(lens["slug"])}" data-col="{esc(lens["slug"])}" '
                          f'title="{esc(lens["title"])}" tabindex="-1">{esc(lens_code(lens))}</button>' for lens in lenses) + "</div>")
    groups = "".join(f'<div class="grp" data-section><h3>{esc(t)}</h3>{"".join(reg_row(d, ctx) for d in ds)}</div>'
                     for t, ds in cover_groups(ctx))
    reg = (f'<section class="sec filterable" id="documents"><div class="sec-head"><h2>The documents</h2><span class="note">'
           f'{"In reading order. " if b and b.get("reading_order") else ""}How much each bears on each lens: '
           f'<span class="d d3 il"></span>central <span class="d d2 il"></span>useful <span class="d d1 il"></span>background.</span></div>'
           f'<div class="toolbar">{lensbar}</div>{lnotes}'
           f'<div class="reg" style="--ncol:{len(lenses)}">{head_row}{groups}</div></section>')
    return head + glance + reg


def memo_html(ctx: dict) -> str:
    cfg, proj, b = ctx["cfg"], ctx["proj"], ctx["brief"]
    by_sid = {d["_source_id"]: d for d in ctx["docs"]}
    lmap = {lens["slug"]: lens for lens in proj["lenses"]}
    toc, body = [], []
    bottom = b.get("bottom_line") or b.get("lede") or ""
    body.append(f'<p class="bottom" id="bottom-line">{esc(bottom)}</p>')
    if b.get("lede") and b.get("bottom_line") and b["lede"] != b["bottom_line"]:
        body.append(f'<p>{esc(b["lede"])}</p>')
    if b.get("themes"):
        toc.append('<li><a href="#findings">Findings</a><ol class="sub">' + "".join(
            f'<li><a href="#theme-{i}">{esc(t.get("title"))}</a></li>' for i, t in enumerate(b["themes"], 1)) + "</ol></li>")
        body.append('<h2 id="findings">Findings</h2>' + "".join(
            f'<section class="point" id="theme-{i}"><h3><span class="n">{i}</span>{esc(t.get("title"))}</h3>'
            f'<p>{esc(t.get("text"))}</p>{ev_list(t["evidence"], ctx)}</section>' for i, t in enumerate(b["themes"], 1)))
    if b.get("tensions"):
        toc.append(f'<li><a href="#tensions">Where the documents pull apart ({len(b["tensions"])})</a></li>')
        body.append('<h2 id="tensions">Where the documents pull apart</h2>' + "".join(
            f'<section class="point" id="tension-{i}"><h3>{esc(t.get("title"))}</h3><p>{esc(t.get("text"))}</p>'
            f'{ev_list(t["evidence"], ctx)}</section>' for i, t in enumerate(b["tensions"], 1)))
    if b.get("for_project"):
        toc.append(f'<li><a href="#by-lens">What it means for {esc(proj["name"])}</a><ol class="sub">' + "".join(
            f'<li><a href="#{esc(t.get("lens"))}">{esc((lmap.get(t.get("lens")) or {}).get("short") or t.get("lens"))}</a></li>'
            for t in b["for_project"]) + "</ol></li>")
        body.append(f'<h2 id="by-lens">What it means for {esc(proj["name"])}</h2>' + "".join(
            f'<section class="point" id="{esc(t.get("lens"))}"><span class="lens-tag">{esc((lmap.get(t.get("lens")) or {}).get("title", ""))}</span>'
            + (f'<h3>{esc(t["headline"])}</h3>' if t.get("headline") else "")
            + f'<p>{esc(t.get("text"))}</p>{ev_list(t.get("evidence") or [], ctx)}</section>' for t in b["for_project"]))
    steps = steps_of(b)
    if steps:
        toc.append('<li><a href="#next-steps">Next steps</a></li>')
        body.append('<h2 id="next-steps">Next steps</h2><ol class="steps-full">' + "".join(
            f'<li id="step-{i}"><b>{esc(s["step"])}</b>' + (f'<span>{esc(s["detail"])}</span>' if s.get("detail") else "")
            + (f' <a href="{esc(see_link(s.get("see"), ctx))}">{esc(SEE_LABEL.get(s.get("see"), "open"))} →</a>' if see_link(s.get("see"), ctx) else "")
            + "</li>" for i, s in enumerate(steps, 1)) + "</ol>")
    order = [r for r in b.get("reading_order") or [] if r["source_id"] in by_sid]
    if order:
        toc.append('<li><a href="#reading-order">Where to start</a></li>')
        body.append('<h2 id="reading-order">Where to start</h2><ol>' + "".join(
            f'<li><a href="docs/{esc(by_sid[r["source_id"]]["_file"])}">{esc(short_title(by_sid[r["source_id"]]))}</a>'
            f'{" — " + esc(r["why"]) if r.get("why") else ""}</li>' for r in order) + "</ol>")
    req = cfg.get("request") or {}
    slip = (f'<div class="slip" id="request"><span class="label">Routing slip</span><p>Resources suggested by '
            f'<b>{esc(req.get("by") or "")}</b>:</p><blockquote>“{esc(req["quote"])}”</blockquote></div>' if req.get("quote") else "")
    memo = (f'<article class="sheet memo">{CLIP}<div class="memo-head"><b>TO</b><span>{esc(cfg["prepared_for"])}</span>'
            f'<b>FROM</b><span>{esc(cfg.get("prepared_by") or "")}</span><b>RE</b><span>{esc(cfg["title"])}</span>'
            f'<b>DATE</b><span>{esc(fmt_day(cfg.get("date")))}</span></div>'
            f'<p class="label">Memo · written by {esc(b["writer"])}</p>{"".join(body)}</article>')
    return (f'<p class="crumbs"><a href="index.html">Cover</a> › Memo</p><div class="memo-grid"><nav class="toc" aria-label="In this memo">'
            f'<span class="label">In this memo</span><ol>{"".join(toc)}</ol>{slip}</nav><div>{memo}</div></div>')


def doc_html(d: dict, ctx: dict, prev: dict | None, nxt: dict | None) -> str:
    proj, lt = ctx["proj"], ctx["lt"]
    by = " · ".join(esc(x) for x in (d.get("authors"), d.get("author_org"), d.get("date"), d.get("identifier")) if x)
    sec = next((s for s in ctx["cfg"].get("sections") or [] if s["slug"] == d.get("section")), None)
    out = [f'<p class="crumbs"><a href="../index.html#documents">← All documents</a> · '
           f'File {esc(d["_no"])} of {len(ctx["docs"])}{" · " + esc(sec["title"]) if sec else ""}</p>',
           f'<header><span class="lens-tag">{esc(KIND_LABEL.get(d.get("kind"), d.get("kind") or "Document"))}</span>'
           f'<h1>{esc(d["title"])}</h1><p class="byline">{by}</p></header>']
    main = []
    matters = d.get("for_us") or first_sentence(d.get("summary"))
    main.append(f'<div class="matters"><span class="label">What matters for {esc(proj["name"])} · {esc(d["writer"])}</span>'
                f'<p class="big">{esc(matters)}</p>'
                + (f'<span class="label">Read: {esc(d["read_this"])}</span>' if d.get("read_this") else "") + "</div>")
    parts = [(len(d.get(k) or []), lab, sid) for k, lab, sid in (
        ("key_points", "key point", "points"), ("figures", "number", "numbers"), ("dates", "date", "dates"),
        ("actors", "person or organization", "people"), ("terms", "term", "terms"), ("_exhibits", "exhibit", "exhibits"),
        ("questions", "question", "questions"), ("cites", "cited work", "cites"))]
    main.append('<nav class="jump" aria-label="In this folder">' + "".join(
        f'<a href="#{sid}">{n} {lab if n == 1 else ("people and organizations" if lab.startswith("person") else lab + "s")}</a>'
        for n, lab, sid in parts if n) + "</nav>")
    sug = d.get("suggested_by")
    if sug:
        main.append(f'<div class="slip"><span class="label">Why it is in the file · {esc(sug.get("by") or "")}</span>'
                    f'<blockquote>“{esc(sug["quote"])}”</blockquote></div>')
    partial = (f'<p class="note">The writer read the first {d["read_chars"]:,} of {d["total_chars"]:,} characters.</p>'
               if d.get("read_chars") and d.get("total_chars") and d["read_chars"] < d["total_chars"] else "")
    main.append(f'<section class="sheet"><p class="label">Summary · by {esc(d["writer"])}</p><p class="prose">{esc(d.get("summary"))}</p>{partial}</section>')
    if d.get("why_it_matters"):
        main.append(f'<h2>Why it matters for {esc(proj["name"])}</h2><div class="postits">' + "".join(
            f'<div class="postit"><span class="label">{esc(lt.get(w.get("lens"), ""))}</span><p style="margin:4px 0 0">{esc(w.get("text"))}</p>'
            f'<span class="who">— {esc(d["writer"])}</span></div>' for w in d["why_it_matters"]) + "</div>")
    if d.get("key_points"):
        main.append(f'<section class="sheet" id="points"><h2>Key points <span class="note">({len(d["key_points"])})</span></h2><ol class="points">' + "".join(
            f'<li>{f"<span class=lens-tag>{esc(lt[k["lens"]])}</span>" if k.get("lens") in lt else ""}<p style="margin:2px 0 0" class="prose">{esc(k.get("text"))}</p>'
            f'<details class="ev"><summary>The passage</summary>{quote_html(k, d)}</details></li>' for k in d["key_points"]) + "</ol></section>")

    def block(title, items, face, cls="", sid=""):
        return (f'<section class="sec" id="{sid}"><div class="sec-head"><h2>{title}</h2><span class="note">{len(items)} · open one for the '
                f'passage</span></div><div class="cards {cls}">' + "".join(card(face(it), quote_html(it, d)) for it in items) + "</div></section>")

    if d.get("figures"):
        main.append(block("Numbers", d["figures"], lambda it: f'<span class="v">{esc(it.get("value"))}</span>'
                          f'<span class="clamp3">{esc(it.get("label"))}</span>', "figs", sid="numbers"))
    if d.get("dates"):
        main.append(block("Dates", sorted(d["dates"], key=lambda it: date_key(it.get("date"))),
                          lambda it: f'<span class="meta">{esc(fmt_when(it.get("date")))}</span><span class="t clamp3">{esc(it.get("event"))}</span>', sid="dates"))
    if d.get("actors"):
        main.append(block("Who's who in it", d["actors"], lambda it: f'<span class="t">{esc(it.get("name"))}</span>'
                          f'<span class="meta">{esc(it.get("type") or "")}</span><span class="clamp3">{esc(it.get("stance"))}</span>', sid="people"))
    if d.get("terms"):
        main.append(block("Terms", d["terms"], lambda it: f'<span class="t">{esc(it.get("term"))}</span>'
                          f'<span class="clamp3">{esc(it.get("definition"))}</span>', sid="terms"))
    if d.get("_exhibits"):
        main.append('<section class="sec" id="exhibits"><h2>Exhibits</h2><div class="exhibits">' + "".join(
            f'<figure class="exhibit"><a href="../{esc(x["img"])}"><img src="../{esc(x["img"])}" alt="{esc(x["caption"])}" loading="lazy"></a>'
            f'<figcaption>{esc(x["caption"])} — page {x["page"]}</figcaption></figure>' for x in d["_exhibits"]) + "</div></section>")
    if d.get("questions"):
        main.append(f'<section class="sheet sec" id="questions"><h2>Questions for the team</h2><p class="label">by {esc(d["writer"])}</p><ul class="check">'
                    + "".join(f"<li>{esc(q)}</li>" for q in d["questions"]) + "</ul></section>")
    if d.get("cites"):
        main.append(block("What it leans on", d["cites"], lambda it: (
            f'<span class="badge {"here" if it.get("_link") else "pull"}">{"In this file" if it.get("_link") else "Not yet gathered"}</span>'
            f'<span class="t clamp3">{esc(it.get("ref"))}</span><span class="note clamp2">{esc(it.get("note") or "")}</span>'
            + (f'<a class="from" href="{esc(it["_link"])}">Open it in this file →</a>' if it.get("_link") else "")), sid="cites"))
    facts = [("Kind", KIND_LABEL.get(d.get("kind"), d.get("kind"))), ("By", d.get("author_org")), ("Date", d.get("date")),
             ("Identifier", d.get("identifier")), ("Pages", d.get("pages")), ("License", d.get("license"))]
    dlh = "".join(f"<dt>{k}</dt><dd>{esc(v)}</dd>" for k, v in facts if v)
    src = d["_src"]
    if src.get("url"):
        dlh += f'<dt>Original</dt><dd><a href="{esc(src["url"])}" target="_blank" rel="noopener">{esc(src["url"])}</a></dd>'
    if d.get("_copy"):
        dlh += f'<dt>Saved copy</dt><dd><a href="../{esc(d["_copy"])}" download>PDF, {d["_copy_size"]}</a></dd>'
    elif d.get("_copy_note"):
        dlh += f'<dt>Saved copy</dt><dd>{esc(d["_copy_note"])}</dd>'
    if src.get("fetched_at"):
        dlh += f"<dt>Saved on</dt><dd>{esc(fmt_day(src['fetched_at']))}</dd>"
    if src.get("sha256"):
        dlh += f'<dt>SHA-256</dt><dd><code>{esc(src["sha256"][:16])}…</code></dd>'
    shot = f'<figure class="shot"><img src="../{esc(d["_cover"])}" alt="First page"></figure>' if d.get("_cover") else ""
    side = f'<aside class="side">{shot}<dl class="facts">{dlh}</dl>{rel_html(d, ctx)}</aside>'
    pn = ('<nav class="prevnext">' + (f'<a href="{esc(prev["_file"])}">← {esc(short_title(prev))}</a>' if prev else "<span></span>")
          + (f'<a href="{esc(nxt["_file"])}">{esc(short_title(nxt))} →</a>' if nxt else "<span></span>") + "</nav>")
    return "".join(out) + f'<div class="docgrid"><div class="main">{"".join(main)}</div>{side}</div>{pn}'


def gathered(ctx: dict, sec: str) -> list[tuple[dict, dict]]:
    return [(it, d) for d in ctx["docs"] for it in d.get(sec) or []]


def date_key(s: str) -> str:
    m = re.match(r"(\d{4})(?:-(\d{2}))?(?:-(\d{2}))?", str(s or ""))
    if m:
        return f"{m.group(1)}-{m.group(2) or '00'}-{m.group(3) or '00'}"
    y = re.search(r"(19|20)\d\d", str(s or ""))
    return f"{y.group(0)}-99-99" if y else "9999"


def threads(ctx: dict):
    """The timeline's colour groups: the config's sections, or else the kinds of document. -> (keys, colour, of)."""
    secs = ctx["cfg"].get("sections") or []
    if secs:
        keys = [(s["slug"], s["title"]) for s in secs]
        known = {k for k, _ in keys}

        def of(d):
            return d.get("section") if d.get("section") in known else "other"
    else:
        keys = []
        for d in ctx["docs"]:
            k = d.get("kind") or "other"
            if k not in [x for x, _ in keys]:
                keys.append((k, KIND_LABEL.get(k, k.capitalize())))

        def of(d):
            return d.get("kind") or "other"
    if any(of(d) == "other" for d in ctx["docs"]) and "other" not in [k for k, _ in keys]:
        keys.append(("other", "Other documents"))
    return keys, {k: f"k{i % 6 + 1}" for i, (k, _) in enumerate(keys)}, of


def timeline_html(ctx: dict) -> str:
    keys, colour, of = threads(ctx)
    today = str(ctx["cfg"].get("date") or "")[:10] or dt.date.today().isoformat()
    evs = sorted(gathered(ctx, "dates"), key=lambda p: (date_key(p[0].get("date")), p[1]["_no"]))
    years: dict[str, list] = {}
    for it, d in evs:
        k = date_key(it.get("date"))
        years.setdefault(k[:4] if k != "9999" else "Undated", []).append((it, d, k != "9999" and k[:10] > today))
    most = max(len(v) for v in years.values())
    bars, prev = [], None
    for y, rows in years.items():
        if y != "Undated" and prev and int(y) - int(prev) > 1:
            bars.append('<span class="gap" aria-hidden="true">⋯</span>')
        counts = {k: sum(1 for _, d, _ in rows if of(d) == k) for k, _ in keys}
        segs = "".join(f'<i class="{colour[k]}" style="height:{max(3, round(n * 84 / most))}px"></i>' for k, _ in keys if (n := counts[k]))
        ahead = all(a for *_, a in rows)
        bars.append(f'<a class="bar{" ahead" if ahead else ""}" href="#y{esc(y)}" title="{esc(y)}: {len(rows)} event{"s" * (len(rows) != 1)}'
                    f'{" (ahead)" if ahead else ""}"><span class="n">{len(rows)}</span><span class="stack">{segs}</span><span class="yr">{esc(y)}</span></a>')
        if y != "Undated":
            prev = y

    def ev_card(it, d, ahead):
        face = (f'<span class="meta">{esc(fmt_when(it.get("date")))}{" · ahead" if ahead else ""}</span>'
                f'<span class="t clamp3">{esc(it.get("event"))}</span>'
                f'<span class="from"><span class="sw"></span>{esc(short_title(d))}</span>')
        more = (f'{quote_html(it, d, "docs/" + d["_file"])}<p class="note">The line above is {esc(d["writer"])}\'s; the quotation '
                f'is the document\'s.</p><p><a href="docs/{esc(d["_file"])}">Open the document\'s folder →</a></p>')
        return card(face, more, cls=f"ev-card {colour[of(d)]}{' ahead' if ahead else ''}",
                    attrs=f'data-g="{esc(of(d))}" data-w="{"ahead" if ahead else "past"}"')

    # Neighbouring years with one or two events share a row, so a sparse decade doesn't fill the page.
    blocks, run = [], []

    def flush():
        if run:
            blocks.append(([y for y, _ in run], [r for _, rows in run for r in rows]))
            run.clear()
    for y, rows in years.items():
        ahead = all(a for *_, a in rows)
        sparse = y != "Undated" and len(rows) < 3
        if run and (not sparse or all(a for *_, a in run[-1][1]) != ahead):
            flush()
        if sparse:
            run.append((y, rows))
        else:
            blocks.append(([y], rows))
    flush()
    body = "".join(
        f'<section class="year" id="y{esc(ys[0])}" data-section>{"".join(f"<span id=y{esc(y)}></span>" for y in ys[1:])}'
        f'<h2 class="yl">{esc(ys[0])}{f"<span>–{esc(ys[-1])}</span>" if len(ys) > 1 else ""}<small>{len(rows)} event{"s" * (len(rows) != 1)}'
        f'</small></h2><div class="cards">{"".join(ev_card(*r) for r in rows)}</div></section>' for ys, rows in blocks)
    n_ahead = sum(1 for rows in years.values() for *_, a in rows if a)
    real = [y for y in years if y != "Undated"]
    span = f"{real[0]} to {real[-1]}" if real else ""
    tools = (f'<div class="toolbar">{chips("g", "Thread", [("", "All")] + [(k, sw(colour[k]) + esc(t)) for k, t in keys if any(of(d) == k for _, d in evs)])}'
             + (chips("w", "When", [("", "All"), ("past", "Happened"), ("ahead", f"Ahead <i>{n_ahead}</i>")]) if n_ahead else "")
             + f'{find_box("events")}</div>')
    return (f'<div class="page-head"><div><h1>Timeline</h1><p class="lede">{len(evs)} dated events from {len({d["_source_id"] for _, d in evs})} '
            f'documents, {esc(span)}. Pick a year, a thread or a word; open an event for the passage behind it.</p></div>'
            f'<div class="stats"><div class="stat"><b>{len(evs)}</b><span>events</span></div><div class="stat"><b>{len(real)}</b>'
            f'<span>years</span></div><div class="stat"><b>{n_ahead}</b><span>still ahead</span></div></div></div>'
            f'<nav class="ribbon" aria-label="Events per year">{"".join(bars)}</nav>'
            f'<p class="tl-note">Each bar is a year, coloured by thread; faded years lie ahead of {esc(fmt_day(today))}. '
            f'Events are the writer\'s words; each opens on the document\'s own.</p>'
            f'<div class="filterable">{tools}<div class="tl">{body}</div></div>')


# Who's who: the writer's free-text kinds of actor, sorted into a few groups for scanning (the card keeps the writer's words).
ROLE_GROUPS = (
    ("local", "Local government", ("local government", "city", "county", "municipal")),
    ("lawmakers", "Lawmakers and officials", ("legislature", "legislator", "elected", "governor", "senator", "assembly")),
    ("agencies", "Agencies and regulators", ("agency", "regulator", "commissioner", "commission", "staff", "department")),
    ("research", "Researchers, consultants and advisers", ("consultant", "research", "university", "think tank", "advis",
                                                           "convener", "academic", "laboratory")),
    ("utilities", "Utilities and energy providers", ("utility", "utilities", "aggregator", "generator", "energy provider",
                                                     "electricity provider", "gas provider")),
    ("advocates", "Advocates", ("advocate", "ratepayer", "coalition", "parties", "justice", "commenter", "ngo")),
    ("market", "Industry, trades and labour", ("industry", "contractor", "middle actor", "retail", "labor", "labour", "trade",
                                               "vendor", "company", "business", "manufacturer", "installer")),
    ("community", "Communities and households", ("community", "resident", "household", "tenant", "renter", "housing", "public")),
)


def role_group(kind: str, name: str) -> str:
    """The group for an actor: by the writer's kind first, then by the name; a word matches at the start of a word."""
    for text in (str(kind).lower(), str(name).lower()):
        for slug_, _, words in ROLE_GROUPS:
            if any(re.search(r"\b" + re.escape(w), text) for w in words):
                return slug_
    return "other"


def _name_keys(name: str, long_paren: bool) -> set[str]:
    """Keys under which two spellings of one name meet: "California Energy Commission (CEC)" and "CEC" share "cec"."""
    name = re.sub(r"\s+", " ", str(name or "")).strip()
    m = re.fullmatch(r"(.*?)\s*\(([^()]*)\)", name)
    base, par = (m.group(1), m.group(2)) if m else (name, "")
    keys = {base.lower().rstrip(".")}
    # A person named with a title after it ("Siva Gunda, CEC Vice Chair"; "Siva Gunda (Vice Chair, CEC)"): the name
    # alone is a key too, when it looks like a person's name (two or three capitalised words, no acronym).
    person = re.split(r",| \(", name, maxsplit=1)[0].strip()
    if person != name and re.fullmatch(r"[A-Z][a-z'’-]+(?: [A-Z][a-z'’.-]+){1,2}", person):
        keys.add(person.lower())
    if par and (re.fullmatch(r"[A-Z][A-Za-z0-9&.\-]{1,9}", par) or (long_paren and len(par.split()) >= 2 and not re.search(r"[,;]", par))):
        keys.add(par.lower())
    return {k for k in keys if k}


def merged(items: list[tuple[dict, dict]], field: str, long_paren: bool = False) -> list[list[tuple[dict, dict]]]:
    groups: list[list] = []
    for it, d in items:
        ks = _name_keys(it.get(field), long_paren)
        hit = [g for g in groups if g[0] & ks]
        if not hit:
            groups.append([set(ks), [(it, d)]])
            continue
        first = hit[0]
        first[0].update(ks)
        first[1].append((it, d))
        for g in hit[1:]:
            first[0].update(g[0])
            first[1].extend(g[1])
            groups.remove(g)
    return [g[1] for g in groups]


def initials(name: str) -> str:
    acr = re.search(r"\(([A-Z][A-Za-z0-9&]{1,4})\)", name)
    if acr:
        return acr.group(1)
    if re.fullmatch(r"[A-Z0-9&]{2,5}", name.strip()):
        return name.strip()
    words = [w for w in re.findall(r"[A-Za-z][\w'&-]*", re.sub(r"\(.*?\)", "", name))
             if w.lower() not in ("the", "of", "and", "for", "a", "an")]
    return "".join(w[0] for w in words[:2]).upper() or "?"


def people_html(ctx: dict) -> str:
    groups = merged(gathered(ctx, "actors"), "name")
    names_ = {s: t for s, t, _ in ROLE_GROUPS} | {"other": "Others"}
    titles = {s: names_[s] for s in ("agencies", "lawmakers", "utilities", "local", "research", "advocates", "community",
                                     "market", "other")}
    colour = {s: f"k{i % 6 + 1}" for i, s in enumerate(titles)}
    by_group: dict[str, list] = {}
    for g in groups:
        name = max((str(it.get("name")) for it, _ in g), key=len)
        kinds = list(dict.fromkeys(str(it.get("type")) for it, _ in g if it.get("type")))
        grp = role_group(" ".join(kinds), name)
        by_group.setdefault(grp, []).append((name, kinds, g))
    order = [s for s in titles if s in by_group]

    def person(name, kinds, g, grp):
        docs = list(dict.fromkeys(d["_source_id"] for _, d in g))
        face = (f'<span class="who"><span class="av">{esc(initials(name))}</span><span><span class="t">{esc(name)}</span>'
                f'<span class="meta clamp2">{esc(" · ".join(kinds))}</span></span></span>'
                f'<span class="clamp3">{esc(g[0][0].get("stance"))}</span>'
                f'<span class="from">In {len(docs)} document{"s" * (len(docs) != 1)}</span>')
        more = '<ul class="ev">' + "".join(f'<li><p class="meta">In {esc(short_title(d))}</p><p>{esc(it.get("stance"))}</p>'
                                           f'{quote_html(it, d, "docs/" + d["_file"])}</li>' for it, d in g) + "</ul>"
        return card(face, more, cls=colour[grp], attrs=f'data-g="{grp}"')

    secs = "".join(
        f'<section class="grp-sec" data-section><h2><span class="sw {colour[s]}"></span>{esc(titles[s])} <small>{len(by_group[s])}</small></h2>'
        f'<div class="cards">' + "".join(person(*p, s) for p in sorted(by_group[s], key=lambda p: (-len({d["_source_id"] for _, d in p[2]}), p[0].lower())))
        + "</div></section>" for s in order)
    n = sum(len(v) for v in by_group.values())
    tools = (f'<div class="toolbar">{chips("g", "Show", [("", "All")] + [(s, sw(colour[s]) + f"{esc(titles[s])} <i>{len(by_group[s])}</i>") for s in order])}'
             f'{find_box("names and positions")}</div>')
    many = sum(1 for v in by_group.values() for p in v if len({d["_source_id"] for _, d in p[2]}) > 1)
    return (f'<div class="page-head"><div><h1>Who\'s who</h1><p class="lede">The people and organizations the documents name, '
            f'grouped by the part they play, most-named first. Open one for the position each document gives them, with its passage.</p></div>'
            f'<div class="stats"><div class="stat"><b>{n}</b><span>named</span></div><div class="stat"><b>{many}</b>'
            f'<span>in more than one document</span></div><div class="stat"><b>{len(order)}</b><span>groups</span></div></div></div>'
            f'<div class="filterable">{tools}{secs}</div>')


def terms_html(ctx: dict) -> str:
    groups = merged(gathered(ctx, "terms"), "term", long_paren=True)
    rows = sorted(((max((str(it.get("term")) for it, _ in g), key=len), g) for g in groups), key=lambda p: p[0].lower())
    letters: dict[str, list] = {}
    for term, g in rows:
        ch = term[:1].upper()
        letters.setdefault(ch if ch.isalpha() else "#", []).append((term, g))

    def term_card(term, g):
        docs = list(dict.fromkeys(d["_source_id"] for _, d in g))
        face = (f'<span class="t">{esc(term)}</span><span class="clamp3">{esc(g[0][0].get("definition"))}</span>'
                f'<span class="from">{esc(short_title(g[0][1])) if len(docs) == 1 else f"In {len(docs)} documents"}</span>')
        more = '<ul class="ev">' + "".join(f'<li><p class="meta">In {esc(short_title(d))}</p><p>{esc(it.get("definition"))}</p>'
                                           f'{quote_html(it, d, "docs/" + d["_file"])}</li>' for it, d in g) + "</ul>"
        return card(face, more)

    bar = "".join(f'<a href="#letter-{esc(ch)}">{esc(ch)}</a>' for ch in letters)
    # One grid for every term, the letters as markers in it, so a letter with one term leaves no gap.
    body = '<div class="cards terms-flow">' + "".join(
        f'<section class="lsec" data-section><h2 class="letter" id="letter-{esc(ch)}">{esc(ch)}</h2>'
        f'{"".join(term_card(*r) for r in rs)}</section>' for ch, rs in letters.items()) + "</div>"
    shared = sum(1 for _, g in rows if len({d["_source_id"] for _, d in g}) > 1)
    return (f'<div class="page-head"><div><h1>Terms</h1><p class="lede">Terms the documents define or lean on. The meanings are '
            f'the writer\'s words; open a term for the passage each comes from.</p></div><div class="stats"><div class="stat">'
            f'<b>{len(rows)}</b><span>terms</span></div><div class="stat"><b>{shared}</b><span>in more than one document</span></div></div></div>'
            f'<div class="filterable"><div class="toolbar"><nav class="letters" aria-label="Jump to a letter">{bar}</nav>{find_box("terms")}</div>{body}</div>')


def leads_html(ctx: dict) -> str:
    groups: dict[tuple, list] = {}
    for it, d in gathered(ctx, "cites"):
        key = ("here", it["_link"]) if it.get("_link") else ("pull", re.sub(r"[^a-z0-9]+", " ", str(it.get("ref")).lower()).strip()[:80])
        groups.setdefault(key, []).append((it, d))
    cites = sum(len(g) for g in groups.values())

    def lead(key, g):
        here = key[0] == "here"
        docs = list(dict.fromkeys(short_title(d) for _, d in g))
        ref = max((str(it.get("ref")) for it, _ in g), key=len)
        note = next((it.get("note") for it, _ in g if it.get("note")), "")
        face = (f'<span class="badge {"here" if here else "pull"}">{"In this file" if here else "Not yet gathered"}</span>'
                f'<span class="t clamp3">{esc(ref)}</span>' + (f'<span class="note clamp2">{esc(note)}</span>' if note else "")
                + f'<span class="from">Cited by {esc(", ".join(docs[:3]))}{f" and {len(docs) - 3} more" if len(docs) > 3 else ""}</span>')
        more = ((f'<p><a href="docs/{esc(key[1])}">Open it in this file →</a></p>' if here else "")
                + '<ul class="ev">' + "".join(f'<li><p class="meta">Cited in {esc(short_title(d))}</p>'
                                              + (f'<p>{esc(it["note"])}</p>' if it.get("note") else "")
                                              + f'{quote_html(it, d, "docs/" + d["_file"])}</li>' for it, d in g) + "</ul>")
        return card(face, more, cls="k1" if here else "k6", attrs=f'data-w="{key[0]}"'), len(docs)

    made = sorted((lead(k, g) + (k,) for k, g in groups.items()), key=lambda x: (x[2][0] != "pull", -x[1], x[0]))
    n_pull = sum(1 for k in groups if k[0] == "pull")
    tools = (f'<div class="toolbar">{chips("w", "Show", [("", f"All <i>{len(groups)}</i>"), ("pull", f"Not yet gathered <i>{n_pull}</i>"), ("here", f"In this file <i>{len(groups) - n_pull}</i>")])}'
             f'{find_box("cited works")}</div>')
    return (f'<div class="page-head"><div><h1>Leads</h1><p class="lede">Earlier work the documents lean on. Those not in this file '
            f'are leads to follow, most-cited first; open one for how each document cites it.</p></div>'
            f'<div class="stats"><div class="stat"><b>{cites}</b><span>citations</span></div><div class="stat"><b>{len(groups)}</b>'
            f'<span>works</span></div><div class="stat"><b>{n_pull}</b><span>not yet gathered</span></div></div></div>'
            f'<div class="filterable" data-w="{"pull" if n_pull else ""}">{tools}<div class="cards">{"".join(c for c, *_ in made)}</div></div>')


def writers_line(ctx: dict) -> str:
    """Who wrote the words: "The writer here: Hermes, using Claude (Opus 5.5)." when a label stands in for the writer."""
    pairs = sorted({(w["writer"], w["_writer"]) for w in ctx["docs"] + ([ctx["brief"]] if ctx["brief"] else [])})
    names = [shown if shown == real else f"{shown}, using {real}" for shown, real in pairs]
    return f"The writer{'s' if len(names) > 1 else ''} here: {'; '.join(names)}." if names else ""


def method_html(ctx: dict) -> str:
    cfg, aud, proj = ctx["cfg"], ctx["aud"], ctx["proj"]
    rows = "".join(
        f'<tr><td>{esc(d["_no"])}</td><td><a href="docs/{esc(d["_file"])}">{esc(d["title"])}</a></td>'
        f'<td>{("<a href=" + chr(34) + esc(d["_src"]["url"]) + chr(34) + " target=_blank rel=noopener>original</a>") if d["_src"].get("url") else "—"}</td>'
        f'<td>{esc(fmt_day(d["_src"].get("fetched_at")))}</td><td><code>{esc((d["_src"].get("sha256") or "")[:16])}</code></td>'
        f'<td>{esc(d["writer"])}</td></tr>' for d in ctx["docs"])
    readers = "".join(f"<li>{esc(r['name'])} — {esc(r['role'])}</li>" for r in aud["readers"])
    lenses = "".join(f"<li><b>{esc(lens['title'])}</b>: {esc(lens['description'])}</li>" for lens in proj["lenses"])
    dropped = ctx["dropped"]
    return f"""<div class="page-head"><div><h1>How this was made</h1><p class="lede">Where each document came from, who wrote
what, and how every quotation was checked.</p></div></div>
<div class="sheet prose"><h2>The method</h2><ul>
<li>Each document was saved as it was served or received, with its exact bytes, its SHA-256 hash and when and where it came from.
The saved copies are kept with the corpus{"; the PDFs are offered for download here" if any(d.get("_copy") for d in ctx["docs"]) else ""}.</li>
<li>For each document, a writer wrote a summary, notes and questions in their own words, labelled as such,
and proposed the passages behind each key point, figure, date, position, term and cited work. {esc(writers_line(ctx))}</li>
<li>A program then looked for every proposed quotation in the saved text, character for character. Only those it found are shown:
{ctx["quotes"]} quotations here{f", and {dropped} proposed quotations that were not found were left out" if dropped else ""}.
It looked again when this page was built.</li>
<li>A page number is the page as a PDF viewer counts it, which can differ from the number printed on the page.</li>
<li>The memo's findings each rest on quotations checked the same way. Relevance (●) is the writer's judgement against the project's lenses.
The groups on the Who's who page sort the writer's own descriptions of each actor; the card keeps those words.</li>
<li>Fixity is not authenticity: a matching hash shows which bytes were read, not that what they say is right.
Provenance is not truth.</li></ul></div>
<div class="sheet prose"><h2>Written for</h2><p><b>{esc(aud['name'])}</b>{", " + esc(aud['organization']) if aud.get('organization') else ""}</p><ul>{readers}</ul>
<h2 style="margin-top:18px">The lenses ({esc(proj['name'])})</h2><ul>{lenses}</ul></div>
<div class="sheet"><h2>The documents</h2><div class="scroll"><table class="ledger"><tr><th>No.</th><th>Document</th><th>Where from</th>
<th>Saved</th><th>SHA-256</th><th>Entry written by</th></tr>{rows}</table></div></div>"""


# ------------------------------------------------------------------ images and copies
def render_page(pdf: Path, page: int, dest: Path, width: int) -> bool:
    """One PDF page as a JPEG (poppler's pdftoppm), unless it is already there."""
    if dest.exists():
        return True
    dest.parent.mkdir(parents=True, exist_ok=True)
    p = subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page), "-singlefile", "-scale-to", str(width), "-jpeg",
                        "-jpegopt", "quality=82", str(pdf), str(dest.with_suffix(""))], capture_output=True, text=True)
    return p.returncode == 0 and dest.exists()


# ------------------------------------------------------------------ the Word file
_MD_SPECIAL = re.compile(r"([\\`*_{}\[\]()#+\-.!<>|~^$\"'@&=:%])")


def md(s) -> str:
    """Text pandoc reads literally: one line, every Markdown character escaped."""
    return _MD_SPECIAL.sub(r"\\\1", re.sub(r"\s+", " ", str(s or "")).strip())


def docx_markdown(ctx: dict, out: Path) -> str:
    cfg, proj, b = ctx["cfg"], ctx["proj"], ctx["brief"]
    lt = ctx["lt"]
    by_sid = {d["_source_id"]: d for d in ctx["docs"]}
    meta = {"title": cfg["title"], "subtitle": cfg.get("subtitle") or "",
            "author": [f"Prepared for {cfg['prepared_for']}"] + ([f"Prepared by {cfg['prepared_by']}"] if cfg.get("prepared_by") else []),
            "date": fmt_day(cfg.get("date"))}
    L = ["---"] + [f"{k}: {json.dumps(v, ensure_ascii=False)}" for k, v in meta.items() if v] + ["---", ""]

    def q(it, d):
        page = f", p. {it['page']}" if it.get("page") else ""
        return [f"> {md(it['quote'])}", ">", f"> — {md(short_title(d))}{md(page)}", ""]

    if b:
        L += ["# At a glance", "", f"*{md('Written by ' + b['writer'] + '. Every quotation matches the saved copy of its source, character for character.')}*", ""]
        L += [md(b.get("bottom_line") or b.get("lede")), ""]
        if b.get("themes"):
            L += ["**Findings**", ""] + [f"{i}. {md(t.get('title'))}" for i, t in enumerate(b["themes"], 1)] + [""]
        if steps_of(b):
            L += ["**Next steps**", ""] + [f"- ☐ {md(s['step'])}" for s in steps_of(b)] + [""]
        L += [f"**{md('The documents, in reading order')}**", "",
              f"| | {md('Document')} | {md('What matters for ' + proj['name'])} |", "|---|---|---|"]
        for i, d in enumerate((d for _, ds in cover_groups(ctx) for d in ds), 1):
            L += [f"| {i} | {md(short_title(d))} | {md(d.get('for_us') or first_sentence(d.get('summary')))}"
                  + (f" {md('Read: ' + d['read_this'] + '.')}" if d.get("read_this") else "") + " |"]
        L += ["", "# Cover memo", ""]
        if b.get("lede") and b.get("bottom_line") and b["lede"] != b["bottom_line"]:
            L += [md(b["lede"]), ""]
        for sec, title in (("themes", "Themes"), ("tensions", "Where the documents pull apart")):
            if b.get(sec):
                L += [f"## {title}", ""]
                for t in b[sec]:
                    L += [f"### {md(t.get('title'))}", "", md(t.get("text")), ""]
                    for ev in t["evidence"]:
                        L += q(ev, by_sid.get(ev["source_id"]) or ctx["others"][ev["source_id"]])
        if b.get("for_project"):
            L += [f"## What it means for {md(proj['name'])}", ""]
            for t in b["for_project"]:
                L += [f"**{md(lt.get(t.get('lens'), ''))}.** " + (f"*{md(t['headline'])}* " if t.get("headline") else "")
                      + md(t.get("text")), ""]
                for ev in t.get("evidence") or []:
                    L += q(ev, by_sid.get(ev["source_id"]) or ctx["others"][ev["source_id"]])
        if steps_of(b):
            L += ["## Next steps", ""] + [f"- ☐ **{md(s['step'])}**" + (f" {md(s['detail'])}" if s.get("detail") else "")
                                          for s in steps_of(b)] + [""]
        if b.get("reading_order"):
            L += ["## Where to start", ""] + [f"{i}. {md(short_title(by_sid[r['source_id']]))}" + (f" — {md(r['why'])}" if r.get("why") else "")
                                             for i, r in enumerate(b["reading_order"], 1) if r["source_id"] in by_sid] + [""]
    L += ["# The documents", ""]
    seen_sec = None
    for d in ctx["docs"]:
        if d.get("section") != seen_sec:
            seen_sec = d.get("section")
            s = next((s for s in ctx["cfg"].get("sections") or [] if s["slug"] == seen_sec), None)
            if s:
                L += [f"*{md(s['title'])}*" + (f" — {md(s['description'])}" if s.get("description") else ""), ""]
        L += [f"## {md(d['_no'])} · {md(d['title'])}", ""]
        facts = " · ".join(str(x) for x in (KIND_LABEL.get(d.get("kind"), d.get("kind")), d.get("author_org"), d.get("date"),
                                            d.get("identifier"), f"{d['pages']} pages" if d.get("pages") else None) if x)
        L += [f"*{md(facts)}*", ""]
        if d["_src"].get("url"):
            L += [f"Original: <{d['_src']['url']}>", ""]
        if d.get("suggested_by"):
            L += [f"**{md('Why it is in the file (' + (d['suggested_by'].get('by') or '') + ')')}**", ""]
            L += [f"> {md(d['suggested_by']['quote'])}", ""]
        if d.get("for_us"):
            L += [f"**{md('What matters for ' + proj['name'])}.** {md(d['for_us'])}"
                  + (f" *{md('Read: ' + d['read_this'] + '.')}*" if d.get("read_this") else ""), ""]
        L += [f"**{md('Summary, by ' + d['writer'])}.** {md(d.get('summary'))}", ""]
        if d.get("why_it_matters"):
            L += [f"**{md('Why it matters for ' + proj['name'])}**", ""]
            L += [f"- **{md(lt.get(w.get('lens'), ''))}:** {md(w.get('text'))}" for w in d["why_it_matters"]] + [""]
        if d.get("key_points"):
            L += ["### Key points", ""]
            for i, k in enumerate(d["key_points"], 1):
                L += [f"**{i}\\.** {md(k.get('text'))}", ""] + q(k, d)
        for sec, head, a, bb in (("figures", "Numbers", "label", "value"), ("dates", "Dates", "date", "event"),
                                 ("actors", "Who's who in it", "name", "stance"), ("terms", "Terms", "term", "definition")):
            if d.get(sec):
                L += [f"### {head}", ""]
                for it in d[sec]:
                    L += [f"**{md(it.get(a))}:** {md(it.get(bb))}", ""] + q(it, d)
        if d.get("questions"):
            L += ["### Questions for the team", ""] + [f"- ☐ {md(x)}" for x in d["questions"]] + [""]
        for x in d.get("_exhibits") or []:
            L += [f"![{md(x['caption'] + ' — ' + short_title(d) + ', page ' + str(x['page']))}]({(out / x['img']).as_posix()}){{width=6in}}", ""]
    leads = gathered(ctx, "cites")
    if leads:
        L += ["# Appendix: leads", "", md("Earlier work the documents lean on, each quoted from the document that cites it."), ""]
        for it, d in leads:
            L += [f"**{md(it.get('ref'))}** {md(('— ' + it['note']) if it.get('note') else '')}", ""] + q(it, d)
    L += ["# Appendix: sources and method", "",
          md("Each document was saved with its exact bytes and their SHA-256 hash. Every quotation in this file was found "
             "again in the saved text, character for character, when the file was made. Summaries, notes and the memo are "
             f"their writers' own words. {writers_line(ctx)} A page number is the page as a PDF viewer counts it, which can "
             "differ from the number printed on the page. Fixity is not authenticity: a match shows which words were read, "
             "not that they are right."), ""]
    for d in ctx["docs"]:
        bits = [x for x in (d["_src"].get("url"), f"saved {fmt_day(d['_src'].get('fetched_at'))}" if d["_src"].get("fetched_at") else None,
                            f"SHA-256 {d['_src']['sha256'][:16]}…" if d["_src"].get("sha256") else None) if x]
        L += [f"- **{md(d['_no'])}** {md(d['title'])}: {md(' · '.join(bits))}"]
    return "\n".join(L) + "\n"


FONT = '<w:rFonts w:ascii="Calibri" w:eastAsia="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/>'


def _sz(hp: int) -> str:
    return f'<w:sz w:val="{hp}"/><w:szCs w:val="{hp}"/>'


def _style(sid, name, based, ppr, rpr, extra=""):
    return (f'<w:style w:type="paragraph" w:styleId="{sid}"><w:name w:val="{name}"/>' + (f'<w:basedOn w:val="{based}"/>' if based else "")
            + extra + "<w:qFormat/>" + (f"<w:pPr>{ppr}</w:pPr>" if ppr else "") + (f"<w:rPr>{rpr}</w:rPr>" if rpr else "") + "</w:style>")


def _heading(level, size, before, color):
    return _style(f"Heading{level}", f"heading {level}", "Normal",
                  f'<w:keepNext/><w:keepLines/><w:spacing w:before="{before}" w:after="100"/><w:outlineLvl w:val="{level - 1}"/>',
                  FONT + f'<w:b/><w:color w:val="{color}"/>' + _sz(size), '<w:next w:val="BodyText"/>')


DOCX_STYLES = {
    "BodyText": _style("BodyText", "Body Text", "Normal", '<w:spacing w:before="0" w:after="140" w:line="276" w:lineRule="auto"/>', ""),
    "FirstParagraph": _style("FirstParagraph", "First Paragraph", "BodyText", "", ""),
    "Compact": _style("Compact", "Compact", "BodyText", '<w:spacing w:before="0" w:after="60"/>', ""),
    "Title": _style("Title", "Title", "Normal", '<w:spacing w:before="1200" w:after="200"/>',
                    FONT + '<w:b/><w:color w:val="3A2F1E"/>' + _sz(48)),
    "Subtitle": _style("Subtitle", "Subtitle", "Title", '<w:spacing w:before="0" w:after="400"/>',
                       '<w:b w:val="0"/><w:color w:val="6B6354"/>' + _sz(28)),
    "Author": _style("Author", "Author", "Normal", '<w:spacing w:after="60"/>', FONT + _sz(24)),
    "Date": _style("Date", "Date", "Normal", '<w:spacing w:after="600"/>', FONT + '<w:color w:val="6B6354"/>' + _sz(24)),
    "Heading1": _heading(1, 34, 480, "3A2F1E"),
    "Heading2": _heading(2, 28, 360, "3A2F1E"),
    "Heading3": _heading(3, 24, 240, "6B5530"),
    "BlockText": _style("BlockText", "Block Text", "BodyText",
                        '<w:pBdr><w:left w:val="single" w:sz="18" w:space="8" w:color="CDB47A"/></w:pBdr>'
                        '<w:spacing w:before="40" w:after="40"/><w:ind w:left="420" w:right="240"/>',
                        '<w:rFonts w:ascii="Georgia" w:hAnsi="Georgia" w:cs="Georgia"/>' + _sz(21)),
    "ImageCaption": _style("ImageCaption", "Image Caption", "BodyText", '<w:spacing w:after="240"/>',
                           '<w:i/><w:color w:val="6B6354"/>' + _sz(18)),
}
DOCX_DEFAULTS = ("<w:docDefaults><w:rPrDefault><w:rPr>" + FONT + _sz(22) +
                 '<w:lang w:val="en-US"/></w:rPr></w:rPrDefault><w:pPrDefault><w:pPr>'
                 '<w:spacing w:after="120" w:line="276" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>')


def reference_docx(dest: Path) -> Path:
    """pandoc's own reference document, restyled (Calibri, warm headings, ruled block quotes)."""
    default = dest.with_name("default-reference.docx")
    subprocess.run(["pandoc", "-o", str(default), "--print-default-data-file", "reference.docx"],
                   check=True, capture_output=True, timeout=60)
    with zipfile.ZipFile(default) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/styles.xml":
                xml = re.sub(r"(?s)<w:docDefaults>.*?</w:docDefaults>", DOCX_DEFAULTS, data.decode("utf-8"), count=1)
                for sid, block in DOCX_STYLES.items():
                    xml, n = re.subn(r'(?s)<w:style [^>]*w:styleId="%s".*?</w:style>' % sid, block, xml)
                    if n == 0:
                        xml = xml.replace("</w:styles>", block + "</w:styles>")
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    return dest


def _norm_ws(s: str) -> str:
    return re.sub(r"\s+", " ", htmlmod.unescape(s)).strip()


def docx_text(path: Path) -> str:
    """The words in a .docx, paragraph by paragraph."""
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    xml = re.sub(r"<w:tab/>|<w:br/>", " ", xml)
    paras = [re.sub(r"<[^>]+>", "", p) for p in re.findall(r"(?s)<w:p[ >].*?</w:p>", xml)]
    return "\n".join(_norm_ws(p) for p in paras)


def make_docx(ctx: dict, out: Path, dest: Path, tmp: Path) -> tuple[str, list[str]]:
    """-> (the Markdown it was made from, quotations missing from the Word file)."""
    C.require_tools("pandoc")
    text = docx_markdown(ctx, out)
    src = tmp / "dossier.md"
    src.write_text(text, encoding="utf-8")
    ref = reference_docx(tmp / "reference.docx")
    p = subprocess.run(["pandoc", str(src), "-f", "markdown-smart", "-t", "docx", "--reference-doc", str(ref),
                        "--resource-path", str(out), "-o", str(dest)], capture_output=True, text=True, timeout=300)
    if p.returncode != 0:
        raise C.Unavailable(f"pandoc failed: {p.stderr.strip()[:400]}")
    words = _norm_ws(docx_text(dest)).replace("\n", " ")
    missing = [it["quote"] for it in all_quotes(ctx) if _norm_ws(it["quote"]) not in words]
    return text, missing


def all_quotes(ctx: dict) -> list[dict]:
    out = [it for d in ctx["docs"] for sec in ITEMS for it in d.get(sec) or []]
    out += [d["suggested_by"] for d in ctx["docs"] if d.get("suggested_by")]
    if ctx["brief"]:
        b = ctx["brief"]
        out += [ev for sec, _ in BRIEF_SECTIONS for t in b.get(sec) or [] for ev in t.get("evidence") or []]
    return out


# ------------------------------------------------------------------ the spreadsheet
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def make_xlsx(ctx: dict, dest: Path) -> tuple[dict[str, int], str]:
    """-> (rows per sheet, every cell's text for the publish checks)."""
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
    wb = Workbook()
    wb.remove(wb.active)
    when = dt.datetime.fromisoformat(f"{str(ctx['cfg'].get('date'))[:10]}T00:00:00")
    wb.properties.created = wb.properties.modified = when
    wb.properties.creator = ctx["cfg"].get("prepared_by") or "upc-corpus dossier"
    wb.properties.title = ctx["cfg"]["title"]
    lt = ctx["lt"]
    head_font, head_fill = Font(bold=True), PatternFill("solid", fgColor="EBE4D5")
    wrap = Alignment(wrap_text=True, vertical="top")
    counts, text = {}, []

    def sheet(name, cols, widths, rows):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            vals = [_ILLEGAL.sub("", v)[:32000] if isinstance(v, str) else v for v in r]
            ws.append(vals)
            text.extend(str(v) for v in vals if v is not None)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "A2"
        counts[name] = len(rows)
        text.extend(cols)

    def src(it, d):
        return [short_title(d), it.get("page"), it["quote"], it.get("extraction_id")]

    docs = ctx["docs"]
    secs = {s["slug"]: s["title"] for s in ctx["cfg"].get("sections") or []}
    sheet("Documents", ["No.", "Section", "Title", f"What matters for {ctx['proj']['name']}", "Read", "Kind", "By", "Date",
                        "Identifier", "Pages", "Original URL", "Saved", "SHA-256",
                        "Suggested by", "Their words", "Entry written by", *[f"{v} (0-3)" for v in lt.values()]],
          [6, 22, 50, 60, 24, 16, 24, 12, 18, 7, 40, 12, 20, 18, 50, 22] + [12] * len(lt),
          [[d["_no"], secs.get(d.get("section"), ""), d["title"], d.get("for_us"), d.get("read_this"),
            KIND_LABEL.get(d.get("kind"), d.get("kind")), d.get("author_org"), d.get("date"),
            d.get("identifier"), d.get("pages"), d["_src"].get("url"), fmt_day(d["_src"].get("fetched_at")),
            d["_src"].get("sha256"), (d.get("suggested_by") or {}).get("by"), (d.get("suggested_by") or {}).get("quote"),
            d["writer"], *[(d.get("relevance") or {}).get(k, 0) for k in lt]] for d in docs])
    tail = ["Document", "Page", "The document says", "Passage id"]
    tw = [24, 7, 60, 18]
    sheet("Key points", ["Lens", "Point", *tail], [16, 50, *tw],
          [[lt.get(it.get("lens"), ""), it.get("text"), *src(it, d)] for it, d in gathered(ctx, "key_points")])
    sheet("Numbers", ["What", "Figure", *tail], [36, 20, *tw],
          [[it.get("label"), it.get("value"), *src(it, d)] for it, d in gathered(ctx, "figures")])
    sheet("Timeline", ["When", "What", *tail], [14, 50, *tw],
          [[it.get("date"), it.get("event"), *src(it, d)]
           for it, d in sorted(gathered(ctx, "dates"), key=lambda p: (date_key(p[0].get("date")), p[1]["_no"]))])
    sheet("Who's who", ["Name", "Type", "Position", *tail], [28, 18, 50, *tw],
          [[it.get("name"), it.get("type"), it.get("stance"), *src(it, d)]
           for it, d in sorted(gathered(ctx, "actors"), key=lambda p: str(p[0].get("name")).lower())])
    sheet("Terms", ["Term", "Meaning", *tail], [24, 50, *tw],
          [[it.get("term"), it.get("definition"), *src(it, d)]
           for it, d in sorted(gathered(ctx, "terms"), key=lambda p: str(p[0].get("term")).lower())])
    sheet("Leads", ["Work", "Note", "In this dossier", *tail], [50, 30, 14, *tw],
          [[it.get("ref"), it.get("note"), "yes" if it.get("_link") else "no", *src(it, d)] for it, d in gathered(ctx, "cites")])
    if ctx["brief"] and steps_of(ctx["brief"]):
        sheet("Next steps", ["Step", "Detail", "Written by"], [60, 90, 22],
              [[s["step"], s.get("detail"), ctx["brief"]["writer"]] for s in steps_of(ctx["brief"])])
    sheet("Questions", ["Document", "Question", "Asked by"], [30, 80, 22],
          [[short_title(d), qn, d["writer"]] for d in docs for qn in d.get("questions") or []])
    sheet("Lenses", ["Lens", "Short", "What it covers"], [40, 16, 80],
          [[lens["title"], lt[lens["slug"]], lens["description"]] for lens in ctx["proj"]["lenses"]])
    wb.save(dest)
    return counts, "\n".join(text)


# ------------------------------------------------------------------ build
def local_paths_in(texts: dict[str, str], corpus: Path) -> list[tuple[str, str]]:
    needles = ["file://", "/Users/", "/home/", str(Path.home()), str(corpus)]
    return [(name, n) for name, t in texts.items() for n in needles if n and n in t]


def order_docs(es: list[dict], cfg: dict, idx: dict) -> list[dict]:
    """By section (the config's order), then the config's document order, then relevance."""
    want = [idx[x]["source_id"] for x in cfg.get("order") or [] if x in idx]
    rank = {sid: i for i, sid in enumerate(want)}
    sec = {s["slug"]: i for i, s in enumerate(cfg.get("sections") or [])}
    return sorted(es, key=lambda e: (sec.get(e.get("section"), len(sec)), rank.get(e["_source_id"], len(rank)),
                                     -sum((e.get("relevance") or {}).values()), e["title"].lower()))


def doc_groups(ctx: dict) -> list[tuple[dict | None, list[dict]]]:
    """The documents by section, in order; one untitled group when the config has none."""
    secs = ctx["cfg"].get("sections") or []
    groups = [(s, [d for d in ctx["docs"] if d.get("section") == s["slug"]]) for s in secs]
    rest = [d for d in ctx["docs"] if d.get("section") not in {s["slug"] for s in secs}]
    if rest:
        groups.append(({"slug": "other", "title": "Other documents", "description": ""} if secs else None, rest))
    return [(s, ds) for s, ds in groups if ds]


def build(corpus: Path, out: Path, *, copies: bool = False, docx: bool = True, xlsx: bool = True,
          images: bool = True) -> dict:
    corpus, out = Path(corpus).resolve(), Path(out).resolve()
    rep = validate(corpus)
    if rep.get("errors"):
        raise SystemExit("the corpus does not validate; a dossier is built only from one that does:\n  "
                         + "\n  ".join(f"{e.get('code')}: {str(e.get('detail'))[:120]}" for e in rep["errors"][:8]))
    cfg = load_cfg(corpus)
    aud, proj = load_profile(cfg["audience"], "audience"), load_profile(cfg["project"], "project")
    es = entries(corpus)
    if not es:
        raise SystemExit("no entries yet (dossier.py add or write)")
    br_obj = current_brief(corpus)
    problems, nq = recheck(corpus, es, br_obj, cfg)
    if problems:
        raise SystemExit("refusing to build: quotations not found again in their sources:\n  " + "\n  ".join(problems[:12]))
    br = br_obj["ext"]["upc_corpus"]["dossier_brief"] if br_obj else None
    # The pages may show a writer under another name (config `writer_labels`); the records keep the real one,
    # and the Method page and the Word file's method note name both.
    shown = cfg.get("writer_labels") or {}
    for w in es + ([br] if br else []):
        w["_writer"], w["writer"] = w["writer"], shown.get(w["writer"], w["writer"])
    idx = source_index(corpus)
    docs = order_docs(es, cfg, idx)
    width = len(str(len(docs)))
    for i, d in enumerate(docs, 1):
        d["_no"] = f"{i:0{max(2, width)}d}"
        d["_src"] = idx.get(d["_source_id"]) or {}
        d["_file"] = f"{d['_no']}-{slug(short_title(d))}.html"
    in_dossier = {d["_source_id"]: d for d in docs}
    for d in docs:
        for c in d.get("cites") or []:
            target = idx.get(c.get("in_corpus") or "") or {}
            if target.get("source_id") in in_dossier:
                c["_link"] = in_dossier[target["source_id"]]["_file"]  # beside the citing page, in docs/
    # A memo quotation from a source with no entry of its own (the request note, say) still needs a name.
    others = {sid: {"title": info["title"], "short_title": info["title"][:60]} for sid, info in idx.items()
              if info["source_id"] == sid and sid not in in_dossier}
    # Fold the build inputs into one id, so an unchanged dossier rebuilds identically.
    digest = hashlib.sha256(json.dumps([cfg, aud, proj, [d["_gen"] for d in docs], (br_obj or {}).get("synthesis_id"),
                                        TOOL], sort_keys=True, default=str).encode()).hexdigest()[:12]
    ctx = {"cfg": cfg, "aud": aud, "proj": proj, "docs": docs, "brief": br, "lt": lens_titles(proj), "others": others,
           "build": f"{str(cfg.get('date'))[:10]}-{digest}", "quotes": nq,
           "dropped": sum(len(d.get("refused") or []) for d in docs) + len((br or {}).get("refused") or []),
           "pages": {"index.html", "method.html"} | ({"memo.html"} if br else set())}
    for sec, name in (("dates", "timeline.html"), ("actors", "people.html"), ("terms", "terms.html"), ("cites", "leads.html")):
        if gathered(ctx, sec):
            ctx["pages"].add(name)

    out.mkdir(parents=True, exist_ok=True)
    produced: set[str] = set()
    img_log: list = []
    known = {i["file"]: i for i in C.load_json(manifest_path(out), {}).get("images", []) if i.get("file")}
    for d in docs:
        src, key = d["_src"], slug(d["_src"].get("slug") or d["_source_id"], 60)
        pdf = src.get("pdf_path")
        if images and pdf and Path(pdf).exists():
            rel = f"img/{key}-cover.jpg"
            if render_page(Path(pdf), 1, out / rel, 640):
                d["_cover"] = rel
                produced.add(rel)
            d["_exhibits"] = []
            for x in d.get("exhibits") or []:
                rel = f"img/{key}-p{x['page']}.jpg"
                if render_page(Path(pdf), x["page"], out / rel, 1500):
                    d["_exhibits"].append({**x, "img": rel})
                    produced.add(rel)
        elif images and src.get("raw_path") and Path(src["raw_path"]).exists():
            og = meta_content(Path(src["raw_path"]).read_text(encoding="utf-8", errors="replace")[:400_000], "og:image", "twitter:image")
            rel = thumbnail(og, out, img_log, known)
            if rel:
                d["_cover"] = rel
                produced.add(rel)
        if copies and pdf and Path(pdf).exists():
            size = Path(pdf).stat().st_size
            if size <= MAX_COPY:
                rel = f"files/{d['_no']}-{slug(short_title(d))}.pdf"
                dest = out / rel
                dest.parent.mkdir(parents=True, exist_ok=True)
                if not dest.exists() or dest.stat().st_size != size:
                    shutil.copyfile(pdf, dest)
                d["_copy"], d["_copy_size"] = rel, size_label(size)
                produced.add(rel)
            else:
                d["_copy_note"] = f"too large to include here ({round(size / 1_048_576)} MB); read it at the original link"

    name = slug(cfg.get("file_name") or cfg["title"], 60)
    downloads = []
    texts: dict[str, str] = {}
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        rows = {}
        if docx:
            md_text, missing = make_docx(ctx, out, tmp / f"{name}.docx", tmp)
            if missing:
                raise SystemExit("refusing to build: quotations missing from the Word file:\n  "
                                 + "\n  ".join(repr(m[:90]) for m in missing[:8]))
            texts[f"{name}.docx (source)"] = md_text.replace(str(out), "")  # image paths are the build's own
            texts[f"{name}.docx"] = docx_text(tmp / f"{name}.docx")
            downloads.append((f"{name}.docx", f"Word (.docx, {size_label((tmp / f'{name}.docx').stat().st_size)})"))
        if xlsx:
            rows, cells = make_xlsx(ctx, tmp / f"{name}.xlsx")
            texts[f"{name}.xlsx"] = cells
            downloads.append((f"{name}.xlsx", f"Spreadsheet (.xlsx, {size_label((tmp / f'{name}.xlsx').stat().st_size)})"))
        ctx["downloads"] = downloads

        pages = {"index.html": page_html(cfg["title"], cover_html(ctx), ctx, "index.html"),
                 "method.html": page_html(f"Method · {cfg['title']}", method_html(ctx), ctx, "method.html")}
        if br:
            pages["memo.html"] = page_html(f"Memo · {cfg['title']}", memo_html(ctx), ctx, "memo.html")
        for i, d in enumerate(docs):
            pages[f"docs/{d['_file']}"] = page_html(f"{short_title(d)} · {cfg['title']}",
                                                    doc_html(d, ctx, docs[i - 1] if i else None, docs[i + 1] if i + 1 < len(docs) else None),
                                                    ctx, "index.html", depth=1)
        for name_, fn in (("timeline.html", timeline_html), ("people.html", people_html), ("terms.html", terms_html),
                          ("leads.html", leads_html)):
            if name_ in ctx["pages"]:
                pages[name_] = page_html(f"{PAGE_TITLE[name_]} · {cfg['title']}", fn(ctx), ctx, name_)
        texts.update(pages)

        never = list(aud.get("never_publish") or []) + list(proj.get("never_publish") or []) + list(cfg.get("suppress") or [])
        hits = suppressed_in(texts, never)
        if hits:
            raise SystemExit("refusing to build: a never-publish phrase would appear:\n  "
                             + "\n  ".join(f"{n}: {w!r}" for n, w in hits[:12]))
        leaks = local_paths_in(texts, corpus)
        if leaks:
            raise SystemExit("refusing to build: a local path would be published:\n  "
                             + "\n  ".join(f"{n}: {w!r}" for n, w in leaks[:12]))
        for rel, text in pages.items():
            (out / rel).parent.mkdir(parents=True, exist_ok=True)
            (out / rel).write_text(text, encoding="utf-8")
            produced.add(rel)
        for f, _ in downloads:
            shutil.copyfile(tmp / f, out / f)
            produced.add(f)
    for f in sorted(out.rglob("*"), reverse=True):  # anything an earlier build made that this one did not
        rel = f.relative_to(out).as_posix()
        if f.is_file() and rel not in produced:
            f.unlink()
        elif f.is_dir() and not any(f.iterdir()):
            f.rmdir()
    C.save_json(manifest_path(out), {"built": ctx["build"], **TOOL, "corpus": str(corpus), "pages": sorted(pages),
                                     "documents": [{"no": d["_no"], "source_id": d["_source_id"], "generation_id": d["_gen"]} for d in docs],
                                     "brief": (br_obj or {}).get("synthesis_id"), "quotations": nq, "images": img_log,
                                     "downloads": [f for f, _ in downloads], "rows": rows})
    C.upc("event", "--corpus", str(corpus), check=False, stdin=json.dumps({
        "activity_type": "export", "tool": TOOL["tool"], "started_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "status": "success", "inputs": {"generation_ids": [d["_gen"] for d in docs]},
        "notes": f"dossier: {len(docs)} documents, {nq} quotations -> {out.name}"}))
    return {"documents": len(docs), "quotations": nq, "pages": len(pages), "downloads": [f for f, _ in downloads],
            "rows": rows, "covers": sum(bool(d.get("_cover")) for d in docs),
            "exhibits": sum(len(d.get("_exhibits") or []) for d in docs),
            "copies": sum(bool(d.get("_copy")) for d in docs), "out": str(out), "build": ctx["build"]}


PAGE_TITLE = {"timeline.html": "Timeline", "people.html": "Who's who", "terms.html": "Terms", "leads.html": "Leads"}


def size_label(n: int) -> str:
    return f"{n / 1_048_576:.1f} MB" if n >= 1_048_576 else f"{max(1, round(n / 1024))} KB"


def manifest_path(out: Path) -> Path:
    """The build record lives beside the site, not in it: it names local paths."""
    return out.parent / f"{out.name}.build.json"


# ------------------------------------------------------------------ main
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("init", help="write or update the dossier's config")
    p.add_argument("--corpus", required=True)
    for f in ("audience", "project", "title", "subtitle", "prepared-for", "prepared-by", "date", "stamp", "file-no",
              "file-name", "order", "request-source", "request-quote", "request-by", "request-date"):
        p.add_argument(f"--{f}")
    p.add_argument("--suppress", action="append", help="a phrase that must never be published (repeatable)")
    p.add_argument("--show-writer-as", action="append", metavar="WRITER=LABEL",
                   help="show a writer under another name on the pages (repeatable; WRITER= drops the label). "
                        "The records keep the real writer; the Method page names both")
    p.add_argument("--sections", help="a JSON file: [{slug, title, description}], the groups documents are shown in")
    for name in ("add", "brief"):
        p = sub.add_parser(name, help=f"store {'an entry' if name == 'add' else 'the cover memo'} written elsewhere")
        p.add_argument("--corpus", required=True)
        p.add_argument("--from", dest="src", required=True, help="the JSON file (references/dossier.md)")
        p.add_argument("--writer", required=True, help="whose words these are, as the pages name them")
        p.add_argument("--model", help="the model's name, when a model wrote it")
        p.add_argument("--method", default="model", help="model, or manual for a person")
        p.add_argument("--plan", action="store_true", help="check the quotations, write nothing")
        if name == "add":
            p.add_argument("--replace", action="store_true", help="rewrite an entry made from the same passages")
    p = sub.add_parser("write", help="have the configured model write the entries")
    p.add_argument("--corpus", required=True)
    p.add_argument("--source")
    p.add_argument("--redo", action="store_true")
    p.add_argument("--plan", action="store_true")
    p.add_argument("--max-chars", type=int, default=48000)
    p = sub.add_parser("build", help="the site, the Word file and the spreadsheet")
    p.add_argument("--corpus", required=True)
    p.add_argument("--out", required=True)
    p.add_argument("--copies", action="store_true", help="offer each PDF (up to 25 MiB) for download")
    p.add_argument("--no-docx", action="store_true")
    p.add_argument("--no-xlsx", action="store_true")
    p.add_argument("--no-images", action="store_true")
    p = sub.add_parser("check", help="find every quotation again; write nothing")
    p.add_argument("--corpus", required=True)
    a = ap.parse_args()
    corpus = Path(a.corpus).expanduser().resolve()
    if a.cmd == "init":
        return cmd_init(corpus, a)
    if a.cmd == "add":
        entry = C.load_json(a.src, None)
        if not isinstance(entry, dict):
            raise SystemExit(f"{a.src}: not a JSON object")
        rep = add_entry(corpus, entry, writer=a.writer, model=a.model, method=a.method, replace=a.replace, plan=a.plan)
        print_entry_report(rep)
        if not a.plan:
            C.upc("regen", str(corpus), check=False)
        return 0
    if a.cmd == "brief":
        b = C.load_json(a.src, None)
        if not isinstance(b, dict):
            raise SystemExit(f"{a.src}: not a JSON object")
        rep = add_brief(corpus, b, writer=a.writer, model=a.model, method=a.method, plan=a.plan)
        C.counts_header("Brief", [("status", rep.get("status")), ("themes", rep["themes"]), ("tensions", rep["tensions"]),
                                  ("for the project", rep["for_project"]), ("quotations", rep["quotations"]),
                                  ("not found", len(rep["refused"]))])
        for r in rep["refused"]:
            print(f"  NOT FOUND {r['where']} ({r['status']}): {r.get('quote', '')[:110]!r}")
        if not a.plan:
            C.upc("regen", str(corpus), check=False)
        return 0
    if a.cmd == "write":
        return cmd_write(corpus, a.source, a.redo, a.plan, a.max_chars)
    if a.cmd == "check":
        return cmd_check(corpus)
    t0 = time.time()
    res = build(corpus, Path(a.out).expanduser(), copies=a.copies, docx=not a.no_docx, xlsx=not a.no_xlsx,
                images=not a.no_images)
    C.counts_header("Dossier", [("documents", res["documents"]), ("quotations, found again", res["quotations"]),
                                ("pages", res["pages"]), ("first-page images", res["covers"]), ("exhibits", res["exhibits"]),
                                ("PDF copies", res["copies"]), ("downloads", ", ".join(res["downloads"]) or "none"),
                                ("spreadsheet rows", ", ".join(f"{k} {v}" for k, v in res["rows"].items()) or "—"),
                                ("site", res["out"]), ("build", res["build"]), ("seconds", round(time.time() - t0, 1))])
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except C.Unavailable as e:
        C.die(f"UNAVAILABLE: {e}")
