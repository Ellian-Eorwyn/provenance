#!/usr/bin/env python3
"""One row per source: its citation, links, codes and annotation, as a spreadsheet and a Word file.

    python3 sources_table.py --corpus <dir> --site-base https://example.org/library \\
        [--front front.md] [--coders scheme=coder,…] [--annotator <coder>] [--title T] \\
        --xlsx <out.xlsx> --docx <out.docx>

Each row (and each Word entry) has:

  * the citation (authors, year, title, venue), the DOI or URL, and the paper's page
    on the published site: `<site-base>/sources/<slug>.html`, where the slug is the
    source's folder name, as the site names its pages;
  * the codes coded on the paper as a whole, one column per scheme. `--coders
    scheme=coder` picks whose codes a scheme shows; a scheme left out shows each of
    its coders in a column of its own. The evidence group (a grade, a citation tier)
    comes last;
  * the newest annotation (`annotate.py`; `--annotator` picks a coder), field by
    field, with each field's checked quotation. The Word file links every passage an
    annotation cites to that passage on the site.

The Word file is the front matter (`--front`, Markdown: methods, search strings, flow
diagram, appendices), then one compact entry per source, made with pandoc and the
dossier's reference document. Before anything is kept, the outputs are read back:
every quotation must equal its stored passage and be found at that passage's place in
the paper, and must appear in the Word file word for word; no local path may appear
in either file. When a check fails, both files are removed and it says why.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import quote

sys.path.insert(0, str(Path(__file__).parent))
import common as C  # noqa: E402
import corpus_view as V  # noqa: E402
import dossier as D  # noqa: E402

EXT_MARK = re.compile(r"\s?\[(ext-[0-9a-f]{12})\]")
URI_SAFE = "-_.!~*'()"  # what encodeURIComponent leaves alone, as the site names its pages
ANNOTATION = "x-annotation"


def citation(s: dict) -> dict:
    bib = s.get("bibliographic") or {}
    names = [", ".join(x for x in (a.get("family") or a.get("literal") or "",
                                   " ".join(f"{g[0]}." for g in re.split(r"[\s.]+", a.get("given") or "") if g)) if x)
             for a in bib.get("authors") or bib.get("author") or []]
    names = [n for n in names if n]  # APA's order: "Family, G. M., Family, G., & Family, G."
    authors = (", ".join(names[:-1]) + ", & " + names[-1] if 1 < len(names) <= 5
               else f"{names[0]}, et al." if len(names) > 5 else "".join(names))
    year = str((((bib.get("issued") or {}).get("date_parts") or [[""]])[0] or [""])[0] or bib.get("year") or "")
    title = s.get("title") or bib.get("title") or s["source_id"]
    venue = bib.get("container_title") or bib.get("container-title") or bib.get("publisher") or ""
    doi = bib.get("doi") or next((i["value"] for i in s.get("identifiers") or [] if i.get("scheme") == "doi"), "")
    ret = s.get("retrieval") or {}
    url = f"https://doi.org/{doi}" if doi else bib.get("url") or ret.get("final_url") or ret.get("original_url") or ""
    text = " ".join(x for x in (f"{authors} ({year or 'n.d.'})." if authors else f"({year or 'n.d.'}).",
                                title.strip() + ("" if title.strip()[-1:] in ".?!" else "."),
                                f"{venue}." if venue else "") if x)
    return {"authors": authors, "year": year, "title": title, "venue": venue, "doi": doi, "url": url, "text": text,
            "sort": ((names[0] if names else title).lower(), year, title.lower())}


def source_codes(corpus: Path, books_by_id: dict) -> dict[str, dict[tuple[str, str], list[str]]]:
    """source_id -> {(scheme slug, coder): code labels} for codings of the paper as a whole."""
    out: dict = {}
    for f in sorted(corpus.glob("codings/*/items.jsonl")):
        for c in V.read_jsonl(f):
            t, b = c.get("target") or {}, books_by_id.get(c.get("codebook_ref"))
            if t.get("kind") != "source" or not b or (c.get("status") or "active") != "active":
                continue
            val = c.get("code") if c.get("code") is not None else c.get("value")
            label = next((x.get("label") or x["code"] for x in b.get("codes", []) if x["code"] == val), str(val))
            cell = out.setdefault(t["id"], {}).setdefault((b["slug"], c["coder"]), [])
            if label not in cell:
                cell.append(label)
    return out


def scheme_columns(codes: dict, books: dict, chosen: dict[str, str]) -> list[tuple[str, str, str]]:
    """(slug, coder, header) per column: the named coder, else every coder of the scheme."""
    by_slug: dict[str, set] = {}
    for per in codes.values():
        for slug, coder in per:
            by_slug.setdefault(slug, set()).add(coder)
    for slug in chosen:
        if slug not in books:
            C.die(f"--coders: {slug} is not a codebook here ({', '.join(sorted(books))})")
    evidence = lambda s: ((books[s].get("ext") or {}).get("upc-site") or {}).get("group") == "evidence"  # noqa: E731
    cols = []
    for slug in sorted(by_slug, key=lambda s: (evidence(s), books[s].get("title") or s)):
        coders = [chosen[slug]] if slug in chosen else sorted(by_slug[slug])
        for coder in coders:
            title = books[slug].get("title") or slug
            cols.append((slug, coder, title if len(coders) == 1 else f"{title} ({coder})"))
    return cols


def annotations(corpus: Path, coder: str | None) -> dict[str, dict]:
    gens = [g for g in V.generations(corpus, ANNOTATION) if not coder or (
        ((g.get("ext") or {}).get("upc-corpus") or {}).get("annotation") or {}).get("coder") == coder]
    return {g["_source_id"]: g for g in V.latest_per_source(gens)}


def place_of(e: dict) -> str:
    """Where the quotation is: "p. 12" in a PDF, "§3.2 ¶4" in a full-text XML (paginate.py), else ""."""
    return C.place_of(e)


def plain(text: str) -> str:
    return EXT_MARK.sub("", text or "").strip()


def build_rows(corpus: Path, args, books: dict, exts: dict) -> tuple[list[dict], list, list[dict]]:
    srcs = V.sources(corpus)
    codes = source_codes(corpus, {b["codebook_id"]: b for b in books.values()})
    chosen = dict((k.strip(), v.strip()) for k, v in (x.split("=", 1) for x in args.coders.split(",") if "=" in x))
    cols = scheme_columns(codes, books, chosen)
    anns = annotations(corpus, args.annotator)
    fields: list[dict] = []
    for g in anns.values():
        for f in g["ext"]["upc-corpus"]["annotation"]["fields"]:
            if f["key"] not in {x["key"] for x in fields}:
                fields.append({"key": f["key"], "label": f["label"]})
    rows = []
    for sid, s in srcs.items():
        cite = citation(s)
        row = {"sid": sid, "slug": s["_dir"].name, **cite,
               "site": f"{args.site_base.rstrip('/')}/sources/{quote(s['_dir'].name, safe=URI_SAFE)}.html",
               "codes": [(h, codes.get(sid, {}).get((slug, coder), [])) for slug, coder, h in cols], "fields": {}}
        g = anns.get(sid)
        if g:
            a = g["ext"]["upc-corpus"]["annotation"]
            row["by"] = f"{a['coder']} (model-written: {a.get('model') or 'a model'})"
            for f in a["fields"]:
                q = f.get("quote")
                if q:
                    e = V.check_extraction(corpus, q["extraction_id"], exts)  # found at its place, or SystemExit
                    if e["direct_quote"] != q["text"]:
                        raise SystemExit(f"{row['slug']}: the {f['label']} quotation differs from its passage {q['extraction_id']}")
                    q = {**q, "place": place_of(e)}
                body = ("; ".join(f.get("code_labels") or []) or "Not reported") + "." + (
                    f" {f['note']}" if f.get("note") else "") if f["kind"] == "code" else f.get("text") or ""
                row["fields"][f["key"]] = {"label": f["label"], "body": body, "quote": q}
        rows.append(row)
    rows.sort(key=lambda r: r["sort"])
    return rows, cols, fields


def write_xlsx(rows: list[dict], cols: list, fields: list[dict], dest: Path, title: str) -> str:
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415
    wb = Workbook()
    ws = wb.active
    ws.title = "Sources"
    wb.properties.title, wb.properties.creator = title, "upc-corpus sources_table"
    wb.properties.created = wb.properties.modified = dt.datetime.combine(dt.date.today(), dt.time())
    schemes = [h for _, _, h in cols]
    head = ["No.", "Citation", "Authors", "Year", "Title", "Venue", "DOI", "URL", "Site page", *schemes,
            *[f"{f['label']} (annotation)" if f["label"] in schemes else f["label"] for f in fields], "Annotation by"]
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="F2E2B6")
    text = list(head)
    for n, r in enumerate(rows, 1):
        cells = [n, r["text"], r["authors"], r["year"], r["title"], r["venue"], r["doi"], r["url"], r["site"],
                 *["; ".join(v) for _, v in r["codes"]]]
        for f in fields:
            fr = r["fields"].get(f["key"])
            q = (fr or {}).get("quote")
            cells.append("" if not fr else plain(fr["body"]) + (
                f" “{q['text']}”" + (f" ({q['place']})" if q.get("place") else "") if q else ""))
        cells.append(r.get("by", ""))
        cells = [D._ILLEGAL.sub("", v)[:32000] if isinstance(v, str) else v for v in cells]
        ws.append(cells)
        text += [str(v) for v in cells]
        for c, v in zip(ws[ws.max_row], cells):
            if isinstance(v, str) and v.startswith("="):  # text, never a formula
                c.data_type = "s"
        for k in (8, 9):  # URL, site page
            if cells[k - 1]:
                ws.cell(ws.max_row, k).hyperlink = cells[k - 1]
    widths = [5, 60, 28, 6, 50, 28, 24, 30, 30] + [22] * len(cols) + [50] * len(fields) + [28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(dest)
    return "\n".join(text)


def esc(s: str) -> str:
    """Text pandoc reads literally (dossier's rule, keeping the spaces at its ends)."""
    return D._MD_SPECIAL.sub(r"\\\1", re.sub(r"\s+", " ", s or ""))


def entry_markdown(n: int, r: dict, fields: list[dict]) -> list[str]:
    L = [f"## {n}\\. {esc(r['text'])}", ""]
    links = " · ".join(x for x in (f"*{esc(r['venue'])}*" if r["venue"] else "", f"<{r['url']}>" if r["url"] else "",
                                   f"[On the site]({r['site']})") if x)
    L += [links, ""]
    codes = " · ".join(f"*{esc(h)}:* {esc('; '.join(v))}" for h, v in r["codes"] if v)
    if codes:
        L += [codes, ""]
    if not r["fields"]:
        return L + ["*No annotation yet.*", ""]
    k = 0
    for f in fields:
        fr = r["fields"].get(f["key"])
        if not fr:
            continue
        parts = EXT_MARK.split(fr["body"])
        body = esc(parts[0])
        for xid, rest in zip(parts[1::2], parts[2::2]):
            k += 1
            body += f" [\\[{k}\\]]({r['site']}#ex={xid})" + esc(rest)
        q = fr.get("quote")
        if q:
            body += f" “{esc(q['text'])}”" + (f" ({esc(q['place'])})" if q.get("place") else "")
        L += [f"**{esc(fr['label'])}:** {body.strip()}", ""]
    return L + [f"*{esc('Annotation by ' + r['by'])}.*", ""]


def docx_markdown(rows: list[dict], fields: list[dict], front: str, title: str, heading: str) -> str:
    L = [] if front.lstrip().startswith("---") else ["---", f"title: {json.dumps(esc(title))}", f"date: {dt.date.today():%-d %B %Y}",
                                                    "---", ""]
    L += [front.rstrip(), ""] if front.strip() else []
    n_ann = sum(1 for r in rows if r["fields"])
    L += [f"# {esc(heading)}", "", f"*{esc(f'{len(rows)} sources, {n_ann} annotated.')} " + esc(
        "Each annotation was written by a model from the paper's opening and its coded passages, and is labelled so; "
        "its quotations were checked against the stored paper. The numbered links open each cited passage on the site.")
        + "*", ""]
    for n, r in enumerate(rows, 1):
        L += entry_markdown(n, r, fields)
    return "\n".join(L)


def zip_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels")))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--corpus", required=True)
    ap.add_argument("--site-base", required=True, help="where the folder site is published, e.g. https://example.org/library")
    ap.add_argument("--front", default="", help="Markdown placed before the entries (methods, search strings, appendices)")
    ap.add_argument("--coders", default="", help="scheme=coder,… : whose codes each scheme shows")
    ap.add_argument("--annotator", default="", help="whose annotations (default: the newest of any coder)")
    ap.add_argument("--title", default="", help="the Word file's title (default: the corpus title)")
    ap.add_argument("--heading", default="Annotated list", help="the heading above the entries")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--docx", required=True)
    args = ap.parse_args(argv)

    corpus = Path(args.corpus).expanduser()
    xlsx, docx = Path(args.xlsx).expanduser(), Path(args.docx).expanduser()
    C.require_tools("pandoc")
    books = {b["slug"]: b for b in (C.load_json(f) for f in sorted((corpus / "codebooks").glob("cbk-*.json"))) if b}
    exts = V.extractions(corpus)
    rows, cols, fields = build_rows(corpus, args, books, exts)
    title = args.title or (C.load_json(corpus / "corpus.json", {}) or {}).get("title") or corpus.name
    front = Path(args.front).expanduser().read_text(encoding="utf-8") if args.front else ""

    xlsx.parent.mkdir(parents=True, exist_ok=True)
    docx.parent.mkdir(parents=True, exist_ok=True)
    sheet_text = write_xlsx(rows, cols, fields, xlsx, title)
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "list.md"
        src.write_text(docx_markdown(rows, fields, front, title, args.heading), encoding="utf-8")
        # Pictures in the front matter (a flow diagram) are found beside it.
        where = str(Path(args.front).expanduser().resolve().parent) if args.front else str(corpus)
        p = subprocess.run(["pandoc", str(src), "-f", "markdown-smart", "-t", "docx", "--reference-doc",
                            str(D.reference_docx(Path(tmp) / "reference.docx")), "--resource-path", where,
                            "-o", str(docx)], capture_output=True, text=True, timeout=600)
    if p.returncode != 0:
        C.die(f"pandoc failed: {p.stderr.strip()[:400]}")
    for line in p.stderr.splitlines():  # a picture it could not find, say
        print(f"  NOTE pandoc: {line.strip()[:200]}")

    # Read both back before keeping them.
    words = D._norm_ws(D.docx_text(docx)).replace("\n", " ")
    quotes = [fr["quote"] for r in rows for fr in r["fields"].values() if fr.get("quote")]
    problems = [f"not in the Word file word for word: {q['extraction_id']} “{q['text'][:60]}”"
                for q in quotes if D._norm_ws(q["text"]) not in words]
    problems += [f"a local path in {name}: {needle!r}"
                 for name, needle in D.local_paths_in({docx.name: zip_text(docx), xlsx.name: sheet_text}, corpus)]
    if problems:
        xlsx.unlink(missing_ok=True)
        docx.unlink(missing_ok=True)
        print("Refusing to keep the list:\n  " + "\n  ".join(problems[:12]))
        print("Nothing was kept.")
        return 1
    C.counts_header("Sources table", [
        ("sources", len(rows)), ("with an annotation", sum(1 for r in rows if r["fields"])),
        ("scheme columns", ", ".join(h for _, _, h in cols) or "none"),
        ("annotation fields", ", ".join(f["label"] for f in fields) or "none"),
        ("quotations checked (passage, place in the paper, Word file)", len(quotes)),
        ("local paths", "none"),
        ("spreadsheet", f"{xlsx} ({xlsx.stat().st_size / 1e3:.0f} KB)"),
        ("Word file", f"{docx} ({docx.stat().st_size / 1e3:.0f} KB)")])
    print("Sources table finished.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except C.Unavailable as exc:
        C.die(str(exc))
