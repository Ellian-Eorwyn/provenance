#!/usr/bin/env python3
"""A folder of PDFs plus a manifest -> a UPC corpus.

    python3 ingest.py --config library.ingest.json [--limit N] [--plan]

The config names where things are and how the manifest's columns map:

    {
      "corpus":   "~/Documents/Corpora/library-upc",
      "title":    "Socio-technical transitions: the literature",
      "manifest": ".../2.1_corpus_manifest.csv",
      "pdf_dir":  ".../PDFs",
      "namespace": "lib",
      "columns": {"id":"id","title":"title","authors":"first_author",
                  "year":"year","doi":"doi","filename":"pdf_filename",
                  "status":"status","note":"source"}
    }

Optional columns: "journal" (container title), and "type" (a CSL item type: report, chapter, …; default
article-journal) and "publisher", read only when the map names them.

A manifest row may name a publisher's full-text XML (`.xml`) instead of a PDF, where the publisher serves the version
of record's full text only as XML (Elsevier). It is kept as the source's original and xmltext.py makes its reading
copy, which has no page numbers; an XML that is metadata only is refused.

Two decisions worth knowing:

  * **The manifest is the bibliographic authority, not the PDF.** Measured on
    this corpus, embedded PDF metadata is wrong often enough to be useless: one
    paper lists the repository depositor as its author, several have no title.
    The DOI printed on page 1 is reliable, so it is used as a cross-check and
    the disagreement is recorded rather than silently resolved.
  * **PDFs are copied, not hardlinked.** A hardlink shares the inode, so editing
    the original would change the corpus's bytes and break every quotation
    anchored in it. Copying is the cheap way to keep the corpus true.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import add_acquired as AQ   # noqa: E402
import common as C          # noqa: E402
import pdftext              # noqa: E402
import xmltext              # noqa: E402

DOI_RE = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Za-z0-9]+\b")


def read_manifest(path: Path, cols: dict) -> list[dict]:
    rows: list[dict] = []
    if path.suffix.lower() in (".csv", ".tsv"):
        delim = "\t" if path.suffix.lower() == ".tsv" else ","
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh, delimiter=delim))
    elif path.suffix.lower() == ".json":
        data = C.load_json(path, {})
        rows = data.get("sources") or data.get("items") or (data if isinstance(data, list) else [])
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl  # noqa: PLC0415
        except ImportError:
            raise C.Unavailable("reading .xlsx needs openpyxl; export the sheet to CSV instead") from None
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        head = [str(h) if h is not None else "" for h in next(it)]
        rows = [dict(zip(head, r)) for r in it]
    else:
        raise C.Unavailable(f"unsupported manifest type {path.suffix}")
    return rows


def parse_authors(raw: str) -> list[dict]:
    """'Geels, F.; Sovacool, B.' -> CSL-shaped names."""
    out = []
    for part in re.split(r"\s*;\s*", str(raw or "").strip()):
        if not part:
            continue
        if "," in part:
            family, given = part.split(",", 1)
            out.append({"family": family.strip(), "given": given.strip()})
        else:
            bits = part.rsplit(" ", 1)
            out.append({"family": bits[-1].strip(), "given": bits[0].strip()} if len(bits) == 2
                       else {"family": part.strip()})
    return out


def is_xml(path: Path) -> bool:
    """A publisher's full-text XML (xmltext.py), listed in the manifest like a PDF."""
    return path.suffix.lower() == ".xml"


def acquisition(pdf: Path) -> tuple[tuple[dict, dict] | None, str]:
    """The pdf-harvest sidecar beside a PDF (`<file>.acq.json`) as UPC `retrieval` plus the
    extras UPC has no field for, or (None, why not). A sidecar that describes other bytes is
    not attached: the record would be false."""
    side = Path(str(pdf) + ".acq.json")
    rec = C.load_json(side) if side.exists() else None
    if not rec:
        return None, ""
    if str(rec.get("sha256", "")).removeprefix("sha256:") != hashlib.sha256(pdf.read_bytes()).hexdigest():
        return None, f"{side.name} describes other bytes than the PDF; not attached"
    return AQ.retrieval_of(rec), ""


def build_record(row: dict, cols: dict, cfg: dict, pdf: Path, txt: Path, info: dict,
                 doi_in_text: str | None, acq: tuple[dict, dict] | None = None) -> dict:
    get = lambda key, default="": str(row.get(cols.get(key, key), default) or "").strip()  # noqa: E731
    rid, title, doi = get("id"), get("title"), get("doi")
    year_raw = get("year")
    year = int(year_raw[:4]) if year_raw[:4].isdigit() else None
    authors = parse_authors(get("authors"))
    if "publisher" in cols and get("authors") and get("authors") == get("publisher"):
        authors = [{"literal": get("authors")}]  # a report by its own organisation: one name, not "given family"
    first = (authors[0].get("family") or authors[0].get("literal", "")).lower() if authors else "anon"
    slug_bits = [rid, re.sub(r"[^a-z0-9]+", "-", first).strip("-"), str(year or "")]

    # A manifest may say what each item is (CSL type) and who published it; only when its column map
    # names them, so an unrelated "type" column in an older manifest changes nothing.
    item_type = get("type") if "type" in cols else ""
    bib = {"item_type": item_type or "article-journal", "title": title}
    if authors:
        bib["authors"] = authors
    if year:
        bib["issued"] = {"date_parts": [[year]]}
    if doi:
        bib["doi"] = doi
    if get("journal"):
        bib["container_title"] = get("journal")
    if "publisher" in cols and get("publisher"):
        bib["publisher"] = get("publisher")

    # Where each bibliographic field came from, and whether the document agrees.
    manifest_name = Path(cfg["manifest"]).name
    fev = {
        "title": {"value": title, "origin": "user_provided", "evidence": manifest_name},
    }
    if doi:
        agree = bool(doi_in_text) and doi_in_text.lower().rstrip(".").startswith(doi.lower().rstrip("."))
        fev["doi"] = {
            "value": doi, "origin": "user_provided",
            "evidence": manifest_name + ("; confirmed in the document text"
                                         if agree else "; NOT found in the document text"),
            # field_evidence.confidence is a SCORE (0..1), not the enum used on
            # extractions and codings. They are different fields with the same name.
            "confidence": 0.95 if agree else 0.4,
        }

    rec = {
        "slug": "-".join(b for b in slug_bits if b),
        "title": title,
        "source_kind": "document",
        "bibliographic": bib,
        "field_evidence": fev,
        "aliases": {cfg.get("namespace", "corpus"): rid},
        "files": [
            {"path": str(pdf), "role": "original", "media_type": "text/xml", "produced_by": "import"}
            if is_xml(pdf) else
            {"path": str(pdf), "role": "document_pdf", "media_type": "application/pdf",
             "produced_by": "import", "page_count": info["pages"]},
            {"path": str(txt), "role": "text", "media_type": "text/plain",
             "produced_by": "conversion", "parent": 0},
        ],
        "provenance": {
            "produced_by": {"tool": "upc-corpus", "tool_version": "1.0.0", "method": "import"},
            "created_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        },
    }
    if doi:
        rec["identifiers"] = [{"scheme": "doi", "value": doi}]
    note = get("note")
    if note:
        rec["ext"] = {cfg.get("namespace", "corpus"): {"manifest_note": note}}
    if acq:  # where the bytes came from; a source with a URL is addressed by it (spec/06)
        rec["retrieval"] = acq[0]
        if acq[1]:
            rec.setdefault("ext", {}).setdefault(cfg.get("namespace", "corpus"), {})["acquisition"] = acq[1]
    return rec


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", required=True)
    ap.add_argument("--limit", type=int, default=0, help="ingest at most N rows (for a pilot)")
    ap.add_argument("--plan", action="store_true", help="report what would happen; write nothing")
    ap.add_argument("--force-text", action="store_true", help="re-extract text even if cached")
    args = ap.parse_args()

    cfg = C.load_json(args.config)
    if not cfg:
        C.die(f"config {args.config} is missing or not JSON")
    corpus = Path(cfg["corpus"]).expanduser()
    pdf_dir = Path(cfg["pdf_dir"]).expanduser()
    manifest = Path(cfg["manifest"]).expanduser()
    cols = cfg.get("columns", {})
    ns = cfg.get("namespace", "corpus")

    try:
        spec = C.check_compat()
        C.require_tools("pdftotext", "pdfinfo")
    except C.Unavailable as e:
        C.die(str(e))
    if not manifest.exists():
        C.die(f"manifest not found: {manifest}")
    if not pdf_dir.is_dir():
        C.die(f"pdf_dir not found: {pdf_dir}")

    rows = read_manifest(manifest, cols)
    done = C.load_json(C.state_dir(corpus) / "ingest.json", {}) if corpus.exists() else {}

    # --- plan: resolve every row to a file before touching anything ---
    todo, skipped, missing = [], [], []
    for row in rows:
        rid = str(row.get(cols.get("id", "id"), "") or "").strip()
        fname = str(row.get(cols.get("filename", "pdf_filename"), "") or "").strip()
        status = str(row.get(cols.get("status", "status"), "") or "").strip().lower()
        if not rid:
            continue
        if not fname:
            skipped.append((rid, f"no file in the manifest (status: {status or 'unknown'})"))
            continue
        pdf = pdf_dir / fname
        if not pdf.exists():
            missing.append((rid, fname))
            continue
        if rid in done and not args.force_text:
            continue
        todo.append((rid, row, pdf))
    if args.limit:
        todo = todo[:args.limit]

    C.counts_header("Ingest plan", [
        ("manifest rows", len(rows)),
        ("already in the corpus", len(done)),
        ("to ingest now", len(todo)),
        ("no PDF listed", len(skipped)),
        ("listed but missing on disk", len(missing)),
        ("corpus", str(corpus)),
        ("upc", spec),
    ])
    for rid, why in skipped[:10]:
        print(f"  skipped {rid}: {why}")
    for rid, fname in missing[:10]:
        print(f"  MISSING {rid}: {fname}")
    if args.plan:
        print("Planning only; nothing was written.")
        return 0
    if not todo:
        print("Nothing to do; the corpus is already up to date.")
        return 0

    # --- corpus skeleton ---
    C.create_corpus(corpus, spec, cfg.get("title", corpus.name), cfg.get("corpus_id"))
    txt_dir = C.state_dir(corpus, "text")

    # --- convert and describe ---
    # A row already in the corpus keeps the retrieval it was added with, so it keeps its id
    # (a source with a URL is addressed by it): re-adding it is then "exists", never a twin.
    prior = {d.get("source_id"): d.get("retrieval") for d in
             (C.load_json(f, {}) or {} for f in corpus.glob("sources/*/source.json"))} if done else {}
    records, converted, doi_mismatch, fetched, t0 = [], 0, [], 0, time.time()
    for rid, row, pdf in todo:
        txt = txt_dir / f"{rid}.txt"
        try:
            if args.force_text or not txt.exists():
                info = xmltext.convert(str(pdf)) if is_xml(pdf) else pdftext.convert(str(pdf))
                if is_xml(pdf) and not info["full_text"]:
                    print(f"  UNAVAILABLE {rid}: the XML is not a full text ({info['format']})")
                    continue
                txt.write_text(info["text"], encoding="utf-8")
                C.save_json(txt_dir / f"{rid}.info.json",
                            {k: v for k, v in info.items() if k != "text"})
                converted += 1
            else:
                info = C.load_json(txt_dir / f"{rid}.info.json", {"pages": None})
        except Exception as e:
            print(f"  UNAVAILABLE {rid}: text extraction failed ({type(e).__name__}: {e})")
            continue
        body = txt.read_text(encoding="utf-8")
        if len(body.strip()) < 500:
            print(f"  UNAVAILABLE {rid}: only {len(body.strip())} characters of text "
                  f"— the PDF is probably a scan and needs OCR")
            continue
        m = DOI_RE.search(body[:6000])
        doi_found = info.get("doi") if is_xml(pdf) else (m.group(0) if m else None)
        if rid in done:
            ret, why = prior.get(done[rid].get("source_id")), ""
            acq = (ret, {}) if ret else None
        else:
            acq, why = acquisition(pdf)
            fetched += bool(acq)
        if why:
            print(f"  NOTE {rid}: {why}")
        rec = build_record(row, cols, cfg, pdf, txt, info, doi_found, acq)
        declared = (rec.get("bibliographic") or {}).get("doi")
        if declared and doi_found and not doi_found.lower().startswith(declared.lower()):
            doi_mismatch.append((rid, declared, doi_found))
        elif declared and not doi_found:
            doi_mismatch.append((rid, declared, "(no DOI printed in the text)"))
        records.append((rid, rec))

    if not records:
        C.die("no rows could be converted")

    # --- one batched write, then regenerate and validate ---
    res = C.upc_jsonl("add", "source", "--corpus", str(corpus), "--batch", "--tool", "upc-corpus",
                      records=[r for _, r in records])
    tally = res.get("tally", {}) if isinstance(res, dict) else {}
    for (rid, _), out in zip(records, res.get("results", [])):
        if out.get("status") in ("added", "exists"):
            done[rid] = {"source_id": out.get("source_id"), "slug": out.get("slug"),
                         "status": out["status"]}
        else:
            print(f"  UNAVAILABLE {rid}: {out.get('detail', 'add failed')}")
    C.save_json(C.state_dir(corpus) / "ingest.json", done)

    C.upc("regen", str(corpus))
    rep = C.upc("validate", str(corpus))
    secs = round(time.time() - t0, 1)
    C.log_timing(corpus, {"stage": "ingest", "n": len(records), "seconds": secs})

    C.counts_header("Ingest result", [
        ("sources added", tally.get("added", 0)),
        ("already present", tally.get("exists", 0)),
        ("failed", tally.get("error", 0)),
        ("PDFs converted to text", converted),
        ("with a fetch record (<pdf>.acq.json)", fetched),
        ("corpus level", f"{rep.get('status')} {rep.get('level')}"),
        ("validation errors", len(rep.get("errors", []))),
        ("DOI cross-check disagreements", len(doi_mismatch)),
        ("seconds", secs),
    ])
    for rid, declared, found in doi_mismatch[:10]:
        print(f"  DOI check {rid}: manifest {declared} / document {found}")
    if doi_mismatch:
        print("  (recorded on each source as field evidence; nothing was overwritten)")
    for e in rep.get("errors", [])[:8]:
        print(f"  ERROR {e.get('code')}: {str(e.get('detail'))[:100]}")
    print("Ingest finished.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except C.Unavailable as exc:
        C.die(str(exc))
