"""sources_table.py: one row per source as a spreadsheet and a Word file, read back before
they are kept. Built on annotate.py's test corpus (two saved pages, one annotated with
checked quotations). Skipped without the upc CLI, pandoc or openpyxl.
"""

import json
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

os.environ["UPC_CORPUS_CONFIG"] = str(Path(__file__).with_name("no-such-config.json"))  # never this machine's config
sys.path.insert(0, str(Path(__file__).parent))
HERE = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(HERE))
import common as C  # noqa: E402
import sources_table as T  # noqa: E402
import test_annotate as TA  # noqa: E402

try:
    import openpyxl
except ImportError:
    openpyxl = None

SITE = "https://example.org/library"
FRONT = "# Methods\n\nWe searched two databases and screened every record twice.\n"


class Citations(unittest.TestCase):
    def test_a_citation_reads_like_one(self):
        s = {"source_id": "src-1", "title": "Why do homes electrify?",
             "bibliographic": {"authors": [{"family": "Rivera", "given": "Ana María"}, {"family": "Okafor", "given": "B."}],
                               "issued": {"date_parts": [[2024]]}, "container_title": "Example Journal", "doi": "10.1000/x"}}
        c = T.citation(s)
        self.assertEqual(c["text"], "Rivera, A. M., & Okafor, B. (2024). Why do homes electrify? Example Journal.")
        self.assertEqual(c["url"], "https://doi.org/10.1000/x")
        many = {"source_id": "src-2", "title": "T", "bibliographic": {"authors": [{"family": f"A{i}"} for i in range(7)]},
                "retrieval": {"original_url": "https://example.org/t"}}
        self.assertEqual((T.citation(many)["authors"], T.citation(many)["url"], T.citation(many)["year"]),
                         ("A0, et al.", "https://example.org/t", ""))
        signed = {"source_id": "src-3", "title": "A report", "retrieval": {
            "original_url": "https://example.org/documents/12/report.pdf",
            "final_url": "https://bucket.s3.amazonaws.com/report.pdf?AWSAccessKeyId=AKIA0&Signature=x&x-amz-security-token=y"}}
        self.assertEqual(T.citation(signed)["url"], "https://example.org/documents/12/report.pdf")

    def test_scheme_columns_put_evidence_last_and_split_coders(self):
        books = {"grade": {"slug": "grade", "title": "Evidence grade", "ext": {"upc-site": {"group": "evidence"}}},
                 "setting": {"slug": "setting", "title": "Setting"}}
        codes = {"s1": {("grade", "ens"): ["High"], ("setting", "a"): ["X"], ("setting", "b"): ["Y"]}}
        self.assertEqual([h for _, _, h in T.scheme_columns(codes, books, {})],
                         ["Setting (a)", "Setting (b)", "Evidence grade"])
        self.assertEqual([h for _, _, h in T.scheme_columns(codes, books, {"setting": "b"})], ["Setting", "Evidence grade"])
        self.assertEqual([h for _, _, h in T.scheme_columns(codes, books, {"setting": "b"}, ["setting"])], ["Setting"])


@unittest.skipUnless(TA.upc_ok() and shutil.which("pandoc") and openpyxl, "needs the upc CLI, pandoc and openpyxl")
class Table(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.fx = TA.make_corpus(cls.tmp)
        with mock.patch.object(TA.A.P, "chat", TA.stub_chat([])):
            TA.A.main(["--corpus", str(cls.fx["corpus"]), "--fields", str(cls.fx["spec"]), "--coder", "t-annot-v1",
                       "--model", "chat"])
        cls.ann = next(json.loads(g.read_text()) for g in cls.fx["corpus"].glob("sources/*/generated/gen-*.json")
                       if json.loads(g.read_text())["provenance"]["derived_from"]["source_ids"][0] == cls.fx["a"]["source_id"])

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, front: str = FRONT, *extra) -> tuple[int, Path, Path]:
        out = Path(tempfile.mkdtemp(dir=self.tmp))
        f = out / "front.md"
        f.write_text(front)
        xlsx, docx = out / "list.xlsx", out / "list.docx"
        code = T.main(["--corpus", str(self.fx["corpus"]), "--site-base", SITE + "/", "--front", str(f),
                       "--xlsx", str(xlsx), "--docx", str(docx), *extra])
        return code, xlsx, docx

    def test_the_spreadsheet_and_the_word_file(self):
        code, xlsx, docx = self.build(FRONT, "--coders", "setting=t-screen")
        self.assertEqual(code, 0)
        ws = openpyxl.load_workbook(xlsx)["Sources"]
        head = [c.value for c in ws[1]]
        # the base columns, the scheme (t-screen's codes only), the annotation's fields, who wrote it
        # a code field is its codes (to filter on) and its note; a text field is one column
        self.assertEqual(head, ["No.", "Citation", "Authors", "Year", "Title", "Venue", "DOI", "URL", "Site page",
                                "Setting", "Setting (annotation)", "Setting (annotation): note", "Technologies",
                                "Annotation by"])
        rows = {r[4]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        a, b = rows["Retrofits in one county"], rows["A survey of apartment owners"]
        slug = self.fx["a"]["slug"]
        self.assertEqual((a[8], a[7]), (f"{SITE}/sources/{slug}.html", "https://doi.org/10.1000/county"))
        self.assertEqual((a[9], b[9]), ("Single-family homes", None))
        self.assertEqual(a[10], "Single-family homes")
        self.assertTrue(a[11].endswith(f"“{TA.UNCODED}”"), a[11])
        self.assertNotIn("[ext-", a[11] + a[12])
        self.assertIn("model-written", a[13])
        self.assertEqual(ws.cell(2, 9).hyperlink.target, rows[ws.cell(2, 5).value][8])
        words = T.D.docx_text(docx)
        for want in ("Methods", "Annotated list", "Rivera, A., & Okafor, B. (2024). Retrofits in one county.", TA.UNCODED, TA.P1,
                     "Annotation by t-annot-v1 (model-written: chat)."):
            self.assertIn(want, words)
        with zipfile.ZipFile(docx) as z:
            rels = z.read("word/_rels/document.xml.rels").decode()
        self.assertIn(f"{SITE}/sources/{slug}.html#ex={self.fx['x1']}", rels)

    def test_without_coders_each_coder_has_a_column(self):
        code, xlsx, _ = self.build()
        self.assertEqual(code, 0)
        head = [c.value for c in openpyxl.load_workbook(xlsx)["Sources"][1]]
        self.assertIn("Setting (t-other)", head)
        self.assertIn("Setting (t-screen)", head)

    def test_a_local_path_keeps_nothing(self):
        code, xlsx, docx = self.build(FRONT + "\nSee /Users/someone/notes/screening.xlsx for the sheet.\n")
        self.assertEqual(code, 1)
        self.assertFalse(xlsx.exists() or docx.exists())

    def test_the_back_matter_follows_the_entries(self):
        back = Path(tempfile.mkdtemp(dir=self.tmp)) / "back.md"
        back.write_text("# Appendix: not retrieved\n\nA report behind a paywall.\n")
        code, _, docx = self.build(FRONT, "--back", str(back))
        self.assertEqual(code, 0)
        words = T.D.docx_text(docx)
        self.assertLess(words.index("Retrofits in one county"), words.index("Appendix: not retrieved"))
        self.assertLess(words.index("A survey of apartment owners"), words.index("A report behind a paywall."))
        back.write_text("# Appendix\n\nThe list is in /Users/someone/notes/missing.csv.\n")
        code, xlsx, docx = self.build(FRONT, "--back", str(back))
        self.assertEqual(code, 1)
        self.assertFalse(xlsx.exists() or docx.exists())

    def test_a_quotation_that_no_longer_matches_is_refused(self):
        gj = next(g for g in self.fx["corpus"].glob("sources/*/generated/gen-*.json")
                  if json.loads(g.read_text())["generation_id"] == self.ann["generation_id"])
        before = gj.read_text()
        try:
            g = json.loads(before)
            for f in g["ext"]["upc-corpus"]["annotation"]["fields"]:
                if f.get("quote"):
                    f["quote"]["text"] = f["quote"]["text"].replace("four hundred", "five hundred")
            gj.write_text(json.dumps(g))
            with self.assertRaises(SystemExit) as cm:
                self.build()
            self.assertIn("differs from its passage", str(cm.exception))
        finally:
            gj.write_text(before)


if __name__ == "__main__":
    unittest.main()
