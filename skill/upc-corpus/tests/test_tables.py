"""Offline tests for the tables route (spreadsheet coding). No model is called:
a fake model answers from canned JSON. One test publishes into a scratch UPC
corpus through the real `upc` CLI and is skipped when it is not installed.

    python3 -m unittest discover -s tests -v
"""

import contextlib
import csv
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

os.environ["UPC_CORPUS_CONFIG"] = str(Path(__file__).with_name("no-such-config.json"))  # never this machine's config
SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

import common as C  # noqa: E402
import tables  # noqa: E402
import tables_data as D  # noqa: E402
import tables_model as M  # noqa: E402
import tables_out as O  # noqa: E402


def write_csv(path: Path, rows: list[dict], bom: bool = False, crlf: bool = True) -> None:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0]), lineterminator="\r\n" if crlf else "\n")
    w.writeheader()
    w.writerows(rows)
    path.write_text(("﻿" if bom else "") + buf.getvalue(), encoding="utf-8", newline="")


FORUM_CODEBOOK = {
    "version": "0.1",
    "fields": {
        "T1": {"type": "text", "def": "Neutral paraphrase."},
        "T4": {"type": "single", "def": "Face-level harm.", "values": {"explicit": "a marker", "none_evident": "none", "unclear": ""}},
        "D_sev": {"type": "int", "def": "0 guidance, 1 instruction, 2 warning, 3 ban."},
        "M_moves": {"type": "multi_with_quotes", "def": "How the mod justifies.", "values": {"closure": "settles", "rule_citation": "cites"}},
        "N_norms": {"type": "multi_with_polarity", "def": "Norms.", "values": {"discussability": "what may be raised"}},
        "R_mapped": {"type": "multi", "def": "Rules covering it.", "values": "CORE"},
        "J_basis": {"type": "single", "def": "What was sanctioned.", "values": {"conduct": "", "content": "", "unclear": ""}},
    },
    "parts": {"B": {"unit": "one sanctioned post", "fields": ["T1", "T4"]},
              "C": {"unit": "one moderator action", "fields": ["D_sev", "M_moves", "N_norms", "R_mapped", "J_basis"]}},
}

FORUM_UNITS = {
    "format": D.UNITS_FORMAT,
    "question": "How moderators treat beliefs.",
    "quote_blocks": {"style": "xenforo", "sources_column": "quoted_post_ids"},
    "lookups": {"CORE": {"table": "core", "code": "core_id", "label": "name", "definition": "summary", "date": "introduced"}},
    "parts": {
        "B": {"table": "posts", "id": "id",
              "where": [{"column": "id", "in": {"table": "targets", "column": "target_post_id"}}],
              "code": "S", "header": [{"column": "thread_title", "label": "Thread"}],
              "text": {"column": "text", "label": "POST", "speaker": "author", "date": "datetime"},
              "context": [{"table": "posts", "id": "id", "text": "text", "speaker": "author", "date": "datetime",
                           "group": "thread_id", "order": "datetime", "before": 2, "after": 0, "require_text": True,
                           "exclude": [{"column": "id", "in": {"table": "actions", "column": "action_post_id"}}]}],
              "blind": True,
              "leak_check": [{"path": [{"from": "posts.id", "to": "targets.target_post_id"},
                                       {"from": "targets.action_id", "to": "actions.action_id"}], "column": "mod_text"}],
              "evidence": {"T4": ["S"]}},
        "C": {"table": "actions", "id": "action_id", "code": "M",
              "header": [{"column": "datetime", "label": "Date", "format": "date"}],
              "text": {"column": "mod_text", "label": "MODERATOR TEXT", "speaker": "mod"},
              "related": [{"label": "SANCTIONED POST", "code": "S", "id": "id", "text": "text", "speaker": "author",
                           "path": [{"from": "actions.action_id", "to": "targets.action_id"},
                                    {"from": "targets.target_post_id", "to": "posts.id"}], "include_codes": "B"}],
              "context": [{"table": "posts", "id": "id", "text": "text", "speaker": "author", "group": "thread_id",
                           "order": "datetime", "anchor_path": [{"from": "actions.action_post_id", "to": "posts.id"}],
                           "before": 1, "after": 2, "require_text": True}],
              "evidence": {"M_moves": ["M"], "N_norms": ["M", "S", "E", "L"]},
              "int_ranges": {"D_sev": [0, 3]},
              "date_checks": [{"name": "predates", "field": "R_mapped", "lookup": "CORE", "unit_date": "datetime"}],
              "export": {"id_column": "action_id", "joined_name": "joined", "columns": ["datetime", "mod"]}},
    },
    "screens": {"posts": {"table": "posts", "id": "id", "text": "text", "speaker": "author", "date": "datetime",
                          "group": "thread_id", "order": "datetime", "window_chars": 220, "row_chars": 120, "overlap": 1,
                          "where": [{"column": "text", "not_empty": True}], "parts": ["B", "new"]}},
}


def forum_files(root: Path) -> Path:
    data = root / "data"
    data.mkdir()
    quote = "Alice said:\n\nPineapple belongs on pizza and everyone knows it.\n\nClick to expand...\n\n"
    posts = [
        {"id": "100", "thread_id": "t1", "thread_title": "Food fights", "author": "Alice", "datetime": "2020-01-01T10:00",
         "text": "Pineapple belongs on pizza and everyone knows it.", "quoted_post_ids": "", "banners": ""},
        {"id": "101", "thread_id": "t1", "thread_title": "Food fights", "author": "Bob", "datetime": "2020-01-01T11:00",
         "text": quote + "Only fools think that. People like you ruin every thread.", "quoted_post_ids": "100", "banners": "Banned"},
        {"id": "102", "thread_id": "t1", "thread_title": "Food fights", "author": "Mod", "datetime": "2020-01-01T12:00",
         "text": "", "quoted_post_ids": "", "banners": "Staff"},
        {"id": "103", "thread_id": "t1", "thread_title": "Food fights", "author": "Cara", "datetime": "2020-01-01T13:00",
         "text": "Fair enough, let us move on.\nThis thread is about toppings.", "quoted_post_ids": "", "banners": ""},
        {"id": "104", "thread_id": "t1", "thread_title": "Food fights", "author": "Dan", "datetime": "2020-01-01T14:00",
         "text": "Do not attack posters. That is the rule here and it is not up for debate.", "quoted_post_ids": "", "banners": ""},
        {"id": "105", "thread_id": "t2", "thread_title": "Other", "author": "Eve", "datetime": "2021-05-01T09:00",
         "text": "Some ordinary chat about dice.", "quoted_post_ids": "", "banners": ""},
    ]
    write_csv(data / "posts.csv", posts, bom=True)
    write_csv(data / "actions.csv", [
        {"action_id": "a102", "action_post_id": "102", "datetime": "2020-01-01T12:00", "mod": "Mod",
         "mod_text": "Calling people fools is a personal attack. This is not up for debate. Leave the thread."},
        {"action_id": "a104", "action_post_id": "104", "datetime": "2019-06-01T12:00", "mod": "Mod",
         "mod_text": "Do not attack posters. That is the rule here and it is not up for debate."},
    ])
    write_csv(data / "targets.csv", [{"action_id": "a102", "target_post_id": "101"},
                                     {"action_id": "a104", "target_post_id": "103"}])
    write_csv(data / "core.csv", [{"core_id": "R1-PA", "name": "Personal attacks", "summary": "No attacks.", "introduced": "2019-12-01"},
                                  {"core_id": "R2", "name": "Other", "summary": "Other rule.", "introduced": ""}])
    (root / "cb").mkdir()
    (root / "cb" / "codebook.json").write_text(json.dumps(FORUM_CODEBOOK), encoding="utf-8")
    return data


class FakeModel:
    """Answers llm_chat from the schema it is given: the first allowed value of
    every field and the first allowed evidence id."""

    def __init__(self):
        self.calls = []

    def __call__(self, system, user, **kw):
        self.calls.append((system, user, kw))
        schema = kw.get("schema") or {}
        if kw.get("max_tokens") == 1:
            return {"content": "OK", "seconds": 0.01}
        if "rows" in schema.get("properties", {}):
            ids = schema["properties"]["rows"]["items"]["properties"]["id"]["enum"]
            return {"content": json.dumps({"rows": [{"id": ids[0], "r": 2, "parts": ["new"], "label": "attacks", "why": "x"}]}),
                    "seconds": 0.01, "finish_reason": "stop"}
        return {"content": json.dumps(self.fill(schema)), "seconds": 0.02, "prompt_tokens": 10, "completion_tokens": 5,
                "finish_reason": "stop"}

    def fill(self, schema):
        t = schema.get("type")
        if t == "object":
            return {k: self.fill(p) for k, p in schema.get("properties", {}).items()}
        if t == "string":
            return schema["enum"][0] if "enum" in schema else "a neutral paraphrase"
        if t == "integer":
            return schema.get("minimum", 0)
        if t == "array":
            items = schema.get("items")
            if not items or schema.get("maxItems") == 0:
                return []
            return [self.fill(items)]
        return None


class RunCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.data = forum_files(self.root)
        self.run_dir = self.root / "runs" / "forum"
        self.fake = FakeModel()
        self.patches = [mock.patch.object(C, "llm_chat", self.fake)]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()
        self.tmp.cleanup()

    def cli(self, *argv):
        with mock.patch.object(sys, "argv", ["tables.py", *argv]), contextlib.redirect_stdout(io.StringIO()) as out:
            try:
                rc = tables.main()
            except SystemExit as e:
                rc = e.code
        return rc, out.getvalue()

    def init_run(self):
        rc, out = self.cli("init", "--run", str(self.run_dir), "--input", str(self.data),
                           "--codebook", str(self.root / "cb" / "codebook.json"), "--question", "q")
        self.assertEqual(rc, 0, out)
        (self.run_dir / "units.json").write_text(json.dumps(FORUM_UNITS), encoding="utf-8")
        return D.Run(self.run_dir)


class LoadAndProfileTests(RunCase):
    def test_csv_bom_crlf_and_newlines_in_cells(self):
        (t,) = D.load_table(self.data / "posts.csv")
        self.assertEqual(t.columns[0], "id")  # the BOM is not part of the first header
        self.assertIn("\n\nClick to expand...", t.rows[1]["text"])
        self.assertNotIn("\r", t.rows[1]["text"])

    def test_jsonl_and_xlsx(self):
        (self.root / "x.jsonl").write_text('{"id": "1", "tags": ["a", "b"], "n": 3}\n{"id": "2", "extra": true}\n')
        (t,) = D.load_table(self.root / "x.jsonl")
        self.assertEqual(t.columns, ["id", "tags", "n", "extra"])
        self.assertEqual(t.rows[0]["tags"], "a;b")
        self.assertEqual(t.rows[1]["extra"], "True")
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["rid", "answer"])
        wb.active.append([1, "I think so."])
        wb.save(self.root / "s.xlsx")
        (x,) = D.load_table(self.root / "s.xlsx")
        self.assertEqual(x.rows, [{"rid": "1", "answer": "I think so."}])

    def test_keys_and_links(self):
        tbs = {tb.name: tb for p in D.discover_inputs([self.data]) for tb in D.load_table(p)}
        prof = D.profile_tables(tbs)
        self.assertIn("id", prof["tables"]["posts"]["keys"])
        links = {(link["from"], link["to"]): link for link in prof["links"]}
        self.assertIn(("targets.target_post_id", "posts.id"), links)
        self.assertIn(("actions.action_post_id", "posts.id"), links)
        self.assertEqual(links[("targets.action_id", "posts.id")]["transform"], {"strip": "a", "add": ""})
        warn = D.check_roles(prof["tables"]["posts"], {"text": "text", "banners": "id", "author": "code_slot"},
                             prof["links"], "posts")
        self.assertTrue(any("banners" in w for w in warn) and any("author" in w for w in warn))


class RenderTests(RunCase):
    def test_quote_blocks_sentences_and_exact_spans(self):
        text = self_text = "Alice said:\n\nA quoted line.\n\nClick to expand...\n\nMy own view. Another point!"
        segs = D.split_quote_blocks(text, "xenforo")
        self.assertEqual([s.kind for s in segs], ["quote", "own"])
        self.assertEqual(segs[0].author, "Alice")
        self.assertEqual(D.quote_body(text, segs[0]), "A quoted line.")
        spans = D.sentences(self_text, segs[1].start, segs[1].end)
        self.assertEqual([self_text[a:b] for a, b in spans], ["My own view.", "Another point!"])

    def test_blind_unit_shows_only_the_whitelist_and_the_past(self):
        run = self.init_run()
        tbs, spec = run.tables(), run.spec()
        r = D.Renderer(tbs, spec, "B")
        units = {u[0]: u for u in r.units()}
        self.assertEqual(sorted(units), ["101", "103"])
        rb = r.render(units["101"][1], unit_id="101")
        self.assertIn("‹quotes E1›", rb.text)            # the quoted post is in the case
        self.assertNotIn("Banned", rb.text)              # banners are not whitelisted
        self.assertNotIn("personal attack", rb.text)     # the moderator's words never appear
        self.assertNotIn("Fair enough", rb.text)         # nothing after the target
        self.assertEqual(rb.evidence["S1.1"].text, "Only fools think that.")
        e = rb.evidence["S1.2"]
        cell = tbs["posts"].rows[1]["text"]
        self.assertEqual(cell[e.start:e.end], "People like you ruin every thread.")
        # a planted leak: the moderator's words inside the post itself
        rc = r.render(units["103"][1], unit_id="103")
        hidden = [row["mod_text"] for _t, row in D.resolve_path(tbs, [units["103"][1]], spec["parts"]["B"]["leak_check"][0]["path"])]
        self.assertEqual(D.leak_check(rc.text, hidden, "xenforo"), [])
        planted = rb.text + "\nDo not attack posters. That is the rule here and it is not up for debate.\n"
        self.assertTrue(D.leak_check(planted, hidden, "xenforo"))

    def test_action_shows_all_sanctioned_posts_context_and_codes(self):
        run = self.init_run()
        tbs, spec = run.tables(), run.spec()
        r = D.Renderer(tbs, spec, "C", coded={"B": {"101": {"T4": "explicit"}}})
        (uid, unit, seg), = [u for u in r.units() if u[0] == "a102"]
        rc = r.render(unit, unit_id=uid)
        self.assertIn("== SANCTIONED POST S1 · Bob", rc.text)
        self.assertIn('CODES FROM B FOR S1: {"T4": "explicit"}', rc.text)
        self.assertIn("LATER POST L1 · Cara", rc.text)
        self.assertIn("[M1.2] This is not up for debate.", rc.text)
        self.assertEqual(set(rc.speakers), {"S", "M", "L", "E"} - {"E"} | ({"E"} if "E" in rc.speakers else set()))

    def test_line_segments(self):
        text = "Rules\nRule 1: Be civil\nDo not attack people or groups.\nok\nBob said:\n\nQuoted rule text here.\n\nClick to expand..."
        segs = D._segment_rows(text, "line", 12, "^Rule \\d+", "xenforo")
        self.assertEqual([text[a:b] for a, b, _h in segs], ["Rule 1: Be civil", "Do not attack people or groups."])
        self.assertEqual(segs[1][2], "Rule 1: Be civil")
        with_q = D._segment_rows(text, "line", 12, "^Rule \\d+", "xenforo", include_quotes=True)
        self.assertIn("Quoted rule text here.", [text[a:b] for a, b, _h in with_q])


class ModelLayerTests(RunCase):
    def test_schema_restricts_evidence_to_speakers(self):
        run = self.init_run()
        tbs, cb, spec = run.tables(), run.codebook(), run.spec()
        r = D.Renderer(tbs, spec, "C")
        (uid, unit, seg), = [u for u in r.units() if u[0] == "a102"]
        rc = r.render(unit, unit_id=uid)
        schema = M.output_schema(spec, tbs, cb, "C", rc)
        m_ids = schema["properties"]["M_moves"]["items"]["properties"]["evidence"]["items"]["enum"]
        self.assertTrue(m_ids and all(i.startswith("M") for i in m_ids))
        self.assertEqual(schema["properties"]["D_sev"], {"type": "integer", "minimum": 0, "maximum": 3})
        self.assertEqual(schema["properties"]["R_mapped"]["items"]["enum"], ["R1-PA", "R2"])
        ans = {"D_sev": 9, "M_moves": [{"code": "closure", "evidence": ["M1.2", "S1.1", "X9.9"]}],
               "N_norms": [{"code": "discussability", "polarity": "enforced", "evidence": ["S1.1"]}],
               "R_mapped": ["R1-PA", "bogus"], "J_basis": "maybe"}
        fields, ev, bad = M.check_answer(ans, spec, tbs, cb, "C", rc)
        self.assertIsNone(fields["D_sev"])
        self.assertEqual(fields["J_basis"], "unclear")
        self.assertEqual(ev["M_moves"][0]["ids"], ["M1.2"])
        self.assertEqual(fields["N_norms"], ["discussability(enforced)"])
        self.assertTrue(any("S1.1 is not from M" in b for b in bad) and any("bogus" in b for b in bad))

    def test_windows_pack_rows_with_one_row_overlap(self):
        run = self.init_run()
        wins = M.screen_windows(run.tables(), run.spec()["screens"]["posts"], "xenforo")
        ids = [i for w in wins for i in w["ids"]]
        self.assertEqual(sorted(ids), sorted(set(ids)))  # every row scored once
        self.assertTrue(any("(context, already screened)" in w["text"] for w in wins))
        self.assertTrue(all(w["group"] in ("t1", "t2") for w in wins))
        self.assertNotIn("Pineapple belongs on pizza and everyone knows it.\n\nClick", "".join(w["text"] for w in wins))

    def test_statistics(self):
        self.assertEqual(M.cohens_kappa([("a", "a"), ("b", "b"), ("a", "b"), ("b", "a")]), 0.0)
        self.assertEqual(M.cohens_kappa([("a", "a"), ("b", "b")]), 1.0)
        self.assertEqual(M.krippendorff_alpha_nominal([("a", "a"), ("b", "b"), ("a", "a")]), 1.0)
        self.assertEqual(M.weighted_kappa([(0, 0), (1, 1), (2, 2)], 0, 2), 1.0)
        lo, hi = M.wilson(8, 10)
        self.assertAlmostEqual(lo, 0.49, places=2)
        self.assertAlmostEqual(hi, 0.943, places=2)


class PipelineTests(RunCase):
    def test_init_guard_code_export_markdown(self):
        # the run may not live inside an input's folder
        rc, out = self.cli("init", "--run", str(self.data / "run"), "--input", str(self.data))
        self.assertNotEqual(rc, 0)
        run = self.init_run()
        with self.assertRaises(D.GuardError):
            run.guard(self.data / "x.csv")
        rc, out = self.cli("units", "--run", str(self.run_dir), "--check")
        self.assertEqual(rc, 0, out)
        self.assertTrue((self.run_dir / "UNITS.md").exists())
        with mock.patch.object(M, "busy_seconds", return_value=0.1):
            rc, out = self.cli("code", "--run", str(self.run_dir), "--part", "C", "--ids", "a102,a104")
        self.assertEqual(rc, 0, out)
        self.assertEqual(sorted(run.coded("B")), ["101", "103"])        # dependencies coded first
        self.assertEqual(sorted(run.coded("C")), ["a102", "a104"])
        calls = len(self.fake.calls)
        with mock.patch.object(M, "busy_seconds", return_value=0.1):
            rc, out = self.cli("code", "--run", str(self.run_dir), "--part", "C", "--ids", "a102,a104")
        self.assertEqual(len(self.fake.calls), calls)                   # resumed: nothing re-coded
        rc, out = self.cli("export", "--run", str(self.run_dir), "--xlsx")
        self.assertEqual(rc, 0, out)
        raw = (self.run_dir / "exports" / "joined.csv").read_bytes()
        self.assertFalse(raw.startswith(b"\xef\xbb\xbf"))                  # find_examples.py reads plain utf-8
        rows = list(csv.DictReader(io.StringIO(raw.decode())))
        self.assertEqual({r["action_id"] for r in rows}, {"a102", "a104"})
        self.assertIn("T4", rows[0])
        by = {r["action_id"]: r for r in rows}
        self.assertEqual(by["a104"]["predates"], "R1-PA")                 # 2019 action, rule from 2019-12
        self.assertEqual(by["a102"]["predates"], "")
        ev = list(csv.DictReader(io.StringIO((self.run_dir / "exports" / "evidence.csv").read_text(encoding="utf-8"))))
        posts = {r["id"]: r for r in run.tables()["posts"].rows}
        actions = {r["action_id"]: r for r in run.tables()["actions"].rows}
        for e in ev:
            cell = (posts if e["table"] == "posts" else actions)[e["row"]][e["column"]]
            self.assertEqual(cell[int(e["start"]):int(e["end"])], e["text"])
        # stale inputs are refused
        snap = self.run_dir / "inputs" / "posts.csv"
        snap.write_text(snap.read_text(encoding="utf-8") + "\n", encoding="utf-8")
        with self.assertRaises(D.GuardError):
            D.Run(self.run_dir).tables()

    def test_markdown_headers_match_rows(self):
        self.init_run()
        with mock.patch.object(O, "LIT_EXTRACTION", Path("/nonexistent")):
            rc, out = self.cli("markdown", "--run", str(self.run_dir), "--part", "B", "--out", str(self.root / "md"))
        self.assertEqual(rc, 0, out)
        files = sorted((self.root / "md").glob("*.md"))
        self.assertEqual(len(files), 2)
        text = files[0].read_text(encoding="utf-8")
        self.assertIn('author: "Bob"', text)
        self.assertIn("Only fools think that.", text)
        self.assertNotIn("[S1.1]", text)

    def test_agree_sheet_and_ingest(self):
        run = self.init_run()
        with mock.patch.object(M, "busy_seconds", return_value=0.1):
            self.cli("code", "--run", str(self.run_dir), "--part", "C", "--ids", "a102,a104")
        rc, out = self.cli("agree", "--run", str(self.run_dir), "--part", "C", "--fields", "J_basis,D_sev",
                           "--also", "B:T4", "--n", "2")
        self.assertEqual(rc, 0, out)
        sheet = next((self.run_dir / "agree").glob("blind-C-*.csv"))
        dossier = sheet.with_suffix(".md").read_text(encoding="utf-8")
        self.assertIn("Part 1. The posts alone", dossier)
        self.assertNotIn("CODES FROM", dossier)                         # the model's codes are not shown
        rows = list(csv.DictReader(io.StringIO(sheet.read_text(encoding="utf-8"))))
        model = run.coded("C")
        for r in rows:
            r["J_basis"] = model[r["unit"]]["fields"]["J_basis"]           # agree on J_basis
            r["D_sev"] = "3"                                             # disagree on the integer
            r["B:T4"] = "explicit" if r["B unit"] else ""
        sheet.write_text(O.csv_text(rows, list(rows[0])), encoding="utf-8")
        rc, out = self.cli("agree", "--run", str(self.run_dir), "--ingest", str(sheet), "--coder", "reviewer")
        self.assertEqual(rc, 0, out)
        report = (self.run_dir / "agree" / "AGREEMENT.md").read_text(encoding="utf-8")
        self.assertIn("| J_basis | 2 | 2/2", report)
        self.assertIn("| D_sev | 2 | 0/2", report)
        self.assertTrue((self.run_dir / "coded" / "C.human-reviewer.jsonl").exists())
        self.assertTrue((self.run_dir / "coded" / "B.human-reviewer.jsonl").exists())

    def test_screen_and_draft_parts(self):
        self.init_run()
        with mock.patch.object(M, "busy_seconds", return_value=0.1):
            rc, out = self.cli("screen", "--run", str(self.run_dir), "--screen", "posts")
        self.assertEqual(rc, 0, out)
        lines = (self.run_dir / "screen" / "posts.windows.jsonl").read_text().splitlines()
        wins = [json.loads(x) for x in lines]
        self.assertTrue(all(w["status"] == "ok" for w in wins))
        rows = tables.load_screen_rows(D.Run(self.run_dir), "posts")
        self.assertEqual(set(rows), {"100", "101", "102", "103", "104", "105"} - {"102"})  # 102 has no text
        self.assertTrue((self.run_dir / "screen" / "RELEVANCE-posts.md").exists())
        calls = len(self.fake.calls)
        with mock.patch.object(M, "busy_seconds", return_value=0.1):
            self.cli("screen", "--run", str(self.run_dir), "--screen", "posts")
        self.assertEqual(len(self.fake.calls), calls)                   # nothing screened twice
        rc, out = self.cli("draft-parts", "--run", str(self.run_dir), "--screen", "posts")
        self.assertEqual(rc, 0, out)
        self.assertTrue((self.run_dir / "drafts" / "parts_draft_posts.md").exists())

    def test_review_sheet_marks_j_fields(self):
        self.init_run()
        with mock.patch.object(M, "busy_seconds", return_value=0.1):
            self.cli("code", "--run", str(self.run_dir), "--part", "C", "--ids", "a102,a104")
        rc, out = self.cli("agree", "--run", str(self.run_dir), "--part", "C", "--fields", "J_basis", "--n", "2", "--review")
        self.assertEqual(rc, 0, out)
        sheet = next((self.run_dir / "agree").glob("review-C-*.csv"))
        rows = list(csv.DictReader(io.StringIO(sheet.read_text(encoding="utf-8"))))
        rows[0]["J_basis decision"] = "accept"
        rows[1]["J_basis decision"] = "change"
        rows[1]["J_basis your value"] = "content"
        sheet.write_text(O.csv_text(rows, list(rows[0])), encoding="utf-8")
        rc, out = self.cli("agree", "--run", str(self.run_dir), "--ingest", str(sheet), "--coder", "reviewer")
        self.assertEqual(rc, 0, out)
        self.cli("export", "--run", str(self.run_dir))
        exported = list(csv.DictReader(io.StringIO((self.run_dir / "exports" / "C_coded.csv").read_text(encoding="utf-8"))))
        self.assertTrue(all(r["J_review"].startswith("reviewed by reviewer") for r in exported))

    def test_calibration_with_fake_embeddings(self):
        self.init_run()
        with mock.patch.object(M, "busy_seconds", return_value=0.1):
            self.cli("screen", "--run", str(self.run_dir), "--screen", "posts")
        vec = lambda texts: [[1.0, float(len(t) % 7)] for t in texts]  # noqa: E731
        with mock.patch.object(C, "embed_resilient", side_effect=vec):
            rc, out = self.cli("screen", "--run", str(self.run_dir), "--screen", "posts", "--calibrate", "--audit", "6")
        self.assertEqual(rc, 0, out)
        cal = (self.run_dir / "screen" / "posts.calibration.md").read_text(encoding="utf-8")
        self.assertIn("Rows re-screened one at a time", cal)
        self.assertIn("Calibration", (self.run_dir / "screen" / "RELEVANCE-posts.md").read_text(encoding="utf-8"))

    def test_pseudonyms_reach_headers_quotes_and_mentions(self):
        run = self.init_run()
        tbs, spec = run.tables(), run.spec()
        mapper, mapping = O.name_mapper(run, tbs, spec, True)
        self.assertIn("Alice", mapping)
        r = D.Renderer(tbs, spec, "C", name_map=mapper)
        (uid, unit, seg), = [u for u in r.units() if u[0] == "a102"]
        text = r.render(unit, unit_id=uid, show_ids=False).clean
        for name in ("Alice", "Bob", "Cara", "Mod"):
            self.assertNotRegex(text, rf"\b{name}\b")
        self.assertEqual(mapper("@Bob said hi to Bobby"), f"@{mapping['Bob']} said hi to Bobby")
        self.assertEqual(O.name_mapper(run, tbs, spec, True)[1], mapping)   # stable across calls (same key)

    def test_upc_tokens(self):
        self.assertEqual(D.slug_token("extends_by_inference"), "extends-by-inference")
        self.assertEqual(D.slug_token("R1-PA"), "r1-pa")


def upc_available():
    try:
        C.check_compat()
        return True
    except C.Unavailable:
        return False


@unittest.skipUnless(upc_available(), "the upc CLI is not installed")
class PublishEndToEndTests(RunCase):
    def test_publish_validates_and_is_idempotent(self):
        run = self.init_run()
        with mock.patch.object(M, "busy_seconds", return_value=0.1):
            rc, out = self.cli("code", "--run", str(self.run_dir), "--part", "C", "--ids", "a102,a104")
        self.assertEqual(rc, 0, out)
        corpus = self.root / "corpus"
        rc, out = self.cli("publish", "--run", str(self.run_dir), "--corpus", str(corpus), "--namespace", "forum")
        self.assertEqual(rc, 0, out)
        report = C.upc("validate", str(corpus))
        self.assertEqual(len(report["errors"]), 0, report["errors"][:3])
        sources = [json.loads(p.read_text()) for p in corpus.glob("sources/*/source.json")]
        self.assertEqual(len(sources), 4)                               # 2 posts + 2 actions, none merged
        self.assertEqual(len({s["source_id"] for s in sources}), 4)
        snapshot = {p: p.read_bytes() for p in corpus.rglob("*.jsonl") if "provenance" not in p.parts}
        rc, out = self.cli("publish", "--run", str(self.run_dir), "--corpus", str(corpus), "--namespace", "forum")
        self.assertEqual(rc, 0, out)
        self.assertEqual(snapshot, {p: p.read_bytes() for p in snapshot})
        items = [json.loads(x) for p in corpus.glob("extractions/*/items.jsonl") for x in p.read_text().splitlines()]
        self.assertTrue(items)
        self.assertTrue(all(any(s["type"] == "section" for s in o.get("secondary_locators", [])) for o in items))
        del run


if __name__ == "__main__":
    unittest.main()
